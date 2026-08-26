"""PDF (pymupdf) and Excel (openpyxl) export for catalog report previews."""

from __future__ import annotations

import html
import io
import re
from dataclasses import dataclass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas.report_catalog import ReportPreview
from app.services.reports.workbook_writer import HEADER_FILL, HEADER_FONT

# Keep this low enough that tests can force a second page with ~50 data rows.
PDF_ROWS_PER_PAGE = 28
_PAGE_WIDTH = 595.0
_PAGE_HEIGHT = 842.0
_MARGIN = 40.0
_ROW_HEIGHT = 16.0
_TITLE_SIZE = 13
_META_SIZE = 9
_CELL_SIZE = 8


@dataclass(frozen=True)
class ReportExportPayload:
    filename: str
    media_type: str
    body: bytes
    data_rows: int
    page_count: int = 1


def _safe_filename(title: str, suffix: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", title).strip("_").lower() or "report"
    return f"{slug}.{suffix}"


def _draw_pdf_text(page, rect: tuple[float, float, float, float], text: str, fontsize: float) -> None:
    """Draw a cell without crashing on unicode; htmlbox keeps CJK/accents readable."""
    value = str(text or "")
    try:
        page.insert_htmlbox(
            rect,
            (
                f'<div style="font-size:{fontsize}px;font-family:sans-serif;">'
                f"{html.escape(value)}</div>"
            ),
        )
    except Exception:
        page.insert_textbox(
            rect,
            value.encode("latin-1", "replace").decode("latin-1"),
            fontsize=fontsize,
            fontname="helv",
        )


def preview_to_xlsx(preview: ReportPreview) -> ReportExportPayload:
    wb = Workbook()
    ws = wb.active
    ws.title = (preview.title or "Report")[:31]
    ws.append([preview.title])
    title_cell = ws.cell(row=1, column=1)
    title_cell.font = Font(bold=True, size=14, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="1F6E7A")
    ws.append([preview.period_label, preview.currency, preview.notes or ""])
    ws.append([])
    ws.append(preview.columns)
    header_row = 4
    for col in range(1, len(preview.columns) + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in preview.rows:
        ws.append(row.cells)
        if row.emphasize:
            excel_row = ws.max_row
            for col in range(1, len(preview.columns) + 1):
                ws.cell(row=excel_row, column=col).font = Font(bold=True)
    for col_idx in range(1, max(len(preview.columns), 1) + 1):
        letter = get_column_letter(col_idx)
        max_len = 12
        for cell in ws[letter]:
            if cell.value is not None:
                max_len = max(max_len, min(len(str(cell.value)), 42))
        ws.column_dimensions[letter].width = max_len + 2
    buf = io.BytesIO()
    wb.save(buf)
    return ReportExportPayload(
        filename=_safe_filename(preview.title, "xlsx"),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        body=buf.getvalue(),
        data_rows=len(preview.rows),
        page_count=1,
    )


def _draw_header_row(page, y: float, columns: list[str], col_width: float) -> None:
    x = _MARGIN
    for col in columns:
        _draw_pdf_text(page, (x, y, x + col_width - 4, y + _ROW_HEIGHT), col, _CELL_SIZE)
        x += col_width


def preview_to_pdf(preview: ReportPreview) -> ReportExportPayload:
    import fitz

    columns = preview.columns or ["Value"]
    usable = _PAGE_WIDTH - 2 * _MARGIN
    col_width = usable / max(len(columns), 1)
    doc = fitz.open()

    def new_page():
        page = doc.new_page(width=_PAGE_WIDTH, height=_PAGE_HEIGHT)
        y = _MARGIN
        _draw_pdf_text(
            page,
            (_MARGIN, y, _PAGE_WIDTH - _MARGIN, y + _TITLE_SIZE + 4),
            preview.title,
            _TITLE_SIZE,
        )
        y += _TITLE_SIZE + 8
        meta = preview.period_label
        if preview.currency:
            meta = f"{meta}  ·  {preview.currency}"
        if preview.notes:
            meta = f"{meta}  ·  {preview.notes}"
        _draw_pdf_text(
            page,
            (_MARGIN, y, _PAGE_WIDTH - _MARGIN, y + _META_SIZE + 4),
            meta,
            _META_SIZE,
        )
        y += _META_SIZE + 14
        _draw_header_row(page, y, columns, col_width)
        y += _ROW_HEIGHT + 4
        return page, y

    page, y = new_page()
    rows_on_page = 0
    bottom = _PAGE_HEIGHT - _MARGIN

    for row in preview.rows:
        if rows_on_page >= PDF_ROWS_PER_PAGE or y + _ROW_HEIGHT > bottom:
            page, y = new_page()
            rows_on_page = 0
        x = _MARGIN
        cells = list(row.cells) + [""] * max(0, len(columns) - len(row.cells))
        for cell in cells[: len(columns)]:
            _draw_pdf_text(
                page,
                (x, y, x + col_width - 4, y + _ROW_HEIGHT),
                str(cell),
                _CELL_SIZE,
            )
            x += col_width
        y += _ROW_HEIGHT
        rows_on_page += 1

    if preview.empty and not preview.rows:
        _draw_pdf_text(
            page,
            (_MARGIN, y + 8, _PAGE_WIDTH - _MARGIN, y + 28),
            "No data for this period.",
            _META_SIZE,
        )

    body = doc.tobytes()
    page_count = doc.page_count
    doc.close()
    return ReportExportPayload(
        filename=_safe_filename(preview.title, "pdf"),
        media_type="application/pdf",
        body=body,
        data_rows=len(preview.rows),
        page_count=page_count,
    )


def export_preview(preview: ReportPreview, fmt: str) -> ReportExportPayload:
    if fmt == "pdf":
        return preview_to_pdf(preview)
    if fmt == "xlsx":
        return preview_to_xlsx(preview)
    raise ValueError(f"Unsupported export format: {fmt}")
