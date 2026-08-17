"""Styled Excel export for Team Expense reports (matches workbook_writer header style)."""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.department_budget import DepartmentBudgetUtilizationRow
from app.schemas.team_expense_reports import (
    EmployeeAdvanceDetailRow,
    EmployeeBudgetUtilizationRow,
    EmployeeExpenseSummaryRow,
    EmployeeMasterReportRow,
    EmployeeSpendDetailRow,
)
from app.services.master_data.department_budget_service import (
    build_department_budget_utilization_rows,
)
from app.services.reports.team_expense_reports_service import (
    build_budget_utilization_rows,
    build_employee_advance_detail_rows,
    build_employee_expense_summary_rows,
    build_employee_spend_detail_rows,
)

HEADER_FILL = PatternFill("solid", fgColor="1F6E7A")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FILL = PatternFill("solid", fgColor="1F6E7A")
TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
SUBTITLE_FONT = Font(size=9, color="334455")
BODY_FONT = Font(size=10)
ZEBRA_FILL = PatternFill("solid", fgColor="F3F7F8")

SHEET_BUDGET = "Employee Limits (Info)"
SHEET_DEPT = "GL Budgets"
SHEET_SUMMARY = "Employee Expense Summary"
SHEET_SPEND_DETAIL = "Employee Spend Detail"
SHEET_ADVANCE_DETAIL = "Employee Advance Detail"

EMPLOYEE_MASTER_HEADERS = [
    "Employee ID",
    "Employee Name",
    "Role",
    "Email",
    "Mobile",
    "Mobile (secondary)",
    "Viber",
    "Date of Joining",
    "Department",
    "Division",
    "Location",
    "Supervisor 1",
    "Supervisor 2",
    "Bank Name",
    "Bank Account Name",
    "Bank Account Number",
    "BSB",
    "SWIFT",
    "IBAN",
    "Advance Parent Ledger",
    "Advance Sub-Ledger",
    "Employee Status",
]

SPEND_DETAIL_HEADERS = [
    *EMPLOYEE_MASTER_HEADERS,
    "Main GL",
    "Sub-Ledger",
    "Sub-GL Budget",
    "Employee Spend (YTD)",
    "% of Sub-GL Used by Employee",
    "No. of Claims",
    "Advance Pending",
    "Cash Reimbursed YTD",
    "Last Claim Date",
    "Monthly Spending Limit",
    "Monthly Spent",
    "Monthly Remaining",
    "Monthly Utilization %",
    "Quarterly Spending Limit",
    "Quarterly Spent",
    "Quarterly Remaining",
    "Quarterly Utilization %",
    "Annual Spending Limit",
    "Annual Spent",
    "Annual Remaining",
    "Annual Utilization %",
]

ADVANCE_DETAIL_HEADERS = [
    *EMPLOYEE_MASTER_HEADERS,
    "Movement Type",
    "Document No.",
    "Document Date",
    "Took",
    "Used",
    "Outstanding After",
    "Pending Claims",
    "Available",
    "Cash Reimbursed",
    "Document Status",
    "Approved By",
    "Approved On",
]

BUDGET_HEADERS = [
    "Employee ID",
    "Employee Name",
    "Role",
    "Email",
    "Mobile",
    "Mobile (secondary)",
    "Viber",
    "Date of Joining",
    "Department",
    "Division",
    "Location",
    "Supervisor 1",
    "Supervisor 2",
    "Bank Name",
    "Bank Account Name",
    "Bank Account Number",
    "BSB",
    "SWIFT",
    "IBAN",
    "Employee Status",
    "Monthly Spending Limit",
    "Monthly Spent",
    "Monthly Remaining",
    "Monthly Utilization %",
    "Quarterly Spending Limit",
    "Quarterly Spent",
    "Quarterly Remaining",
    "Quarterly Utilization %",
    "Annual Spending Limit",
    "Annual Spent",
    "Annual Remaining",
    "Annual Utilization %",
    "Advance Float",
    "Monthly Cash Committed",
    "Monthly Cash Remaining",
    "Monthly Cash Utilization %",
    "Quarterly Cash Committed",
    "Quarterly Cash Remaining",
    "Quarterly Cash Utilization %",
    "Annual Cash Committed",
    "Annual Cash Remaining",
    "Annual Cash Utilization %",
    "Category Caps",
    "Claim Count",
    "Last Claim Date",
]

DEPT_BUDGET_HEADERS = [
    "Parent GL",
    "Sub-GL",
    "Period Kind",
    "Period Key",
    "Budget",
    "Spent So Far",
    "Left",
    "Utilization %",
    "Enforcement",
    "Notes",
]

SUMMARY_HEADERS = [
    "Employee ID",
    "Employee Name",
    "Role",
    "Email",
    "Mobile",
    "Department",
    "Division",
    "Location",
    "Document No.",
    "Invoice Date",
    "Expense Type",
    "Finance Role",
    "Document Type",
    "Line Item",
    "Qty",
    "Line Amount",
    "Currency",
    "Ledger Code",
    "Main GL",
    "Sub-Ledger",
    "Document Status",
    "Evaluation Status",
]


@dataclass(frozen=True)
class TeamExpenseExcelExport:
    xlsx_bytes: bytes
    filename: str
    data_rows: int


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _write_title_block(ws: Worksheet, title: str, subtitle: str, col_count: int) -> int:
    """Return the 1-based header row index after title + subtitle + spacer."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(col_count, 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(col_count, 1))
    sub_cell = ws.cell(row=2, column=1, value=subtitle)
    sub_cell.font = SUBTITLE_FONT
    sub_cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 28

    # Spacer row
    return 4


def _write_header_row(
    ws: Worksheet,
    headers: Sequence[str],
    header_row: int,
    *,
    freeze_col: int = 1,
) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30
    freeze_letter = get_column_letter(max(1, freeze_col))
    ws.freeze_panes = f"{freeze_letter}{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(headers))}{header_row}"
    )


def _autosize(ws: Worksheet, *, max_width: int = 36) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 60))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), max_width)


def _append_data_rows(
    ws: Worksheet,
    *,
    start_row: int,
    rows: Sequence[Sequence[Any]],
) -> int:
    row_idx = start_row
    for values in rows:
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col, value=_cell_value(value))
            cell.font = BODY_FONT
            if row_idx % 2 == 0:
                cell.fill = ZEBRA_FILL
        row_idx += 1
    return row_idx - start_row


def _master_values(row: EmployeeMasterReportRow) -> list[Any]:
    return [
        row.employee_id,
        row.name,
        row.role,
        row.email,
        row.whatsapp_number,
        row.whatsapp_number_2,
        row.viber_number or "",
        row.date_of_joining,
        row.department,
        row.division,
        row.location,
        row.supervisor_1,
        row.supervisor_2,
        row.bank_name,
        row.bank_account_name,
        row.bank_account_number,
        row.bank_bsb,
        row.bank_swift,
        row.bank_iban,
        row.advance_parent_ledger,
        row.advance_sub_ledger,
        row.status,
    ]


def _spend_detail_values(row: EmployeeSpendDetailRow) -> list[Any]:
    return [
        *_master_values(
            EmployeeMasterReportRow(
                employee_id=row.employee_id,
                name=row.name,
                role=row.role,
                email=row.email,
                whatsapp_number=row.whatsapp_number,
                whatsapp_number_2=row.whatsapp_number_2,
                viber_number=row.viber_number,
                date_of_joining=row.date_of_joining,
                department=row.department,
                location=row.location,
                division=row.division,
                supervisor_1=row.supervisor_1,
                supervisor_2=row.supervisor_2,
                bank_name=row.bank_name,
                bank_account_name=row.bank_account_name,
                bank_account_number=row.bank_account_number,
                bank_bsb=row.bank_bsb,
                bank_swift=row.bank_swift,
                bank_iban=row.bank_iban,
                advance_parent_ledger=row.advance_parent_ledger,
                advance_sub_ledger=row.advance_sub_ledger,
                status=row.status,
            )
        ),
        row.main_gl,
        row.sub_ledger,
        row.sub_gl_budget,
        row.employee_spend_ytd,
        row.pct_of_sub_gl_used if row.pct_of_sub_gl_used is not None else "",
        row.claim_count,
        row.advance_pending,
        row.cash_reimbursed_ytd,
        row.last_claim_date,
        row.budget_monthly,
        row.mtd_spent,
        row.monthly_remaining if row.monthly_remaining is not None else "",
        row.monthly_utilization_pct if row.monthly_utilization_pct is not None else "",
        row.budget_quarterly,
        row.qtd_spent,
        row.quarterly_remaining if row.quarterly_remaining is not None else "",
        row.quarterly_utilization_pct if row.quarterly_utilization_pct is not None else "",
        row.budget_annual,
        row.ytd_spent_total,
        row.annual_remaining if row.annual_remaining is not None else "",
        row.annual_utilization_pct if row.annual_utilization_pct is not None else "",
    ]


def _advance_detail_values(row: EmployeeAdvanceDetailRow) -> list[Any]:
    return [
        *_master_values(
            EmployeeMasterReportRow(
                employee_id=row.employee_id,
                name=row.name,
                role=row.role,
                email=row.email,
                whatsapp_number=row.whatsapp_number,
                whatsapp_number_2=row.whatsapp_number_2,
                viber_number=row.viber_number,
                date_of_joining=row.date_of_joining,
                department=row.department,
                location=row.location,
                division=row.division,
                supervisor_1=row.supervisor_1,
                supervisor_2=row.supervisor_2,
                bank_name=row.bank_name,
                bank_account_name=row.bank_account_name,
                bank_account_number=row.bank_account_number,
                bank_bsb=row.bank_bsb,
                bank_swift=row.bank_swift,
                bank_iban=row.bank_iban,
                advance_parent_ledger=row.advance_parent_ledger,
                advance_sub_ledger=row.advance_sub_ledger,
                status=row.status,
            )
        ),
        row.movement_type,
        row.document_no,
        row.document_date.isoformat() if row.document_date else "",
        row.took,
        row.used,
        row.outstanding_after,
        row.pending_claims,
        row.available,
        row.cash_reimbursed,
        row.document_status,
        row.approved_by,
        row.approved_on,
    ]


def _budget_values(row: EmployeeBudgetUtilizationRow) -> list[Any]:
    return [
        row.employee_id,
        row.name,
        row.role,
        row.email,
        row.whatsapp_number,
        row.whatsapp_number_2,
        row.viber_number or "",
        row.date_of_joining,
        row.department,
        row.division,
        row.location,
        row.supervisor_1,
        row.supervisor_2,
        row.bank_name,
        row.bank_account_name,
        row.bank_account_number,
        row.bank_bsb,
        row.bank_swift,
        row.bank_iban,
        row.status,
        row.budget_monthly,
        row.mtd_spent,
        row.monthly_remaining if row.monthly_remaining is not None else "",
        row.monthly_utilization_pct if row.monthly_utilization_pct is not None else "",
        row.budget_quarterly,
        row.qtd_spent,
        row.quarterly_remaining if row.quarterly_remaining is not None else "",
        row.quarterly_utilization_pct if row.quarterly_utilization_pct is not None else "",
        row.budget_annual,
        row.ytd_spent,
        row.annual_remaining if row.annual_remaining is not None else "",
        row.annual_utilization_pct if row.annual_utilization_pct is not None else "",
        row.advance_float,
        row.monthly_cash_committed,
        row.monthly_cash_remaining if row.monthly_cash_remaining is not None else "",
        row.monthly_cash_utilization_pct
        if row.monthly_cash_utilization_pct is not None
        else "",
        row.quarterly_cash_committed,
        row.quarterly_cash_remaining if row.quarterly_cash_remaining is not None else "",
        row.quarterly_cash_utilization_pct
        if row.quarterly_cash_utilization_pct is not None
        else "",
        row.annual_cash_committed,
        row.annual_cash_remaining if row.annual_cash_remaining is not None else "",
        row.annual_cash_utilization_pct
        if row.annual_cash_utilization_pct is not None
        else "",
        row.category_caps,
        row.claim_count,
        row.last_claim,
    ]


def _dept_budget_values(row: DepartmentBudgetUtilizationRow) -> list[Any]:
    return [
        row.gl_ledger,
        "",
        row.period_kind,
        row.period_key,
        row.allocated,
        row.consumed,
        row.remaining if row.remaining is not None else "",
        row.utilization_pct if row.utilization_pct is not None else "",
        getattr(row, "enforcement", None) or "soft",
        row.notes or "",
    ]


def _dept_budget_sub_values(
    parent: DepartmentBudgetUtilizationRow,
    sub_gl: str,
    consumed: float,
    pct: float,
) -> list[Any]:
    return [
        parent.gl_ledger,
        sub_gl,
        parent.period_kind,
        parent.period_key,
        "",
        consumed,
        "",
        pct,
        getattr(parent, "enforcement", None) or "soft",
        "",
    ]


def _kind_label(kind: str) -> str:
    mapping = {
        "advance_requisition": "Advance requisition",
        "expense_claim": "Expense claim",
        "direct_payment": "Direct payment",
        # Legacy stored value — display as claim.
        "expense_against_advance": "Expense claim",
    }
    return mapping.get((kind or "").strip(), kind or "")


def _finance_role(kind: str) -> str:
    """How finance should treat the document in books."""
    token = (kind or "").strip().lower()
    if token == "advance_requisition":
        return "Balance-sheet float (no GL budget)"
    if token == "direct_payment":
        return "Company direct spend (no GL budget / no advance netting)"
    return "P&L / GL budget spend"


def _status_label(status: str) -> str:
    token = (status or "").strip()
    if not token:
        return ""
    return token.replace("_", " ").title()


def _summary_values(row: EmployeeExpenseSummaryRow) -> list[Any]:
    return [
        row.employee_id,
        row.employee_name,
        row.role,
        row.employee_email,
        row.mobile,
        row.department,
        row.division,
        row.location,
        row.document_no,
        row.invoice_date.isoformat() if row.invoice_date else "",
        _kind_label(row.team_expense_kind),
        _finance_role(row.team_expense_kind),
        row.document_type_code,
        row.line_description,
        row.line_qty if row.line_qty is not None else "",
        row.line_amount if row.line_amount is not None else "",
        row.currency,
        row.ledger_code,
        row.main_gl,
        row.sub_ledger,
        _status_label(row.status),
        row.evaluation_status,
    ]


def _slug_filename(prefix: str, tenant_slug: str) -> str:
    safe = re.sub(r"[^a-zA-Z0-9_-]+", "-", (tenant_slug or "tenant").strip()) or "tenant"
    return f"{prefix}_{safe}.xlsx"


def _write_sheet(
    ws: Worksheet,
    *,
    title: str,
    subtitle: str,
    headers: Sequence[str],
    data_rows: Sequence[Sequence[Any]],
    freeze_col: int = 1,
) -> int:
    header_row = _write_title_block(ws, title, subtitle, len(headers))
    _write_header_row(ws, headers, header_row, freeze_col=freeze_col)
    count = _append_data_rows(ws, start_row=header_row + 1, rows=data_rows)
    if count > 0:
        last = header_row + count
        ws.auto_filter.ref = (
            f"A{header_row}:{get_column_letter(len(headers))}{last}"
        )
    _autosize(ws)
    ws.sheet_view.showGridLines = False
    return count


async def build_team_expense_excel_export(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    *,
    report: str,
    tenant_slug: str = "tenant",
    date_from: date | None = None,
    date_to: date | None = None,
) -> TeamExpenseExcelExport:
    """
    Build a styled workbook for Team Expense finance reporting.

    report:
      department-budget-utilization | expense-summary |
      budget-utilization | spending-limit-utilization |
      employee-spend-detail | employee-advance-detail
    """
    wb = Workbook()
    default = wb.active
    wb.remove(default)

    total_rows = 0
    kind = (report or "").strip().lower()

    if kind == "employee-advance-detail":
        detail_rows = await build_employee_advance_detail_rows(session, tenant_id)
        ws = wb.create_sheet(SHEET_ADVANCE_DETAIL)
        total_rows += _write_sheet(
            ws,
            title="Employee Advance Detail",
            subtitle=(
                "Employee Staff Advance movement ledger (not GL budget). "
                "Advance row = Took; Claim row = Used when a claim nets float. "
                "Outstanding After = running float; Pending Claims / Available = current snapshot."
            ),
            headers=ADVANCE_DETAIL_HEADERS,
            data_rows=[_advance_detail_values(r) for r in detail_rows],
        )
        return TeamExpenseExcelExport(
            xlsx_bytes=_workbook_bytes(wb),
            filename=_slug_filename("te_employee_advance_detail", tenant_slug),
            data_rows=total_rows,
        )

    if kind == "employee-spend-detail":
        spend_rows = await build_employee_spend_detail_rows(session, tenant_id)
        ws_spend = wb.create_sheet(SHEET_SPEND_DETAIL)
        total_rows += _write_sheet(
            ws_spend,
            title="Employee Spend Detail",
            subtitle=(
                "One row per employee per expense Sub-GL. "
                "A–V = employee master; then Sub-GL spend vs GL budget; "
                "then employee spending limits (MTD/QTD/YTD claim spend, advances excluded). "
                "Cash Reimbursed = settlement credits on claims."
            ),
            headers=SPEND_DETAIL_HEADERS,
            data_rows=[_spend_detail_values(r) for r in spend_rows],
        )
        return TeamExpenseExcelExport(
            xlsx_bytes=_workbook_bytes(wb),
            filename=_slug_filename("te_employee_spend_detail", tenant_slug),
            data_rows=total_rows,
        )

    if kind == "department-budget-utilization":
        gl_rows = await build_department_budget_utilization_rows(session, tenant_id)
        ws = wb.create_sheet(SHEET_DEPT)
        total_rows += _write_sheet(
            ws,
            title="GL Budget Utilization",
            subtitle=(
                "GL account budgets (finance control). "
                "Expense claims consume budget on the full claim amount. Advances do not. "
                "Enforcement soft = manager can approve overruns; hard = raise budget first. "
                "Sub-GL spend rolls into the parent wallet."
            ),
            headers=DEPT_BUDGET_HEADERS,
            data_rows=[
                value
                for r in gl_rows
                for value in (
                    [_dept_budget_values(r)]
                    + [
                        _dept_budget_sub_values(
                            r,
                            sub.gl_ledger,
                            float(sub.consumed),
                            float(sub.pct_of_budget),
                        )
                        for sub in (r.sub_breakdown or [])
                    ]
                )
            ],
        )
        return TeamExpenseExcelExport(
            xlsx_bytes=_workbook_bytes(wb),
            filename=_slug_filename("te_gl_budget_utilization", tenant_slug),
            data_rows=total_rows,
        )

    if kind == "expense-summary":
        summary_rows = await build_employee_expense_summary_rows(
            session, tenant_id, date_from=date_from, date_to=date_to
        )
        period = "All dates"
        if date_from and date_to:
            period = f"{date_from.isoformat()} to {date_to.isoformat()}"
        elif date_from:
            period = f"From {date_from.isoformat()}"
        elif date_to:
            period = f"Through {date_to.isoformat()}"
        ws = wb.create_sheet(SHEET_SUMMARY)
        total_rows += _write_sheet(
            ws,
            title="Employee Expense Summary",
            subtitle=(
                f"One row per team expense document line. "
                f"Claims = P&L / GL budget spend (full claim amount). "
                f"Advances = employee float only (no GL budget). Period: {period}."
            ),
            headers=SUMMARY_HEADERS,
            data_rows=[_summary_values(r) for r in summary_rows],
        )
        return TeamExpenseExcelExport(
            xlsx_bytes=_workbook_bytes(wb),
            filename=_slug_filename("te_employee_expense_summary", tenant_slug),
            data_rows=total_rows,
        )

    if kind in {"budget-utilization", "spending-limit-utilization"}:
        # Informational only — employee limits are not the finance budget control.
        rows = await build_budget_utilization_rows(session, tenant_id)
        ws = wb.create_sheet(SHEET_BUDGET)
        count = _write_sheet(
            ws,
            title="Employee Spending Limits (Informational)",
            subtitle=(
                "NOT the primary finance control. GL Budgets enforce claim spend. "
                "This sheet shows optional employee spending-limit counters vs claim MTD/QTD/YTD. "
                "Advances do not count as spend."
            ),
            headers=BUDGET_HEADERS,
            data_rows=[_budget_values(r) for r in rows],
        )
        return TeamExpenseExcelExport(
            xlsx_bytes=_workbook_bytes(wb),
            filename=_slug_filename("te_employee_limits_info", tenant_slug),
            data_rows=count,
        )

    raise ValueError(f"Unknown team expense report: {report}")


def _workbook_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
