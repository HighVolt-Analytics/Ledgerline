"""Excel export — documents bundle matrix for auditors.

Single sheet — Documents Bundle
    Rows: transactional posting anchors (Rule Book DT matrix) plus
    vision soft-bundle invoice-family anchors (vault type-folder matrix).

Fixed columns
    Timestamp, Uploaded by, Source, DT type, Invoice date,
    Counterparty, Total, Currency, Linkage, Invoice no. (vault hyperlink),
    Proforma invoice no., PO reference, SO reference, Universal match.

Dynamic columns
    Org Rule Book document types (DT codes), then vault type-folder names
    for soft-bundle siblings (Air Waybill, Packing List, …).
"""

from __future__ import annotations

import io
import re
import uuid
from dataclasses import dataclass
from datetime import date, timezone, tzinfo
from typing import Literal
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.invoice import Invoice, InvoiceStatus
from app.schemas.document_type import DocumentTypeDefinition
from app.schemas.dossier import DossierLinkedDocumentResponse, DossierLinkedDocumentsResponse
from app.services.audit.audit_export_service import (
    collect_linked_doc_entries,
    excel_hyperlink,
    format_vault_links,
    linked_docs_by_dt_code,
    vault_view_path,
)
from app.services.classification.document_type_catalog import (
    org_document_types_for_bundle_export,
)
from app.services.classification.document_type_klass import (
    is_trans_posting,
    normalize_document_type_identity,
)
from app.services.classification.document_type_playbook_service import (
    resolve_definition_for_invoice,
)
from app.services.dossier.dossier_linked_documents_service import build_dossier_linked_documents
from app.services.dossier.dossier_service import build_linkage_sibling_cache
from app.services.dossier.vision_bundle_linkage import should_use_vision_bundle_linkage
from app.services.extraction.document_heading_utils import is_invoice_family_vault_label
from app.services.invoice.invoice_evaluation_service import load_posting_config_for_tenant
from app.services.vault.vault_invoice_paths import vault_type_folder_from_llm_or_heading
from app.services.vault.vault_paths import is_standard_vault_book, normalize_vault_type_book_label

BundleCellFormat = Literal["excel", "plain"]

_BUNDLE_EXPORT_STATUSES = frozenset({
    InvoiceStatus.PROCESSED,
    InvoiceStatus.EXCEPTION,
    InvoiceStatus.VALIDATING,
    InvoiceStatus.MAPPING,
    InvoiceStatus.JOURNALING,
    InvoiceStatus.RECONCILING,
})

_FIXED_COLUMNS = [
    "Timestamp",
    "Uploaded by",
    "Source",
    "DT type",
    "Invoice date",
    "Counterparty",
    "Total",
    "Currency",
    "Linkage",
    "Invoice no.",
    "Proforma invoice no.",
    "PO reference",
    "SO reference",
    "Universal match",
]

_MATCH_COLUMNS = frozenset({"Universal match"})

_TITLE_FILL = PatternFill("solid", fgColor="1F6E7A")
_TITLE_FONT = Font(color="FFFFFF", bold=True, size=14)
_HEADER_FILL = PatternFill("solid", fgColor="1F6E7A")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_ZEBRA_FILL = PatternFill("solid", fgColor="F3F7F8")
_MISSING_FILL = PatternFill("solid", fgColor="F8D7DA")
_ADVISORY_FILL = PatternFill("solid", fgColor="FFF3CD")
_YES_FILL = PatternFill("solid", fgColor="D1E7DD")
_NO_FILL = PatternFill("solid", fgColor="E9ECEF")
_LINK_FONT = Font(color="0563C1", underline="single", size=10)
_BODY_FONT = Font(size=10)
_LEGEND_FONT = Font(size=9, color="334455")

_SINGLE_HYPERLINK_RE = re.compile(
    r'^=HYPERLINK\("((?:[^"]|"")*)","((?:[^"]|"")*)"\)$'
)
_HYPERLINK_PART_RE = re.compile(
    r'HYPERLINK\("((?:[^"]|"")*)","((?:[^"]|"")*)"\)'
)

_TITLE_ROW = 1
_LEGEND_ROW = 2
_HEADER_ROW = 4
_FIRST_DATA_ROW = 5

_SHEET_POSTING = "Documents Bundle"


@dataclass(frozen=True)
class DocumentsBundleExportPayload:
    xlsx_bytes: bytes
    filename: str
    data_rows: int


@dataclass(frozen=True)
class _AnchorExportContext:
    invoice: Invoice
    definition: DocumentTypeDefinition
    linked: DossierLinkedDocumentsResponse
    invoice_dt_code_by_id: dict[int, str]


@dataclass(frozen=True)
class _UnderstoodExportContext:
    invoice: Invoice
    linked: DossierLinkedDocumentsResponse
    vault_folder_by_invoice_id: dict[int, str]


def _effective_invoice_date():
    return func.coalesce(Invoice.invoice_date, func.date(Invoice.created_at))


async def _load_bundle_export_invoices(
    db: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[Invoice]:
    """Load anchor candidates: active workflow statuses within the date window."""
    stmt = select(Invoice).where(
        Invoice.tenant_id == tenant_id,
        Invoice.status.in_(_BUNDLE_EXPORT_STATUSES),
    )
    effective = _effective_invoice_date()
    if date_from is not None:
        stmt = stmt.where(effective >= date_from)
    if date_to is not None:
        stmt = stmt.where(effective <= date_to)
    stmt = stmt.order_by(effective.desc(), Invoice.id.desc())
    return list((await db.execute(stmt)).scalars().all())


def documents_bundle_filename(*, date_from: date | None) -> str:
    if date_from is not None:
        return f"documents_bundle_{date_from.year:04d}-{date_from.month:02d}.xlsx"
    return "documents_bundle.xlsx"


def _period_label(*, date_from: date | None, date_to: date | None) -> str:
    if date_from is not None and date_to is not None:
        if date_from == date_to:
            return date_from.isoformat()
        return f"{date_from.isoformat()} → {date_to.isoformat()}"
    if date_from is not None:
        return f"From {date_from.isoformat()}"
    if date_to is not None:
        return f"Through {date_to.isoformat()}"
    return "All dates"


def _is_transactional_posting(defn: DocumentTypeDefinition | None) -> bool:
    if defn is None:
        return False
    _, posting = normalize_document_type_identity(defn)
    return is_trans_posting(defn) and posting == "Yes"


def _dt_type_label(defn: DocumentTypeDefinition | None) -> str:
    if defn is None:
        return "Unclassified"
    short = (defn.short_title or "").strip()
    if short:
        return short
    return (defn.title or defn.code or "Unclassified").strip() or "Unclassified"


def vault_folder_label_for_export(invoice: Invoice) -> str:
    """Vault type-folder label for understood-path matrix columns / DT type."""
    folder = vault_type_folder_from_llm_or_heading(invoice)
    if folder:
        return folder
    route = (invoice.route_target or "").strip()
    if route and not is_standard_vault_book(route):
        return normalize_vault_type_book_label(route) or route
    return ""


def is_understood_invoice_family_anchor(invoice: Invoice) -> bool:
    """True when this document is a vision soft-bundle invoice-family row anchor."""
    if not should_use_vision_bundle_linkage(invoice):
        return False
    return is_invoice_family_vault_label(vault_folder_label_for_export(invoice))


def _yes_no(value: bool) -> str:
    return "Yes" if value else "No"


def _match_flags(linked: DossierLinkedDocumentsResponse) -> tuple[str, str, str]:
    status = (linked.match_summary.status if linked.match_summary else "").strip()
    two_way = _yes_no(status == "2-Way Match")
    three_way = _yes_no(status == "3-Way Match")

    universal = any(
        doc.link_kind == "invoice_no" and doc.present and doc.invoice_id is not None
        for doc in linked.documents
        if not doc.is_anchor
    )

    return two_way, three_way, _yes_no(universal)


def _invoice_date_label(invoice: Invoice) -> str:
    if invoice.invoice_date is not None:
        return invoice.invoice_date.isoformat()
    if invoice.created_at is not None:
        return invoice.created_at.date().isoformat()
    return ""


def resolve_display_timezone(tz_name: str | None) -> tzinfo:
    """IANA zone from the browser; invalid/missing → UTC (DB stays UTC)."""
    token = (tz_name or "").strip()
    if not token:
        return timezone.utc
    try:
        return ZoneInfo(token)
    except (ZoneInfoNotFoundError, ValueError, TypeError, KeyError):
        return timezone.utc


def _timestamp_label(
    invoice: Invoice,
    *,
    display_tz: tzinfo | None = None,
) -> str:
    """When the document was received / uploaded, shown in browser/local TZ."""
    if invoice.created_at is None:
        return ""
    created = invoice.created_at
    if created.tzinfo is None:
        created = created.replace(tzinfo=timezone.utc)
    else:
        created = created.astimezone(timezone.utc)
    local = created.astimezone(display_tz or timezone.utc)
    return local.strftime("%Y-%m-%d %H:%M:%S")


def _source_label(invoice: Invoice) -> str:
    src = (invoice.capture_source or "").strip().lower()
    if src == "email":
        return "Email"
    if src == "whatsapp":
        return "WhatsApp"
    if src == "viber":
        return "Viber"
    if src == "upload":
        return "Direct upload"
    sender = (invoice.email_sender or "").lower()
    if "onedrive" in sender or "sharepoint" in sender:
        return "OneDrive"
    if invoice.connected_mailbox_id or invoice.email_sender:
        return "Email"
    return "Direct upload"


def _uploaded_by_label(invoice: Invoice, *, upload_actor: str | None = None) -> str:
    """Manual uploader for direct upload; email/messaging sender otherwise."""
    name = (invoice.uploaded_by_name or "").strip()
    email = (invoice.uploaded_by_email or "").strip()
    if name or email:
        return name or email
    sender = (invoice.email_sender or "").strip()
    if sender:
        return sender
    return (upload_actor or "").strip()


def _total_label(invoice: Invoice) -> str:
    if invoice.total is not None:
        return str(invoice.total)
    if invoice.subtotal is not None:
        return str(invoice.subtotal)
    return ""


def _invoice_no_cell(invoice: Invoice, *, cell_format: BundleCellFormat) -> str:
    label = (invoice.invoice_no or "").strip()
    if not label:
        return ""
    url = vault_view_path(invoice.id)
    if cell_format == "plain":
        return f"{label} | {url}" if url else label
    return excel_hyperlink(url, label)


def _proforma_invoice_no_label(invoice: Invoice) -> str:
    fields = invoice.extracted_fields if isinstance(invoice.extracted_fields, dict) else {}
    raw = fields.get("proforma_invoice_no")
    if isinstance(raw, str):
        return raw.strip()
    if raw is None:
        return ""
    return str(raw).strip()


def _bundle_slots_by_dt_code(
    linked: DossierLinkedDocumentsResponse,
) -> dict[str, DossierLinkedDocumentResponse]:
    slots: dict[str, DossierLinkedDocumentResponse] = {}
    for doc in linked.documents:
        if doc.is_anchor:
            continue
        code = (doc.document_type_code or "").strip().upper()
        if not code:
            continue
        existing = slots.get(code)
        if existing is None:
            slots[code] = doc
            continue
        if doc.requirement == "mandatory" and existing.requirement != "mandatory":
            slots[code] = doc
    return slots


def _manual_link_cell(
    manual_link,
    *,
    cell_format: BundleCellFormat,
) -> str:
    ref = (manual_link.document_ref or "").strip() or f"DOC-{manual_link.invoice_id}"
    url = vault_view_path(manual_link.invoice_id)
    if cell_format == "plain":
        return f"{ref} | {url}" if url else ref
    return excel_hyperlink(url, ref)


def bundle_dt_cells_by_code(
    anchor_invoice_id: int,
    linked: DossierLinkedDocumentsResponse,
    dt_codes: list[str],
    *,
    invoice_dt_code_by_id: dict[int, str] | None = None,
    cell_format: BundleCellFormat = "excel",
) -> dict[str, str]:
    """DT column cells with hyperlinks/plain links plus Missing/Advisory slot markers."""
    linked_cells = linked_docs_by_dt_code(
        anchor_invoice_id,
        linked,
        label_fn="bundle",
        invoice_dt_code_by_id=invoice_dt_code_by_id,
        cell_format=cell_format,
    )
    slots_by_dt = _bundle_slots_by_dt_code(linked)

    out: dict[str, str] = {}
    for code in dt_codes:
        if linked_cells.get(code):
            out[code] = linked_cells[code]
            continue
        slot = slots_by_dt.get(code)
        if slot is None:
            out[code] = ""
        elif slot.manual_link is not None:
            out[code] = _manual_link_cell(slot.manual_link, cell_format=cell_format)
        elif slot.present:
            out[code] = linked_cells.get(code, "")
        elif slot.requirement == "mandatory":
            out[code] = "Missing"
        elif slot.requirement == "advisory":
            out[code] = "Advisory"
        else:
            out[code] = ""
    return out


def _export_link_label(*, document_ref: str | None, invoice_no: str | None, invoice_id: int) -> str:
    ref = (document_ref or "").strip()
    if ref:
        return ref
    inv_no = (invoice_no or "").strip()
    if inv_no:
        return inv_no
    return f"DOC-{invoice_id}"


def bundle_vault_folder_cells(
    anchor_invoice_id: int,
    linked: DossierLinkedDocumentsResponse,
    vault_folders: list[str],
    *,
    vault_folder_by_invoice_id: dict[int, str],
    cell_format: BundleCellFormat = "excel",
) -> dict[str, str]:
    """Vault-folder column cells for understood soft-bundle siblings (no Missing/Advisory)."""
    by_folder: dict[str, list[tuple[str, str]]] = {}
    for entry in collect_linked_doc_entries(anchor_invoice_id, linked):
        folder = (vault_folder_by_invoice_id.get(entry.invoice_id) or "").strip()
        if not folder:
            continue
        url = vault_view_path(entry.invoice_id)
        label = _export_link_label(
            document_ref=entry.document_ref,
            invoice_no=entry.invoice_no,
            invoice_id=entry.invoice_id,
        )
        by_folder.setdefault(folder, []).append((url, label))

    return {
        folder: (
            format_vault_links(by_folder.get(folder, []), cell_format=cell_format)
            if by_folder.get(folder)
            else ""
        )
        for folder in vault_folders
    }


async def _invoice_dt_code_by_id(
    db: AsyncSession,
    linked: DossierLinkedDocumentsResponse,
) -> dict[int, str]:
    invoice_ids = {
        doc.invoice_id
        for doc in linked.documents
        if doc.invoice_id is not None and not doc.is_anchor
    }
    for doc in linked.documents:
        if doc.manual_link is not None:
            invoice_ids.add(doc.manual_link.invoice_id)
    if not invoice_ids:
        return {}
    rows = (
        await db.execute(
            select(Invoice.id, Invoice.document_type_code).where(
                Invoice.id.in_(invoice_ids)
            )
        )
    ).all()
    return {
        row.id: (row.document_type_code or "").strip().upper()
        for row in rows
        if (row.document_type_code or "").strip()
    }


async def _invoices_by_id(
    db: AsyncSession,
    invoice_ids: set[int],
) -> dict[int, Invoice]:
    if not invoice_ids:
        return {}
    rows = (
        await db.execute(select(Invoice).where(Invoice.id.in_(invoice_ids)))
    ).scalars().all()
    return {row.id: row for row in rows}


def _sorted_tenant_document_types(
    document_types: list[DocumentTypeDefinition],
) -> list[DocumentTypeDefinition]:
    return sorted(
        document_types,
        key=lambda row: (row.code or "").upper(),
    )


def _dt_column_headers(document_types: list[DocumentTypeDefinition]) -> list[str]:
    headers: list[str] = []
    seen: set[str] = set()
    for row in _sorted_tenant_document_types(document_types):
        label = (row.short_title or row.title or row.code).strip() or row.code
        if label in seen:
            label = f"{label} ({row.code})"
        seen.add(label)
        headers.append(label)
    return headers


def _dt_codes_ordered(document_types: list[DocumentTypeDefinition]) -> list[str]:
    return [row.code.upper() for row in _sorted_tenant_document_types(document_types)]


def build_documents_bundle_row(
    invoice: Invoice,
    *,
    definition: DocumentTypeDefinition | None,
    linked: DossierLinkedDocumentsResponse,
    dt_codes: list[str],
    invoice_dt_code_by_id: dict[int, str] | None = None,
    cell_format: BundleCellFormat = "excel",
    upload_actor: str | None = None,
    display_tz: tzinfo | None = None,
) -> list[str]:
    _, _, universal = _match_flags(linked)
    by_dt = bundle_dt_cells_by_code(
        invoice.id,
        linked,
        dt_codes,
        invoice_dt_code_by_id=invoice_dt_code_by_id,
        cell_format=cell_format,
    )

    row: list[str] = [
        _timestamp_label(invoice, display_tz=display_tz),
        _uploaded_by_label(invoice, upload_actor=upload_actor),
        _source_label(invoice),
        _dt_type_label(definition),
        _invoice_date_label(invoice),
        (invoice.vendor or "").strip(),
        _total_label(invoice),
        (invoice.currency or "").strip(),
        (linked.linkage_kind or "").strip(),
        _invoice_no_cell(invoice, cell_format=cell_format),
        _proforma_invoice_no_label(invoice),
        (invoice.po_reference or "").strip(),
        (invoice.so_reference or "").strip(),
        universal,
    ]
    row.extend(by_dt.get(code, "") for code in dt_codes)
    return row


def build_understood_bundle_row(
    invoice: Invoice,
    *,
    linked: DossierLinkedDocumentsResponse,
    vault_folders: list[str],
    vault_folder_by_invoice_id: dict[int, str],
    cell_format: BundleCellFormat = "excel",
    upload_actor: str | None = None,
    display_tz: tzinfo | None = None,
) -> list[str]:
    _, _, universal = _match_flags(linked)
    by_folder = bundle_vault_folder_cells(
        invoice.id,
        linked,
        vault_folders,
        vault_folder_by_invoice_id=vault_folder_by_invoice_id,
        cell_format=cell_format,
    )
    type_label = vault_folder_label_for_export(invoice) or "Unclassified"

    row: list[str] = [
        _timestamp_label(invoice, display_tz=display_tz),
        _uploaded_by_label(invoice, upload_actor=upload_actor),
        _source_label(invoice),
        type_label,
        _invoice_date_label(invoice),
        (invoice.vendor or "").strip(),
        _total_label(invoice),
        (invoice.currency or "").strip(),
        (linked.linkage_kind or "").strip(),
        _invoice_no_cell(invoice, cell_format=cell_format),
        _proforma_invoice_no_label(invoice),
        (invoice.po_reference or "").strip(),
        (invoice.so_reference or "").strip(),
        universal,
    ]
    row.extend(by_folder.get(folder, "") for folder in vault_folders)
    return row


async def _upload_actors_by_invoice_id(
    db: AsyncSession,
    invoice_ids: list[int],
) -> dict[int, str]:
    """Map invoice id → uploader from invoice_uploaded audit (name preferred, else email)."""
    if not invoice_ids:
        return {}
    from app.models.audit import AuditLog

    rows = (
        await db.execute(
            select(AuditLog)
            .where(
                AuditLog.invoice_id.in_(invoice_ids),
                AuditLog.event == "invoice_uploaded",
            )
            .order_by(AuditLog.created_at.asc())
        )
    ).scalars().all()

    actors: dict[int, str] = {}
    for log in rows:
        if log.invoice_id is None or log.invoice_id in actors:
            continue
        detail = log.detail if isinstance(log.detail, dict) else {}
        name = str(detail.get("actor_name") or "").strip()
        email = str(detail.get("actor_email") or "").strip()
        label = name or email
        if label:
            actors[log.invoice_id] = label
    return actors


def _unescape_excel_csv(value: str) -> str:
    return value.replace('""', '"')


def _parse_single_hyperlink_formula(value: str) -> tuple[str, str] | None:
    match = _SINGLE_HYPERLINK_RE.match((value or "").strip())
    if match is None:
        return None
    url = _unescape_excel_csv(match.group(1))
    label = _unescape_excel_csv(match.group(2))
    return url, label


def _parse_hyperlink_entries(value: str) -> list[tuple[str, str]]:
    """Parse one or more HYPERLINK(...) parts from an Excel formula cell."""
    raw = (value or "").strip()
    if not raw.startswith("=") or "HYPERLINK(" not in raw:
        return []
    return [
        (_unescape_excel_csv(url), _unescape_excel_csv(label))
        for url, label in _HYPERLINK_PART_RE.findall(raw)
    ]


def _expand_cell_slots(raw_value: str, link_mode: BundleCellFormat) -> list[str]:
    """Split multi-doc link cells into one single-link value per slot (same row)."""
    raw = raw_value or ""
    if link_mode == "excel":
        entries = _parse_hyperlink_entries(raw)
        if len(entries) >= 2:
            return [excel_hyperlink(url, label) for url, label in entries]
        return [raw]
    lines = [line.strip() for line in raw.split("\n") if line.strip()]
    if len(lines) >= 2 and all(" | " in line for line in lines):
        return lines
    return [raw]


def _logical_column_widths(
    rows: list[list[str]],
    *,
    headers: list[str],
    link_mode: BundleCellFormat,
) -> list[int]:
    widths = [1] * len(headers)
    for row in rows:
        for idx in range(len(headers)):
            raw = row[idx] if idx < len(row) else ""
            widths[idx] = max(widths[idx], len(_expand_cell_slots(raw, link_mode)))
    return widths


def _physical_header_labels(headers: list[str], widths: list[int]) -> list[str]:
    labels: list[str] = []
    for header, width in zip(headers, widths):
        for slot in range(width):
            labels.append(header if slot == 0 else f"{header} ({slot + 1})")
    return labels


def _autosize_columns(ws: Worksheet, *, max_width: int = 36) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx, min_row=_HEADER_ROW):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, min(len(str(cell.value)), 48))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 11), max_width)


def _write_title_and_period(
    ws: Worksheet,
    *,
    title: str,
    col_count: int,
    date_from: date | None,
    date_to: date | None,
) -> None:
    ws.merge_cells(start_row=_TITLE_ROW, start_column=1, end_row=_TITLE_ROW, end_column=max(col_count, 1))
    title_cell = ws.cell(row=_TITLE_ROW, column=1, value=title)
    title_cell.fill = _TITLE_FILL
    title_cell.font = _TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    for col in range(2, col_count + 1):
        cell = ws.cell(row=_TITLE_ROW, column=col)
        cell.fill = _TITLE_FILL
    ws.row_dimensions[_TITLE_ROW].height = 28

    period = _period_label(date_from=date_from, date_to=date_to)
    ws.cell(row=_LEGEND_ROW, column=1, value=f"Period: {period}").font = _LEGEND_FONT
    ws.row_dimensions[_LEGEND_ROW].height = 18

    # Spacer row keeps title/period visually separate from the table header.
    ws.row_dimensions[3].height = 8


def _apply_data_cell_style(
    cell,
    *,
    header: str,
    raw_value: str,
    zebra: bool,
    link_mode: BundleCellFormat,
) -> None:
    cell.font = _BODY_FONT
    cell.alignment = Alignment(vertical="center", wrap_text=True)

    value = raw_value
    if link_mode == "excel":
        parsed = _parse_single_hyperlink_formula(raw_value)
        if parsed is not None:
            url, label = parsed
            cell.value = label
            cell.hyperlink = url
            cell.font = _LINK_FONT
            value = label
        else:
            cell.value = raw_value
    else:
        cell.value = raw_value
        if " | " in raw_value and "/vault?" in raw_value:
            cell.font = _LINK_FONT

    fill = None
    if value == "Missing":
        fill = _MISSING_FILL
    elif value == "Advisory":
        fill = _ADVISORY_FILL
    elif header in _MATCH_COLUMNS:
        if value == "Yes":
            fill = _YES_FILL
        elif value == "No":
            fill = _NO_FILL
    elif zebra:
        fill = _ZEBRA_FILL

    if fill is not None:
        cell.fill = fill


def _write_bundle_sheet(
    ws: Worksheet,
    *,
    title: str,
    rows: list[list[str]],
    dynamic_headers: list[str],
    date_from: date | None,
    date_to: date | None,
    cell_format: BundleCellFormat,
) -> None:
    headers = [*_FIXED_COLUMNS, *dynamic_headers]
    widths = _logical_column_widths(rows, headers=headers, link_mode=cell_format)
    physical_headers = _physical_header_labels(headers, widths)
    col_count = len(physical_headers)

    _write_title_and_period(
        ws,
        title=title,
        col_count=col_count,
        date_from=date_from,
        date_to=date_to,
    )

    for col_idx, header in enumerate(physical_headers, start=1):
        cell = ws.cell(row=_HEADER_ROW, column=col_idx, value=header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[_HEADER_ROW].height = 24

    for row_offset, row in enumerate(rows):
        excel_row = _FIRST_DATA_ROW + row_offset
        zebra = row_offset % 2 == 1
        physical_col = 1
        for logical_idx, header in enumerate(headers):
            raw = row[logical_idx] if logical_idx < len(row) else ""
            slots = _expand_cell_slots(raw, cell_format)
            width = widths[logical_idx]
            if len(slots) < width:
                slots = [*slots, *[""] * (width - len(slots))]
            for slot_value in slots[:width]:
                cell = ws.cell(row=excel_row, column=physical_col)
                _apply_data_cell_style(
                    cell,
                    header=header,
                    raw_value=slot_value,
                    zebra=zebra,
                    link_mode=cell_format,
                )
                physical_col += 1

    ws.freeze_panes = ws.cell(row=_FIRST_DATA_ROW, column=1)
    if rows:
        last_col = get_column_letter(col_count)
        last_row = _FIRST_DATA_ROW + len(rows) - 1
        ws.auto_filter.ref = f"A{_HEADER_ROW}:{last_col}{last_row}"
    else:
        last_col = get_column_letter(max(col_count, 1))
        ws.auto_filter.ref = f"A{_HEADER_ROW}:{last_col}{_HEADER_ROW}"

    _autosize_columns(ws)
    ws.sheet_view.showGridLines = False


def documents_bundle_rows_to_xlsx(
    rows: list[list[str]],
    *,
    dt_column_headers: list[str],
    date_from: date | None = None,
    date_to: date | None = None,
    cell_format: BundleCellFormat = "excel",
    understood_rows: list[list[str]] | None = None,
    vault_folder_headers: list[str] | None = None,
) -> bytes:
    """Build a single-sheet workbook: posting + understood rows together."""
    vault_headers = list(vault_folder_headers or [])
    understood = list(understood_rows or [])
    n_dt = len(dt_column_headers)
    n_vault = len(vault_headers)
    combined_headers = [*dt_column_headers, *vault_headers]

    combined_rows: list[list[str]] = []
    fixed_n = len(_FIXED_COLUMNS)
    for row in rows:
        # posting row = fixed + DT cells → pad vault columns
        body = list(row)
        if len(body) < fixed_n + n_dt:
            body.extend([""] * (fixed_n + n_dt - len(body)))
        combined_rows.append([*body[: fixed_n + n_dt], *[""] * n_vault])
    for row in understood:
        # understood row = fixed + vault cells → pad DT columns in the middle
        body = list(row)
        fixed = body[:fixed_n]
        vault_part = body[fixed_n:]
        if len(vault_part) < n_vault:
            vault_part = [*vault_part, *[""] * (n_vault - len(vault_part))]
        combined_rows.append([*fixed, *[""] * n_dt, *vault_part[:n_vault]])

    wb = Workbook()
    ws = wb.active
    ws.title = _SHEET_POSTING
    _write_bundle_sheet(
        ws,
        title=_SHEET_POSTING,
        rows=combined_rows,
        dynamic_headers=combined_headers,
        date_from=date_from,
        date_to=date_to,
        cell_format=cell_format,
    )

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def build_documents_bundle_export(
    db: AsyncSession,
    *,
    tenant_id: uuid.UUID,
    date_from: date | None = None,
    date_to: date | None = None,
    cell_format: BundleCellFormat = "excel",
    display_timezone: str | None = None,
) -> DocumentsBundleExportPayload:
    if date_from is not None and date_to is not None and date_from > date_to:
        raise ValueError("date_from must be on or before date_to")

    display_tz = resolve_display_timezone(display_timezone)

    config = await load_posting_config_for_tenant(db, tenant_id)
    invoices = await _load_bundle_export_invoices(
        db,
        tenant_id=tenant_id,
        date_from=date_from,
        date_to=date_to,
    )

    catalogue = org_document_types_for_bundle_export(config.document_types)
    export_dt_types = _sorted_tenant_document_types(catalogue)
    dt_codes = _dt_codes_ordered(export_dt_types)
    dt_headers = _dt_column_headers(export_dt_types)

    anchors: list[Invoice] = []
    anchor_definitions: dict[int, DocumentTypeDefinition] = {}
    for invoice in invoices:
        definition = resolve_definition_for_invoice(invoice, catalogue)
        if not _is_transactional_posting(definition):
            continue
        assert definition is not None
        anchors.append(invoice)
        anchor_definitions[invoice.id] = definition

    linkage_cache = await build_linkage_sibling_cache(db, tenant_id=tenant_id, anchors=anchors)

    need_upload_actor = [
        inv.id for inv in anchors if not (inv.email_sender or "").strip()
    ]
    upload_actors = await _upload_actors_by_invoice_id(db, need_upload_actor)

    contexts: list[_AnchorExportContext] = []
    for invoice in anchors:
        definition = anchor_definitions[invoice.id]
        linked = await build_dossier_linked_documents(
            db,
            invoice,
            definition=definition,
            document_types=catalogue,
            linkage_cache=linkage_cache,
        )
        invoice_dt_codes = await _invoice_dt_code_by_id(db, linked)
        contexts.append(
            _AnchorExportContext(
                invoice=invoice,
                definition=definition,
                linked=linked,
                invoice_dt_code_by_id=invoice_dt_codes,
            )
        )

    data_rows: list[list[str]] = []
    for ctx in contexts:
        data_rows.append(
            build_documents_bundle_row(
                ctx.invoice,
                definition=ctx.definition,
                linked=ctx.linked,
                dt_codes=dt_codes,
                invoice_dt_code_by_id=ctx.invoice_dt_code_by_id,
                cell_format=cell_format,
                upload_actor=upload_actors.get(ctx.invoice.id),
                display_tz=display_tz,
            )
        )

    # --- Understood soft-bundle rows (same sheet) ---
    understood_anchors = [
        inv for inv in invoices if is_understood_invoice_family_anchor(inv)
    ]
    understood_linkage_cache = await build_linkage_sibling_cache(
        db, tenant_id=tenant_id, anchors=understood_anchors
    )
    understood_upload_actors = await _upload_actors_by_invoice_id(
        db,
        [inv.id for inv in understood_anchors if not (inv.email_sender or "").strip()],
    )

    understood_contexts: list[_UnderstoodExportContext] = []
    vault_folders_seen: set[str] = set()
    for invoice in understood_anchors:
        linked = await build_dossier_linked_documents(
            db,
            invoice,
            definition=None,
            document_types=catalogue,
            linkage_cache=understood_linkage_cache,
        )
        sibling_ids: set[int] = set()
        for doc in linked.documents:
            if doc.is_anchor:
                continue
            if doc.invoice_id is not None:
                sibling_ids.add(doc.invoice_id)
            if doc.manual_link is not None:
                sibling_ids.add(doc.manual_link.invoice_id)
        siblings = await _invoices_by_id(db, sibling_ids)
        vault_folder_by_id: dict[int, str] = {}
        for sib_id, sib in siblings.items():
            label = vault_folder_label_for_export(sib)
            if label:
                vault_folder_by_id[sib_id] = label
                vault_folders_seen.add(label)
        understood_contexts.append(
            _UnderstoodExportContext(
                invoice=invoice,
                linked=linked,
                vault_folder_by_invoice_id=vault_folder_by_id,
            )
        )

    vault_folder_headers = sorted(vault_folders_seen, key=lambda s: s.casefold())
    understood_rows: list[list[str]] = []
    for ctx in understood_contexts:
        understood_rows.append(
            build_understood_bundle_row(
                ctx.invoice,
                linked=ctx.linked,
                vault_folders=vault_folder_headers,
                vault_folder_by_invoice_id=ctx.vault_folder_by_invoice_id,
                cell_format=cell_format,
                upload_actor=understood_upload_actors.get(ctx.invoice.id),
                display_tz=display_tz,
            )
        )

    xlsx_bytes = documents_bundle_rows_to_xlsx(
        data_rows,
        dt_column_headers=dt_headers,
        date_from=date_from,
        date_to=date_to,
        cell_format=cell_format,
        understood_rows=understood_rows,
        vault_folder_headers=vault_folder_headers,
    )
    filename = documents_bundle_filename(date_from=date_from)
    return DocumentsBundleExportPayload(
        xlsx_bytes=xlsx_bytes,
        filename=filename,
        data_rows=len(data_rows) + len(understood_rows),
    )
