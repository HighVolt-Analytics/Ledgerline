"""Export pipeline data to multi-sheet Excel workbook (output_workbook format)."""

from __future__ import annotations

import json
from collections import defaultdict
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import get_settings
from app.services.tenant_storage_paths import tenant_local_dir
from app.models.audit import AuditLog
from app.models.invoice import Invoice, InvoiceStatus
from app.models.journal import JournalEntry
from app.models.line_item import LineItem
from app.models.tenant import Tenant
from app.schemas.rule_book_config import RuleBookConfigPayload
from app.services.reconciliation_service import reconcile_daily
from app.services.rule_book_config_io import tenant_rule_book_config_path
from app.services.rule_book_mapper import (
    FALLBACK_RULE_TYPE,
    load_classification_config,
    map_invoice_with_details,
)
from app.utils.logger import get_logger

logger = get_logger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F6E7A")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="D9E8EC")
TOTAL_FONT = Font(bold=True)

SHEET_README = "README"
SHEET_INVOICES = "Invoices"
SHEET_LINE_ITEMS = "Line Items"
SHEET_LEDGER_MAPPING = "Ledger Mapping"
SHEET_JOURNAL_ENTRIES = "Journal Entries"
SHEET_DAILY_RECON = "Daily Reconciliation"
SHEET_EXPENSE_SUMMARY = "Expense Summary"
SHEET_PROCESSING_STATUS = "Processing Status"
SHEET_RULE_BOOK = "Rule Book"


def invoice_display_id(invoice: Invoice) -> str:
    return f"INV-{invoice.id:03d}"


def _money(value: Decimal | float | int | None) -> float:
    if value is None:
        return 0.0
    return float(value)


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autosize_columns(ws: Worksheet, max_width: int = 42) -> None:
    for col_idx in range(1, ws.max_column + 1):
        letter = get_column_letter(col_idx)
        max_len = 0
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), max_width)


def _extract_po_ref(invoice: Invoice, config: RuleBookConfigPayload) -> str:
    if invoice.po_reference:
        return invoice.po_reference
    invoice_no = invoice.invoice_no
    if not invoice_no:
        return "—"
    inv_text = invoice_no.upper()
    for rule in config.purchase_rules:
        if not rule.enabled:
            continue
        prefix = rule.match_on.po_prefix
        if prefix and prefix.upper() in inv_text:
            return prefix.rstrip("-") + "…"
    return "—"


def _validation_label(invoice: Invoice) -> str:
    if invoice.status == InvoiceStatus.EXCEPTION:
        if invoice.validation_results:
            try:
                results = json.loads(invoice.validation_results)
                failed = [r for r in results if not r.get("passed")]
                if failed:
                    return f"✗ {failed[0].get('rule', 'FAIL')}: {failed[0].get('message', '')}"
            except (json.JSONDecodeError, TypeError):
                pass
        return "✗ Exception"
    if invoice.status == InvoiceStatus.PROCESSED:
        return "✓ Passed"
    if invoice.status == InvoiceStatus.DUPLICATE_SKIPPED:
        return "Duplicate skipped"
    if invoice.status == InvoiceStatus.REJECTED:
        return "Rejected"
    return "—"


def _line_rows_for_export(invoice: Invoice) -> list[dict[str, object]]:
    """Line items from DB, or synthetic split when parser did not persist lines."""
    if invoice.line_items:
        rows: list[dict[str, object]] = []
        for idx, line in enumerate(sorted(invoice.line_items, key=lambda li: li.id), start=1):
            subtotal = line.amount or Decimal("0")
            gst = line.tax_amount
            if gst is None:
                gst = (subtotal * Decimal("0.10")).quantize(Decimal("0.01"))
            rows.append(
                {
                    "line_no": idx,
                    "description": line.description or f"Line {idx}",
                    "qty": line.qty or Decimal("1"),
                    "unit_price": line.unit_price or subtotal,
                    "subtotal": subtotal,
                    "gst": gst,
                    "total": subtotal + gst,
                }
            )
        return rows

    subtotal = invoice.subtotal or Decimal("0")
    if subtotal <= 0:
        return []

    half = (subtotal / 2).quantize(Decimal("0.01"))
    remainder = subtotal - half
    parts = [half, remainder]
    vendor = invoice.vendor or "Invoice"
    rows = []
    for idx, part in enumerate(parts, start=1):
        if part <= 0:
            continue
        gst = (part * Decimal("0.10")).quantize(Decimal("0.01"))
        rows.append(
            {
                "line_no": idx,
                "description": f"{vendor} — line {idx}",
                "qty": Decimal("1"),
                "unit_price": part,
                "subtotal": part,
                "gst": gst,
                "total": part + gst,
            }
        )
    return rows


def _write_readme(
    ws: Worksheet,
    *,
    tenant_name: str,
    tenant_slug: str,
    tenant_id: int,
    date_from: date | None,
    date_to: date | None,
    invoice_count: int,
) -> None:
    rule_path = tenant_rule_book_config_path(tenant_id)
    priority_label = "Purchase rule → Expense rule → Vendor master → Fallback"
    if date_from and date_to:
        range_label = (
            date_from.isoformat()
            if date_from == date_to
            else f"{date_from.isoformat()} to {date_to.isoformat()}"
        )
    elif date_from:
        range_label = f"from {date_from.isoformat()}"
    elif date_to:
        range_label = f"through {date_to.isoformat()}"
    else:
        range_label = "all dates"

    lines = [
        "Invoice Processing Pipeline — Output Workbook",
        "",
        f"Tenant: {tenant_name} ({tenant_slug})",
        f"Invoice date filter: {range_label}",
        f"Invoices exported: {invoice_count}",
        "",
        "Deterministic GL mapping uses this organisation's rule book:",
        f"  {rule_path}",
        f"  Priority order: {priority_label}",
        "",
        "Sheets:",
        "  Invoices — header-level invoice data and validation status",
        "  Line Items — per-line amounts (from DB or estimated split)",
        "  Ledger Mapping — rule book mapping per line (deterministic priority)",
        "  Journal Entries — expense, GST, and AP postings",
        "  Daily Reconciliation — RC1 / RC2 balance checks per day",
        "  Expense Summary — totals by ledger account (ex-GST)",
        "  Processing Status — pipeline stage checkmarks",
        "  Rule Book — active rules for this organisation (same source as mapping)",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 88


def _write_invoices_sheet(ws: Worksheet, invoices: list[Invoice], config: RuleBookConfigPayload) -> None:
    headers = [
        "ID",
        "Vendor",
        "ABN",
        "Invoice No",
        "Invoice Date",
        "Due Date",
        "Currency",
        "PO / Ref",
        "Cost Centre",
        "Subtotal",
        "GST",
        "Total",
        "Validation",
    ]
    _write_header(ws, headers)
    sum_sub = Decimal("0")
    sum_gst = Decimal("0")
    sum_total = Decimal("0")

    for inv in sorted(invoices, key=lambda i: (i.invoice_date or date.min, i.id)):
        sub = inv.subtotal or Decimal("0")
        gst = inv.gst or Decimal("0")
        total = inv.total or sub + gst
        sum_sub += sub
        sum_gst += gst
        sum_total += total
        ws.append(
            [
                invoice_display_id(inv),
                inv.vendor,
                inv.abn,
                inv.invoice_no,
                inv.invoice_date.isoformat() if inv.invoice_date else None,
                inv.due_date.isoformat() if inv.due_date else None,
                inv.currency,
                _extract_po_ref(inv, config),
                "—",
                sub,
                gst,
                total,
                _validation_label(inv),
            ]
        )

    total_row = ws.max_row + 1
    ws.append(
        [
            "GRAND TOTAL",
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            sum_sub,
            sum_gst,
            sum_total,
            None,
        ]
    )
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
    _autosize_columns(ws)


def _write_line_items_sheet(ws: Worksheet, invoices: list[Invoice]) -> None:
    headers = [
        "Invoice ID",
        "Vendor",
        "Line #",
        "Description",
        "Qty",
        "Unit Price",
        "Line Subtotal",
        "Line GST",
        "Line Total",
    ]
    _write_header(ws, headers)
    for inv in sorted(invoices, key=lambda i: (i.invoice_date or date.min, i.id)):
        for line in _line_rows_for_export(inv):
            ws.append(
                [
                    invoice_display_id(inv),
                    inv.vendor,
                    line["line_no"],
                    line["description"],
                    line["qty"],
                    line["unit_price"],
                    line["subtotal"],
                    line["gst"],
                    line["total"],
                ]
            )
    _autosize_columns(ws)


def _write_ledger_mapping_sheet(
    ws: Worksheet, invoices: list[Invoice], config: RuleBookConfigPayload
) -> None:
    headers = [
        "Invoice ID",
        "Invoice Date",
        "Vendor",
        "Line Description",
        "Line Total",
        "Mapped Account",
        "Rule Type",
        "Match Reason",
    ]
    _write_header(ws, headers)
    for inv in sorted(invoices, key=lambda i: (i.invoice_date or date.min, i.id)):
        for line in _line_rows_for_export(inv):
            desc = str(line["description"])
            detail = map_invoice_with_details(inv, config=config, line_description=desc)
            ws.append(
                [
                    invoice_display_id(inv),
                    inv.invoice_date.isoformat() if inv.invoice_date else None,
                    inv.vendor,
                    desc,
                    line["total"],
                    detail.expense_category,
                    detail.rule_type,
                    detail.match_reason,
                ]
            )
            if detail.rule_type == FALLBACK_RULE_TYPE:
                row = ws.max_row
                for col in range(1, len(headers) + 1):
                    ws.cell(row=row, column=col).fill = PatternFill(
                        "solid", fgColor="FCE4E4"
                    )
    _autosize_columns(ws)


def _write_journal_entries_sheet(
    ws: Worksheet, invoices: list[Invoice], config: RuleBookConfigPayload
) -> None:
    headers = ["Date", "Invoice ID", "JE #", "Account", "Description", "Debit", "Credit"]
    _write_header(ws, headers)
    for inv in sorted(invoices, key=lambda i: (i.invoice_date or date.min, i.id)):
        entries = sorted(inv.journal_entries, key=lambda e: e.id)
        if not entries and inv.status == InvoiceStatus.PROCESSED:
            mapping = map_invoice_with_details(inv, config=config)
            sub = inv.subtotal or Decimal("0")
            gst = inv.gst or Decimal("0")
            total = inv.total or sub + gst
            entry_date = inv.invoice_date or date.today()
            desc = f"{inv.vendor} — {inv.invoice_no or ''}"
            entries_data = [
                (mapping.expense_category, sub, Decimal("0")),
                ("GST Paid", gst, Decimal("0")),
                ("Accounts Payable", Decimal("0"), total),
            ]
            for account, dr, cr in entries_data:
                ws.append(
                    [
                        entry_date.isoformat(),
                        invoice_display_id(inv),
                        f"JE-{inv.id:04d}",
                        account,
                        desc.strip(" —"),
                        dr,
                        cr,
                    ]
                )
            continue

        for entry in entries:
            ws.append(
                [
                    entry.date.isoformat() if entry.date else None,
                    invoice_display_id(inv),
                    f"JE-{inv.id:04d}",
                    entry.account_name,
                    f"{inv.vendor} — {inv.invoice_no or ''}".strip(" —"),
                    entry.debit,
                    entry.credit,
                ]
            )
    _autosize_columns(ws)


async def _write_daily_reconciliation_sheet(
    ws: Worksheet,
    session: AsyncSession,
    invoices: list[Invoice],
    *,
    tenant_id: int,
) -> None:
    headers = [
        "Date",
        "Invoices",
        "Σ Invoice Totals",
        "Σ Debit Postings",
        "Σ Credit Postings",
        "Dr - Cr (Δ)",
        "Inv vs Cr (Δ)",
        "Status",
    ]
    _write_header(ws, headers)

    dates = sorted({inv.invoice_date for inv in invoices if inv.invoice_date})
    for day in dates:
        recon = await reconcile_daily(session, day, tenant_id=tenant_id)
        inv_total = sum(
            (inv.total or Decimal("0"))
            for inv in invoices
            if inv.invoice_date == day and inv.status == InvoiceStatus.PROCESSED
        )
        dr_cr = recon.total_debits - recon.total_credits
        inv_vs_cr = inv_total - recon.total_ap_credits
        status = "✓ Balanced" if recon.is_balanced and not recon.halted else "✗ Halted"
        ws.append(
            [
                day.isoformat(),
                recon.total_invoices,
                inv_total,
                recon.total_debits,
                recon.total_credits,
                dr_cr,
                inv_vs_cr,
                status,
            ]
        )
    _autosize_columns(ws)


def _write_expense_summary_sheet(
    ws: Worksheet,
    invoices: list[Invoice],
    config: RuleBookConfigPayload,
) -> None:
    headers = ["Ledger Account", "# Lines", "Total (ex-GST)", "% of Total"]
    _write_header(ws, headers)

    totals: dict[str, Decimal] = defaultdict(lambda: Decimal("0"))
    counts: dict[str, int] = defaultdict(int)

    for inv in invoices:
        for line in _line_rows_for_export(inv):
            desc = str(line["description"])
            detail = map_invoice_with_details(inv, config=config, line_description=desc)
            totals[detail.expense_category] += Decimal(str(line["subtotal"]))
            counts[detail.expense_category] += 1

    grand = sum(totals.values(), Decimal("0"))
    for account in sorted(totals.keys(), key=lambda k: totals[k], reverse=True):
        amount = totals[account]
        pct = (amount / grand * 100) if grand else Decimal("0")
        ws.append([account, counts[account], amount, f"{pct.quantize(Decimal('0.1'))}%"])

    total_row = ws.max_row + 1
    ws.append(["TOTAL", sum(counts.values()), grand, "100.0%"])
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=total_row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = TOTAL_FONT
    _autosize_columns(ws)


async def _audit_flags(
    session: AsyncSession,
    invoice_ids: list[int],
) -> dict[int, set[str]]:
    if not invoice_ids:
        return {}
    rows = (
        await session.execute(
            select(AuditLog.invoice_id, AuditLog.event).where(
                AuditLog.invoice_id.in_(invoice_ids)
            )
        )
    ).all()
    flags: dict[int, set[str]] = defaultdict(set)
    for inv_id, event in rows:
        if inv_id is not None:
            flags[inv_id].add(event)
    return flags


def _stage_check(inv: Invoice, events: set[str], stage: str) -> str:
    if stage == "received":
        return "✓" if "email_ingested" in events or "invoice_uploaded" in events else "—"
    if stage == "parsed":
        return "✓" if "parse_completed" in events or inv.vendor else "—"
    if stage == "validated":
        if inv.status == InvoiceStatus.PROCESSED:
            return "✓"
        if inv.status == InvoiceStatus.EXCEPTION and inv.validation_results:
            return "✗"
        return "—"
    if stage == "mapped":
        return "✓" if inv.account_code or inv.status == InvoiceStatus.PROCESSED else "—"
    if stage == "journaled":
        return "✓" if inv.journal_entries or inv.status == InvoiceStatus.PROCESSED else "—"
    if stage == "notified":
        return "✓" if inv.status == InvoiceStatus.PROCESSED else "—"
    return "—"


def _write_processing_status_sheet(
    ws: Worksheet,
    invoices: list[Invoice],
    audit_flags: dict[int, set[str]],
) -> None:
    headers = [
        "Invoice ID",
        "Vendor",
        "Received",
        "Parsed",
        "Validated",
        "Mapped",
        "Journaled",
        "Notified",
        "Status",
    ]
    _write_header(ws, headers)
    for inv in sorted(invoices, key=lambda i: (i.invoice_date or date.min, i.id)):
        events = audit_flags.get(inv.id, set())
        status_label = inv.status.value.replace("_", " ").title()
        if inv.status == InvoiceStatus.PROCESSED:
            status_label = "Complete"
        ws.append(
            [
                invoice_display_id(inv),
                inv.vendor,
                _stage_check(inv, events, "received"),
                _stage_check(inv, events, "parsed"),
                _stage_check(inv, events, "validated"),
                _stage_check(inv, events, "mapped"),
                _stage_check(inv, events, "journaled"),
                _stage_check(inv, events, "notified"),
                status_label,
            ]
        )
    _autosize_columns(ws)


def _write_rule_book_sheet(ws: Worksheet, config: RuleBookConfigPayload) -> None:
    ws.append(["Classification Rule Book — Purchase → Expense → Vendor → Fallback"])
    ws.append([])

    sections: list[tuple[str, list[tuple[str, str]]]] = [
        (
            "Purchase Rules",
            [
                (rule.name, rule.post_to.ledger)
                for rule in config.purchase_rules
                if rule.enabled
            ],
        ),
        (
            "Expense Rules",
            [
                (rule.name, rule.post_to.ledger)
                for rule in config.expense_rules
                if rule.enabled
            ],
        ),
        (
            "Vendor Masters",
            [
                (master.name, master.default_ledger)
                for master in config.vendor_masters
                if master.default_ledger and master.default_ledger != "—"
            ],
        ),
    ]
    for title, rows in sections:
        ws.append([title])
        ws.append(["Rule", "Maps To"])
        header_row = ws.max_row
        for col in (1, 2):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
        for pattern, category in rows:
            ws.append([pattern, category])
        ws.append([])

    ws.append(["Posting defaults"])
    ws.append(["Setting", "Account"])
    header_row = ws.max_row
    for col in (1, 2):
        cell = ws.cell(row=header_row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    defaults = config.posting_defaults
    ws.append(["Tax account", defaults.tax_account])
    ws.append(["Payable account", defaults.payable_account])
    ws.append(["Fallback (no match)", defaults.fallback_account])
    _autosize_columns(ws)


def workbook_filename(
    tenant_slug: str,
    date_from: date | None = None,
    date_to: date | None = None,
) -> str:
    """Stable on-disk name for a generate/download pair (scoped per organisation)."""
    slug = tenant_slug.strip() or "org"
    if date_from is None and date_to is None:
        return f"output_workbook_{slug}.xlsx"
    if date_from is not None and date_to is not None and date_from == date_to:
        return f"output_workbook_{slug}_{date_from.isoformat()}.xlsx"
    if date_from is not None and date_to is not None:
        return (
            f"output_workbook_{slug}_{date_from.isoformat()}_to_{date_to.isoformat()}.xlsx"
        )
    if date_from is not None:
        return f"output_workbook_{slug}_from_{date_from.isoformat()}.xlsx"
    return f"output_workbook_{slug}_to_{date_to.isoformat()}.xlsx"


async def _load_invoices(
    session: AsyncSession,
    *,
    tenant_id: int,
    date_from: date | None = None,
    date_to: date | None = None,
) -> list[Invoice]:
    stmt = (
        select(Invoice)
        .where(
            Invoice.tenant_id == tenant_id,
            Invoice.status != InvoiceStatus.DUPLICATE_SKIPPED,
            Invoice.status != InvoiceStatus.REJECTED,
        )
        .options(
            selectinload(Invoice.line_items),
            selectinload(Invoice.journal_entries),
        )
        .order_by(Invoice.invoice_date, Invoice.id)
    )
    if date_from is not None:
        stmt = stmt.where(Invoice.invoice_date >= date_from)
    if date_to is not None:
        stmt = stmt.where(Invoice.invoice_date <= date_to)
    return list((await session.execute(stmt)).scalars().all())


def _reports_dir(tenant_id: int) -> Path:
    path = tenant_local_dir(tenant_id, "reports")
    path.mkdir(parents=True, exist_ok=True)
    return path


async def write_workbook(
    session: AsyncSession,
    tenant_id: int,
    workbook_date: date | None = None,
    *,
    date_from: date | None = None,
    date_to: date | None = None,
) -> Path:
    """
    Write multi-sheet Excel workbook matching output_workbook.xlsx layout.

    Scoped to tenant_id — invoices and Rule Book sheet use that organisation's data.
    Filter by invoice_date inclusive range (date_from / date_to).
    workbook_date sets both bounds to the same day (legacy single-day export).
    """
    if workbook_date is not None:
        date_from = date_from or workbook_date
        date_to = date_to or workbook_date
    if date_from is not None and date_to is not None and date_from > date_to:
        raise ValueError("date_from must be on or before date_to")

    org = await session.get(Tenant, tenant_id)
    tenant_name = org.name if org else f"Tenant {tenant_id}"
    tenant_slug = org.slug if org else str(tenant_id)

    invoices = await _load_invoices(
        session, tenant_id=tenant_id, date_from=date_from, date_to=date_to
    )
    config = await load_classification_config(session, tenant_id)
    wb = Workbook()
    wb.remove(wb.active)

    sheets: list[tuple[str, object]] = [
        (
            SHEET_README,
            lambda ws: _write_readme(
                ws,
                tenant_name=tenant_name,
                tenant_slug=tenant_slug,
                tenant_id=tenant_id,
                date_from=date_from,
                date_to=date_to,
                invoice_count=len(invoices),
            ),
        ),
        (SHEET_INVOICES, lambda ws: _write_invoices_sheet(ws, invoices, config)),
        (SHEET_LINE_ITEMS, lambda ws: _write_line_items_sheet(ws, invoices)),
        (SHEET_LEDGER_MAPPING, lambda ws: _write_ledger_mapping_sheet(ws, invoices, config)),
        (SHEET_JOURNAL_ENTRIES, lambda ws: _write_journal_entries_sheet(ws, invoices, config)),
    ]
    for name, writer in sheets:
        ws = wb.create_sheet(name)
        writer(ws)

    ws_recon = wb.create_sheet(SHEET_DAILY_RECON)
    await _write_daily_reconciliation_sheet(ws_recon, session, invoices, tenant_id=tenant_id)

    ws_exp = wb.create_sheet(SHEET_EXPENSE_SUMMARY)
    _write_expense_summary_sheet(ws_exp, invoices, config)

    audit_flags = await _audit_flags(session, [inv.id for inv in invoices])
    ws_status = wb.create_sheet(SHEET_PROCESSING_STATUS)
    _write_processing_status_sheet(ws_status, invoices, audit_flags)

    ws_rules = wb.create_sheet(SHEET_RULE_BOOK)
    _write_rule_book_sheet(ws_rules, config)

    filename = workbook_filename(tenant_slug, date_from, date_to)
    dest = _reports_dir(tenant_id) / filename
    wb.save(dest)
    logger.info("workbook_written", path=str(dest), invoices=len(invoices))

    return dest


async def write_workbook_for_invoice(
    session: AsyncSession,
    invoice: Invoice,
) -> Path | None:
    """Regenerate workbook for the invoice's posting date."""
    if not invoice.invoice_date:
        return None
    return await write_workbook(session, invoice.tenant_id, invoice.invoice_date)
