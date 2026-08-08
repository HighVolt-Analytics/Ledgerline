"""Bulk GL budget import from Excel / CSV templates."""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import Any, Literal

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.department_budget import DepartmentBudget
from app.schemas.department_budget import (
    GlBudgetSubAllocation,
    ParentGlBudgetTreeUpsert,
    PeriodKind,
)
from app.schemas.rule_book_config import ChartOfAccountEntry
from app.services.invoice.invoice_evaluation_service import load_config_for_tenant
from app.services.master_data.chart_of_accounts_service import (
    _find_coa_entry,
    sub_ledgers_for_ledger,
)
from app.services.master_data.department_budget_service import upsert_parent_gl_budget_tree
from app.services.purchase.team_expense_spend_service import (
    current_period_keys,
    period_key_bounds,
)

_MAX_IMPORT_ROWS = 2000
_TEMPLATE_DATA_PAD = 20

_FIELDS = (
    "parent_gl",
    "gl_ledger",
    "period_kind",
    "period_key",
    "allocated",
    "enforcement",
    "notes",
)

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "parent_gl": (
        "parent_gl",
        "parent gl",
        "parent ledger",
        "parent account",
        "wallet",
        "parent",
    ),
    "gl_ledger": (
        "gl_ledger",
        "gl ledger",
        "gl account",
        "ledger",
        "account",
        "sub gl",
        "sub-gl",
        "sub_gl",
    ),
    "period_kind": (
        "period_kind",
        "period kind",
        "period type",
        "frequency",
    ),
    "period_key": (
        "period_key",
        "period key",
        "period",
        "budget period",
    ),
    "allocated": (
        "allocated",
        "budget",
        "amount",
        "budget amount",
        "allocation",
    ),
    "enforcement": (
        "enforcement",
        "over budget",
        "budget mode",
        "mode",
    ),
    "notes": ("notes", "note", "comment", "comments"),
}

_TEMPLATE_COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("parent_gl", "Parent GL *", 28),
    ("gl_ledger", "GL ledger *", 28),
    ("period_kind", "Period kind *", 14),
    ("period_key", "Period key *", 14),
    ("allocated", "Allocated *", 14),
    ("enforcement", "Enforcement", 12),
    ("notes", "Notes", 28),
)

_TEMPLATE_BRAND = "1F6E7A"
_TEMPLATE_HEADER_FILL = PatternFill("solid", fgColor=_TEMPLATE_BRAND)
_TEMPLATE_HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
_TEMPLATE_TITLE_FONT = Font(size=16, bold=True, color=_TEMPLATE_BRAND)
_TEMPLATE_SUBTITLE_FONT = Font(size=10, color="666666")
_TEMPLATE_NOTE_FONT = Font(size=9, italic=True, color="888888")
_TEMPLATE_EXAMPLE_FILL = PatternFill("solid", fgColor="F4F7F8")
_TEMPLATE_PARENT_FILL = PatternFill("solid", fgColor="EEF6F7")
_TEMPLATE_BORDER = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE"),
)

BudgetImportAction = Literal["create", "update", "skip"]


@dataclass
class BudgetImportRowError:
    row_number: int
    parent_gl: str | None
    message: str


@dataclass
class BudgetImportRowPreview:
    row_number: int
    parent_gl: str
    period_kind: str
    period_key: str
    action: BudgetImportAction
    detail: str


@dataclass
class BudgetImportResult:
    dry_run: bool
    created: int = 0
    updated: int = 0
    skipped: int = 0
    errors: list[BudgetImportRowError] = field(default_factory=list)
    previews: list[BudgetImportRowPreview] = field(default_factory=list)


@dataclass
class _ParsedLine:
    row_number: int
    parent_gl: str
    gl_ledger: str
    period_kind: PeriodKind
    period_key: str
    allocated: Decimal | None
    enforcement: str | None
    notes: str | None
    is_parent_row: bool


def _tenant_id(tenant_id: uuid.UUID | int | str) -> uuid.UUID:
    if isinstance(tenant_id, uuid.UUID):
        return tenant_id
    return uuid.UUID(str(tenant_id))


def _normalize_header(value: str) -> str:
    cleaned = re.sub(r"[\s_\-]+", " ", (value or "").strip().lower())
    return cleaned.rstrip(" *").strip()


def _map_headers(headers: list[str], fields: tuple[str, ...]) -> dict[str, str]:
    normalized = {_normalize_header(header): header for header in headers if header.strip()}
    mapping: dict[str, str] = {}
    for field_name in fields:
        aliases = _HEADER_ALIASES.get(field_name, (field_name,))
        for alias in aliases:
            key = _normalize_header(alias)
            if key in normalized:
                mapping[field_name] = normalized[key]
                break
    return mapping


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_amount(raw: str) -> tuple[Decimal | None, str | None]:
    cleaned = (raw or "").strip().replace(",", "").replace("$", "")
    if not cleaned:
        return None, None
    try:
        value = Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None, f"invalid amount {raw!r}"
    if value < 0:
        return None, "allocated must be >= 0"
    return value.quantize(Decimal("0.01")), None


def _normalize_period_kind(raw: str) -> PeriodKind | None:
    token = (raw or "").strip().lower()
    if token in {"monthly", "month", "m"}:
        return "monthly"
    if token in {"quarterly", "quarter", "q"}:
        return "quarterly"
    if token in {"annual", "yearly", "year", "a", "y"}:
        return "annual"
    return None


def _normalize_enforcement(raw: str) -> str | None:
    token = (raw or "").strip().lower()
    if not token:
        return None
    if token in {"soft", "s", "manager", "approve"}:
        return "soft"
    if token in {"hard", "h", "block", "strict"}:
        return "hard"
    return None


def _detect_budget_headers(headers: list[str]) -> bool:
    mapping = _map_headers(headers, _FIELDS)
    return "parent_gl" in mapping and "gl_ledger" in mapping and "period_kind" in mapping


def _find_header_row_in_sheet(sheet) -> tuple[int, list[str]] | None:
    for row_idx, cells in enumerate(sheet.iter_rows(min_row=1, max_row=25, values_only=True), start=1):
        headers = [_cell_str(cell) for cell in cells]
        if not any(headers):
            continue
        if _detect_budget_headers(headers):
            return row_idx, headers
    return None


def _parse_rows_from_csv(data: bytes) -> list[dict[str, str]]:
    text = data.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValueError("CSV file has no header row")
    rows: list[dict[str, str]] = []
    for raw in reader:
        if not any(_cell_str(v) for v in raw.values()):
            continue
        rows.append({k: _cell_str(v) for k, v in raw.items() if k})
    return rows


def _parse_rows_from_xlsx(data: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    header_info: tuple[int, list[str], Any] | None = None
    for sheet in workbook.worksheets:
        found = _find_header_row_in_sheet(sheet)
        if found:
            header_info = (found[0], found[1], sheet)
            break
    if not header_info:
        raise ValueError("Excel file has no recognisable GL budget header row")

    header_row_idx, headers, sheet = header_info
    rows: list[dict[str, str]] = []
    for cells in sheet.iter_rows(min_row=header_row_idx + 1, values_only=True):
        values = [_cell_str(cell) for cell in cells]
        if not any(values):
            continue
        row = {
            headers[i]: values[i] if i < len(values) else ""
            for i in range(len(headers))
            if headers[i]
        }
        rows.append(row)
    return rows


def parse_budget_import_file(data: bytes, filename: str) -> list[dict[str, str]]:
    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        rows = _parse_rows_from_csv(data)
    elif lower.endswith(".xlsx"):
        rows = _parse_rows_from_xlsx(data)
    else:
        raise ValueError("Accepted formats: .csv, .xlsx")
    if len(rows) > _MAX_IMPORT_ROWS:
        raise ValueError(f"Maximum {_MAX_IMPORT_ROWS} data rows per import")
    return rows


def _team_expense_coa_entries(coa: list[ChartOfAccountEntry]) -> list[ChartOfAccountEntry]:
    """Parent GLs suitable for Team Expense budget wallets.

    Matches route-target guidance: Expense and Asset (advances / floats may be Asset).
    """
    allowed = {"expense", "asset"}
    out: list[ChartOfAccountEntry] = []
    for entry in coa:
        if (entry.type or "").strip().lower() not in allowed:
            continue
        if not (entry.name or "").strip():
            continue
        out.append(entry)
    return out


def _apply_cell_border(ws, row: int, col: int) -> None:
    ws.cell(row=row, column=col).border = _TEMPLATE_BORDER


def _autosize_template_columns(ws) -> None:
    for idx, (_field, _label, width) in enumerate(_TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_instructions_sheet(ws) -> None:
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 78

    ws["A1"] = "GL budget import"
    ws["A1"].font = _TEMPLATE_TITLE_FONT
    ws.merge_cells("A1:B1")

    ws["A2"] = "Team Expenses · Parent GL wallets and Sub-GL slices"
    ws["A2"].font = _TEMPLATE_SUBTITLE_FONT
    ws.merge_cells("A2:B2")

    rows = [
        ("Purpose", "Create or update Parent GL budgets and their Sub-GL allocations in bulk."),
        (
            "Rows",
            "One row for the Parent GL (GL ledger = Parent GL), then one row per Sub-GL under that parent.",
        ),
        (
            "Sum rule",
            "When a parent has Sub-GLs in the chart of accounts, Sub-GL amounts must sum to the parent allocated amount.",
        ),
        ("Period kind", "monthly, quarterly, or annual."),
        ("Period key", "monthly=YYYY-MM · quarterly=YYYY-Qn · annual=YYYY (example: 2026-08, 2026-Q3, 2026)."),
        ("Enforcement", "soft (manager can approve overruns) or hard (block until budget raised). Leave blank for soft."),
        ("Notes", "Optional; taken from the parent row."),
        ("Matching", "GL names must match your chart of accounts (case-insensitive)."),
        ("Blank rows", "Leave Allocated blank on unused template rows — those groups are skipped."),
        ("Limits", f"Maximum {_MAX_IMPORT_ROWS} data rows per upload."),
        ("Workflow", "1. Fill Budget data sheet  2. Review import  3. Confirm upload"),
    ]
    start = 4
    for offset, (label, value) in enumerate(rows):
        row = start + offset
        ws.cell(row=row, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=row, column=2, value=value).font = Font(size=10)
        ws.cell(row=row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    example_row = start + len(rows) + 2
    ws.cell(row=example_row, column=1, value="Example rows").font = Font(bold=True, size=10)
    ws.merge_cells(start_row=example_row, start_column=1, end_row=example_row, end_column=2)

    header_row = example_row + 1
    example_headers = [label for _f, label, _w in _TEMPLATE_COLUMNS]
    for col, header in enumerate(example_headers, start=1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = _TEMPLATE_HEADER_FILL
        cell.font = _TEMPLATE_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    examples = [
        ["Marketing Expenses", "Marketing Expenses", "monthly", "2026-08", "10000", "soft", "Q1 campaign"],
        ["Marketing Expenses", "Hotel", "monthly", "2026-08", "4000", "", ""],
        ["Marketing Expenses", "Traveling", "monthly", "2026-08", "6000", "", ""],
    ]
    for offset, values in enumerate(examples):
        value_row = header_row + 1 + offset
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=value_row, column=col, value=value)
            cell.fill = _TEMPLATE_EXAMPLE_FILL
            cell.font = Font(size=10, italic=True)

    note_row = header_row + len(examples) + 2
    ws.cell(
        row=note_row,
        column=1,
        value="Enter your data on the Budget data sheet. Do not rename header cells.",
    ).font = _TEMPLATE_NOTE_FONT
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=2)


def _write_data_sheet(
    ws,
    *,
    period_kind: PeriodKind,
    period_key: str,
    coa_rows: list[tuple[str, str, Decimal | None]],
) -> None:
    col_count = len(_TEMPLATE_COLUMNS)
    last_col = get_column_letter(col_count)

    ws["A1"] = "GL budgets"
    ws["A1"].font = _TEMPLATE_TITLE_FONT
    ws.merge_cells(f"A1:{last_col}1")

    ws["A2"] = (
        f"Period: {period_kind} · {period_key}. "
        "Parent row: GL ledger equals Parent GL. Sub rows list each Sub-GL. "
        "Fill Allocated; delete unused sample groups before upload if needed."
    )
    ws["A2"].font = _TEMPLATE_SUBTITLE_FONT
    ws.merge_cells(f"A2:{last_col}2")

    header_row = 4
    for col, (_field, label, _width) in enumerate(_TEMPLATE_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.fill = _TEMPLATE_HEADER_FILL
        cell.font = _TEMPLATE_HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _apply_cell_border(ws, header_row, col)

    kind_col = next(i for i, (f, _, _) in enumerate(_TEMPLATE_COLUMNS, start=1) if f == "period_kind")
    enf_col = next(i for i, (f, _, _) in enumerate(_TEMPLATE_COLUMNS, start=1) if f == "enforcement")
    first_data_row = header_row + 1

    kind_letter = get_column_letter(kind_col)
    enf_letter = get_column_letter(enf_col)
    last_data_row = first_data_row + max(len(coa_rows), 1) + _TEMPLATE_DATA_PAD - 1

    kind_dv = DataValidation(
        type="list",
        formula1='"monthly,quarterly,annual"',
        allow_blank=False,
        showDropDown=False,
    )
    kind_dv.error = "Choose monthly, quarterly, or annual"
    kind_dv.errorTitle = "Invalid period kind"
    kind_dv.add(f"{kind_letter}{first_data_row}:{kind_letter}{last_data_row}")
    ws.add_data_validation(kind_dv)

    enf_dv = DataValidation(
        type="list",
        formula1='"soft,hard"',
        allow_blank=True,
        showDropDown=False,
    )
    enf_dv.error = "Choose soft or hard"
    enf_dv.errorTitle = "Invalid enforcement"
    enf_dv.add(f"{enf_letter}{first_data_row}:{enf_letter}{last_data_row}")
    ws.add_data_validation(enf_dv)

    write_rows = list(coa_rows)
    if not write_rows:
        write_rows = [
            ("Marketing Expenses", "Marketing Expenses", None),
            ("Marketing Expenses", "Hotel", None),
            ("Marketing Expenses", "Traveling", None),
        ]

    for offset, (parent_gl, gl_ledger, allocated) in enumerate(write_rows):
        row = first_data_row + offset
        values = [
            parent_gl,
            gl_ledger,
            period_kind,
            period_key,
            "" if allocated is None else float(allocated),
            "soft" if parent_gl == gl_ledger else "",
            "",
        ]
        is_parent = parent_gl.strip().lower() == gl_ledger.strip().lower()
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            _apply_cell_border(ws, row, col)
            cell.alignment = Alignment(vertical="center")
            if is_parent:
                cell.fill = _TEMPLATE_PARENT_FILL

    pad_start = first_data_row + len(write_rows)
    for row in range(pad_start, last_data_row + 1):
        for col in range(1, col_count + 1):
            _apply_cell_border(ws, row, col)
            ws.cell(row=row, column=col).alignment = Alignment(vertical="center")

    ws.freeze_panes = ws.cell(row=first_data_row, column=1)
    ws.row_dimensions[header_row].height = 28
    _autosize_template_columns(ws)


async def build_budget_import_template(
    session: AsyncSession,
    tenant_id: uuid.UUID | int | str,
    *,
    period_kind: PeriodKind = "monthly",
    period_key: str | None = None,
    prefill_coa: bool = True,
) -> bytes:
    tid = _tenant_id(tenant_id)
    kind = period_kind if period_kind in {"monthly", "quarterly", "annual"} else "monthly"
    key = (period_key or "").strip() or current_period_keys(date.today())[kind]
    if period_key_bounds(kind, key) is None:
        raise ValueError(f"Invalid period key {key!r} for {kind}")

    coa_rows: list[tuple[str, str, Decimal | None]] = []
    if prefill_coa:
        config = await load_config_for_tenant(session, tid)
        for entry in _team_expense_coa_entries(list(config.chart_of_accounts or [])):
            parent = entry.name.strip()
            coa_rows.append((parent, parent, None))
            for sub in entry.sub_ledgers or []:
                sub_name = (sub.name or "").strip()
                if sub_name:
                    coa_rows.append((parent, sub_name, None))

    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Budget data"
    _write_data_sheet(data_sheet, period_kind=kind, period_key=key, coa_rows=coa_rows)

    instructions = workbook.create_sheet("Instructions")
    _write_instructions_sheet(instructions)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _extract_line(
    raw: dict[str, str],
    header_map: dict[str, str],
    row_number: int,
) -> tuple[_ParsedLine | None, str | None]:
    def get(field: str) -> str:
        header = header_map.get(field)
        if not header:
            return ""
        return _cell_str(raw.get(header, ""))

    parent_gl = get("parent_gl")
    gl_ledger = get("gl_ledger")
    period_kind_raw = get("period_kind")
    period_key = get("period_key")
    allocated_raw = get("allocated")
    enforcement_raw = get("enforcement")
    notes_raw = get("notes")

    # Completely blank leftover template row
    if not any([parent_gl, gl_ledger, period_kind_raw, period_key, allocated_raw, enforcement_raw, notes_raw]):
        return None, None

    if not parent_gl:
        return None, "parent_gl is required"
    if not gl_ledger:
        return None, "gl_ledger is required"

    period_kind = _normalize_period_kind(period_kind_raw)
    if period_kind is None:
        return None, "period_kind must be monthly, quarterly, or annual"
    if not period_key:
        return None, "period_key is required"
    if period_key_bounds(period_kind, period_key) is None:
        return None, f"invalid period_key {period_key!r} for {period_kind}"

    allocated, amount_error = _parse_amount(allocated_raw)
    if amount_error:
        return None, amount_error

    enforcement = None
    if enforcement_raw:
        enforcement = _normalize_enforcement(enforcement_raw)
        if enforcement is None:
            return None, "enforcement must be soft or hard"

    is_parent = parent_gl.strip().lower() == gl_ledger.strip().lower()
    return (
        _ParsedLine(
            row_number=row_number,
            parent_gl=parent_gl,
            gl_ledger=gl_ledger,
            period_kind=period_kind,
            period_key=period_key,
            allocated=allocated,
            enforcement=enforcement,
            notes=notes_raw or None,
            is_parent_row=is_parent,
        ),
        None,
    )


async def _parent_budget_exists(
    session: AsyncSession,
    tenant_id: uuid.UUID,
    *,
    parent_gl: str,
    period_kind: str,
    period_key: str,
) -> bool:
    row = (
        await session.execute(
            select(DepartmentBudget.id).where(
                DepartmentBudget.tenant_id == tenant_id,
                DepartmentBudget.gl_ledger == parent_gl,
                DepartmentBudget.period_kind == period_kind,
                DepartmentBudget.period_key == period_key,
            )
        )
    ).scalar_one_or_none()
    return row is not None


async def import_department_budgets(
    session: AsyncSession,
    tenant_id: uuid.UUID | int | str,
    *,
    rows: list[dict[str, str]],
    dry_run: bool = False,
) -> BudgetImportResult:
    tid = _tenant_id(tenant_id)
    if not rows:
        raise ValueError("No data rows found")

    first_keys = list(rows[0].keys())
    header_map = _map_headers(first_keys, _FIELDS)
    required = ("parent_gl", "gl_ledger", "period_kind", "period_key")
    missing = [name for name in required if name not in header_map]
    if missing:
        raise ValueError("Import file must include columns: " + ", ".join(missing))
    if "allocated" not in header_map:
        raise ValueError("Import file must include an allocated / budget column")

    result = BudgetImportResult(dry_run=dry_run)
    lines: list[_ParsedLine] = []

    # Excel template header is row 4 → data starts at row 5; CSV header is row 1 → data at 2.
    # Use enumerate offset 2 as a stable human row when CSV; for xlsx the absolute sheet row
    # is unknown after parse, so report 2-based data index (same as employee import).
    for index, raw in enumerate(rows, start=2):
        parsed, error = _extract_line(raw, header_map, index)
        if error:
            result.errors.append(
                BudgetImportRowError(
                    row_number=index,
                    parent_gl=_cell_str(raw.get(header_map.get("parent_gl", ""), "")) or None,
                    message=error,
                )
            )
            result.skipped += 1
            continue
        if parsed is None:
            continue
        lines.append(parsed)

    if not lines and not result.errors:
        raise ValueError("No data rows found")

    groups: dict[tuple[str, str, str], list[_ParsedLine]] = {}
    for line in lines:
        key = (line.parent_gl.strip().lower(), line.period_kind, line.period_key)
        groups.setdefault(key, []).append(line)

    config = await load_config_for_tenant(session, tid)
    coa = list(config.chart_of_accounts or [])

    for (_parent_key, period_kind, period_key), group_lines in groups.items():
        first_row = min(line.row_number for line in group_lines)
        display_parent = group_lines[0].parent_gl.strip()

        parent_lines = [line for line in group_lines if line.is_parent_row]
        sub_lines = [line for line in group_lines if not line.is_parent_row]

        # Skip untouched template groups (no amounts filled).
        amounts = [line.allocated for line in group_lines]
        if all(amount is None for amount in amounts):
            result.skipped += 1
            result.previews.append(
                BudgetImportRowPreview(
                    row_number=first_row,
                    parent_gl=display_parent,
                    period_kind=period_kind,
                    period_key=period_key,
                    action="skip",
                    detail="No allocated amounts — skipped",
                )
            )
            continue

        if len(parent_lines) > 1:
            result.errors.append(
                BudgetImportRowError(
                    row_number=first_row,
                    parent_gl=display_parent,
                    message="multiple parent rows for the same Parent GL and period",
                )
            )
            result.skipped += 1
            continue

        parent_line = parent_lines[0] if parent_lines else None
        if parent_line is None:
            result.errors.append(
                BudgetImportRowError(
                    row_number=first_row,
                    parent_gl=display_parent,
                    message="missing parent row (GL ledger must equal Parent GL)",
                )
            )
            result.skipped += 1
            continue

        if parent_line.allocated is None:
            result.errors.append(
                BudgetImportRowError(
                    row_number=parent_line.row_number,
                    parent_gl=display_parent,
                    message="parent allocated amount is required",
                )
            )
            result.skipped += 1
            continue

        entry = _find_coa_entry(parent_line.parent_gl, coa)
        if entry is None:
            result.errors.append(
                BudgetImportRowError(
                    row_number=parent_line.row_number,
                    parent_gl=display_parent,
                    message=f"Parent GL {parent_line.parent_gl!r} is not in your chart of accounts",
                )
            )
            result.skipped += 1
            continue

        parent_name = (entry.name or "").strip() or parent_line.parent_gl.strip()
        catalog_subs = [
            sub.name.strip()
            for sub in sub_ledgers_for_ledger(parent_name, coa)
            if (sub.name or "").strip()
        ]
        catalog_lookup = {name.lower(): name for name in catalog_subs}

        sub_allocations: dict[str, Decimal] = {}
        group_error = False
        for sub_line in sub_lines:
            sub_key = sub_line.gl_ledger.strip().lower()
            canonical = catalog_lookup.get(sub_key)
            if canonical is None:
                result.errors.append(
                    BudgetImportRowError(
                        row_number=sub_line.row_number,
                        parent_gl=parent_name,
                        message=f"Unknown Sub-GL for this parent: {sub_line.gl_ledger}",
                    )
                )
                group_error = True
                continue
            if canonical in sub_allocations:
                result.errors.append(
                    BudgetImportRowError(
                        row_number=sub_line.row_number,
                        parent_gl=parent_name,
                        message=f"duplicate Sub-GL row for {canonical}",
                    )
                )
                group_error = True
                continue
            sub_allocations[canonical] = (
                sub_line.allocated if sub_line.allocated is not None else Decimal("0")
            )

        if group_error:
            result.skipped += 1
            continue

        # Fill missing catalog Sub-GLs as 0 so upsert can validate the full tree.
        for name in catalog_subs:
            sub_allocations.setdefault(name, Decimal("0"))

        if catalog_subs and sub_allocations and not sub_lines:
            # Parent-only file when COA has children — still allow if we zero-fill,
            # but sum must equal parent; warn via upsert error if non-zero parent.
            pass

        body = ParentGlBudgetTreeUpsert(
            parent_gl=parent_name,
            period_kind=period_kind,  # type: ignore[arg-type]
            period_key=period_key,
            allocated=parent_line.allocated,
            sub_allocations=[
                GlBudgetSubAllocation(gl_ledger=name, allocated=sub_allocations[name])
                for name in catalog_subs
            ],
            enforcement=(parent_line.enforcement or "soft"),  # type: ignore[arg-type]
            notes=parent_line.notes,
        )

        exists = await _parent_budget_exists(
            session,
            tid,
            parent_gl=parent_name,
            period_kind=period_kind,
            period_key=period_key,
        )
        action: BudgetImportAction = "update" if exists else "create"
        sub_count = len(catalog_subs)
        detail = (
            f"{'Update' if exists else 'Create'} parent wallet"
            + (f" + {sub_count} Sub-GL(s)" if sub_count else "")
            + f" · {body.allocated}"
        )

        try:
            if not dry_run:
                await upsert_parent_gl_budget_tree(session, tid, body)
            else:
                # Validate without writing: reuse upsert rules via a throwaway check.
                # Calling upsert then rolling back is risky mid-batch; mirror key checks.
                parent_allocated = body.allocated
                if catalog_subs:
                    sub_total = sum(sub_allocations.values(), Decimal("0"))
                    if sub_total != parent_allocated:
                        raise ValueError(
                            f"Sum of Sub-GL budgets ({sub_total}) must equal parent budget ({parent_allocated})"
                        )
                elif sub_lines:
                    raise ValueError(
                        f"Parent GL {parent_name!r} has no Sub-GLs in the chart of accounts"
                    )
        except ValueError as exc:
            result.errors.append(
                BudgetImportRowError(
                    row_number=parent_line.row_number,
                    parent_gl=parent_name,
                    message=str(exc),
                )
            )
            result.skipped += 1
            continue

        if action == "create":
            result.created += 1
        else:
            result.updated += 1
        result.previews.append(
            BudgetImportRowPreview(
                row_number=first_row,
                parent_gl=parent_name,
                period_kind=period_kind,
                period_key=period_key,
                action=action,
                detail=detail,
            )
        )

    return result
