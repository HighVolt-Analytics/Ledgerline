"""Documents bundle Excel export API tests."""

from __future__ import annotations

import io
from datetime import date
from decimal import Decimal

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.invoice import Invoice, InvoiceStatus, PurchaseDocumentType
from app.models.line_item import LineItem
from app.schemas.dossier import DossierLinkedDocumentsResponse, DossierMatchSummaryResponse
from app.services.audit.audit_export_service import (
    collect_linked_doc_entries,
    format_linked_docs_export,
    linked_docs_by_dt_code,
    vault_view_path,
)
from app.services.invoice.invoice_evaluation_service import ROUTE_PURCHASE
from app.services.purchase.purchase_match_service import sync_purchase_order_from_invoice
from app.services.reports.documents_bundle_export_service import (
    _HEADER_ROW,
    _match_flags,
    build_documents_bundle_export,
    build_documents_bundle_row,
    bundle_dt_cells_by_code,
)
from app.tenant_ids import TESTING_TENANT_UUID
from tests.test_audit_export_api import assert_csv_hyperlink


def _read_xlsx(content: bytes) -> tuple[list[str], list[list[object]], object]:
    wb = load_workbook(io.BytesIO(content), data_only=False)
    ws = wb.active
    header = [
        ws.cell(row=_HEADER_ROW, column=col).value
        for col in range(1, ws.max_column + 1)
    ]
    assert header and header[0] is not None
    headers = [str(value) for value in header]
    data: list[list[object]] = []
    for row_idx in range(_HEADER_ROW + 1, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
        if all(value is None or value == "" for value in values):
            continue
        data.append(values)
    return headers, data, ws


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _invoice_no_from_cell(value: object) -> str:
    text = _cell_text(value)
    if text.startswith("=HYPERLINK("):
        label_start = text.rfind(',"')
        if label_start >= 0:
            return text[label_start + 2 : -2]
    if " | " in text:
        return text.split(" | ", 1)[0]
    return text


def _invoice_nos_from_rows(header: list[str], data: list[list[object]]) -> set[str]:
    col = header.index("Invoice no.")
    return {_invoice_no_from_cell(row[col]) for row in data}


def _assert_hyperlink_cell(ws, *, header: list[str], data_row_index: int, column_name: str, url: str) -> None:
    """data_row_index is 0-based among data rows under the header."""
    col = header.index(column_name) + 1
    excel_row = _HEADER_ROW + 1 + data_row_index
    cell = ws.cell(row=excel_row, column=col)
    assert cell.hyperlink is not None, f"expected hyperlink in {column_name}"
    target = getattr(cell.hyperlink, "target", None) or str(cell.hyperlink)
    assert url in target
    assert _cell_text(cell.value)


def test_match_flags_yes_no_only() -> None:
    linked = DossierLinkedDocumentsResponse(
        linkage_kind="po_reference",
        linkage_label="PO reference",
        enforce_bundle=True,
        match_summary=DossierMatchSummaryResponse(
            status="3-Way Match",
            currency="AUD",
            po_value=100.0,
            invoice_total=110.0,
        ),
    )
    two, three, universal = _match_flags(linked)
    assert two == "No"
    assert three == "Yes"
    assert universal == "No"


@pytest.mark.asyncio
async def test_documents_bundle_export_excludes_non_posting_documents(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    db_session.add_all(
        [
            Invoice(
                tenant_id=TESTING_TENANT_UUID,
                vendor="Supplier",
                document_type_code="DT-01",
                invoice_no="INV-POST-1",
                invoice_date=date(2026, 5, 10),
                status=InvoiceStatus.PROCESSED,
                file_hash="bundle-posting-yes",
            ),
            Invoice(
                tenant_id=TESTING_TENANT_UUID,
                vendor="Supplier",
                document_type_code="DT-02",
                invoice_no="PO-SUPPORT-1",
                invoice_date=date(2026, 5, 10),
                status=InvoiceStatus.PROCESSED,
                file_hash="bundle-posting-no",
            ),
        ]
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-05-01&date_to=2026-05-31"
    )
    assert res.status_code == 200
    assert res.headers.get("x-data-rows") == "1"
    header, data, _ws = _read_xlsx(res.content)
    assert "Class" not in header
    assert "Posting" not in header
    assert "Status" not in header
    assert "2 way match" not in header
    assert "3 way match" not in header
    assert "Universal match" in header
    assert "Invoice date" in header
    assert "Counterparty" in header
    assert "Linkage" in header
    assert "PO goods invoice" in header
    assert "GRN (supporting)" in header

    invoice_nos = _invoice_nos_from_rows(header, data)
    assert "INV-POST-1" in invoice_nos
    assert "PO-SUPPORT-1" not in invoice_nos

    posting_row = next(
        row for row in data if _invoice_no_from_cell(row[header.index("Invoice no.")]) == "INV-POST-1"
    )
    assert posting_row[header.index("Universal match")] == "No"


def test_documents_bundle_row_universal_match_flag() -> None:
    """Universal match column is Yes/No; 2/3-way are not exported as columns."""
    from datetime import datetime, timezone

    from app.schemas.document_type import DocumentTypeDefinition
    from app.schemas.dossier import DossierLinkedDocumentResponse

    invoice = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        document_type_code="DT-01",
        invoice_no="INV-FLAG-1",
        status=InvoiceStatus.PROCESSED,
        email_sender="vendor@example.com",
        capture_source="email",
    )
    invoice.id = 99
    invoice.created_at = datetime(2026, 5, 12, 9, 30, 0, tzinfo=timezone.utc)
    definition = DocumentTypeDefinition(
        code="DT-01",
        title="PO goods invoice",
        short_title="PO goods invoice",
        klass="Transactional",
        posting="Yes",
        recognition_mode="signals",
        recognition_signals=[],
        llm_prompt="",
        route_target="Purchase Management",
    )
    linked = DossierLinkedDocumentsResponse(
        linkage_kind="po_reference",
        linkage_label="PO reference",
        enforce_bundle=True,
        match_summary=DossierMatchSummaryResponse(
            status="3-Way Match",
            currency="AUD",
            po_value=200.0,
            invoice_total=220.0,
        ),
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-01",
                label="Anchor",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=99,
            ),
        ],
    )
    row = build_documents_bundle_row(
        invoice,
        definition=definition,
        linked=linked,
        dt_codes=[],
    )
    # Fixed column order: Timestamp(0), Uploaded by(1), Source(2), … Universal(12)
    assert row[0] == "2026-05-12 09:30:00"
    assert row[1] == "vendor@example.com"
    assert row[2] == "Email"
    assert row[12] == "No"


def test_documents_bundle_row_manual_uploader_name() -> None:
    """Direct uploads prefer persisted uploaded_by_* over empty email_sender."""
    from datetime import datetime, timezone

    from app.schemas.document_type import DocumentTypeDefinition

    invoice = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        document_type_code="DT-01",
        invoice_no="INV-UPLOAD-1",
        status=InvoiceStatus.PROCESSED,
        capture_source="upload",
        uploaded_by_name="Vishnu Admin",
        uploaded_by_email="vishnu@example.com",
    )
    invoice.id = 101
    invoice.created_at = datetime(2026, 5, 12, 10, 0, 0, tzinfo=timezone.utc)
    definition = DocumentTypeDefinition(
        code="DT-01",
        title="PO goods invoice",
        short_title="PO goods invoice",
        klass="Transactional",
        posting="Yes",
        recognition_mode="signals",
        recognition_signals=[],
        llm_prompt="",
        route_target="Purchase Management",
    )
    linked = DossierLinkedDocumentsResponse(
        linkage_kind="po_reference",
        linkage_label="PO reference",
        enforce_bundle=False,
        match_summary=None,
        documents=[],
    )
    row = build_documents_bundle_row(
        invoice,
        definition=definition,
        linked=linked,
        dt_codes=[],
    )
    assert row[1] == "Vishnu Admin"
    assert row[2] == "Direct upload"


def test_documents_bundle_row_upload_actor_and_source() -> None:
    from app.schemas.document_type import DocumentTypeDefinition
    from app.schemas.dossier import DossierLinkedDocumentResponse

    invoice = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        document_type_code="DT-01",
        invoice_no="INV-UPLOAD-1",
        status=InvoiceStatus.PROCESSED,
        capture_source="upload",
    )
    invoice.id = 100
    definition = DocumentTypeDefinition(
        code="DT-01",
        title="PO goods invoice",
        short_title="PO goods invoice",
        klass="Transactional",
        posting="Yes",
        recognition_mode="signals",
        recognition_signals=[],
        llm_prompt="",
        route_target="Purchase Management",
    )
    linked = DossierLinkedDocumentsResponse(
        linkage_kind="po_reference",
        linkage_label="PO reference",
        enforce_bundle=True,
        match_summary=DossierMatchSummaryResponse(
            status="2-Way Match",
            currency="AUD",
            po_value=100.0,
            invoice_total=100.0,
        ),
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-01",
                label="Anchor",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=100,
            ),
        ],
    )
    row = build_documents_bundle_row(
        invoice,
        definition=definition,
        linked=linked,
        dt_codes=[],
        upload_actor="Ada Lovelace",
    )
    assert row[1] == "Ada Lovelace"
    assert row[2] == "Direct upload"


@pytest.mark.asyncio
async def test_documents_bundle_export_linked_po_hyperlink(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    po_doc = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme Supplies",
        document_type_code="DT-02",
        po_reference="PO-BUNDLE-1",
        invoice_no="PO-BUNDLE-1",
        invoice_date=date(2026, 5, 12),
        route_target=ROUTE_PURCHASE,
        purchase_document_type=PurchaseDocumentType.PO.value,
        subtotal=Decimal("200.00"),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-po-doc",
    )
    db_session.add(po_doc)
    await db_session.flush()
    db_session.add(
        LineItem(
            invoice_id=po_doc.id,
            description="Paper",
            qty=Decimal("4"),
            unit_price=Decimal("50.00"),
            amount=Decimal("200.00"),
        )
    )
    await db_session.flush()
    await sync_purchase_order_from_invoice(db_session, po_doc)

    inv = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme Supplies",
        document_type_code="DT-01",
        po_reference="PO-BUNDLE-1",
        invoice_no="INV-BUNDLE-1",
        invoice_date=date(2026, 5, 15),
        route_target=ROUTE_PURCHASE,
        purchase_document_type=PurchaseDocumentType.INVOICE.value,
        subtotal=Decimal("200.00"),
        gst=Decimal("20.00"),
        total=Decimal("220.00"),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-inv-doc",
    )
    db_session.add(inv)
    await db_session.flush()
    db_session.add(
        LineItem(
            invoice_id=inv.id,
            description="Paper",
            qty=Decimal("4"),
            unit_price=Decimal("50.00"),
            amount=Decimal("200.00"),
        )
    )
    await db_session.flush()
    po = await sync_purchase_order_from_invoice(db_session, inv)
    assert po is not None
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-05-01&date_to=2026-05-31"
    )
    assert res.status_code == 200
    assert "documents_bundle_2026-05.xlsx" in res.headers.get("content-disposition", "")
    assert res.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    header, data, ws = _read_xlsx(res.content)
    row_index = next(
        i
        for i, r in enumerate(data)
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-BUNDLE-1"
    )
    row = data[row_index]
    assert row[header.index("Universal match")] == "No"

    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="PO (supporting)",
        url=vault_view_path(po_doc.id),
    )


@pytest.mark.asyncio
async def test_documents_bundle_export_universal_match_yes(
    db_session: AsyncSession,
) -> None:
    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Importer",
        status=InvoiceStatus.PROCESSED,
        document_type_code="DT-01",
        invoice_no="SHARED-INV-77",
        invoice_date=date(2026, 6, 3),
        file_hash="bundle-universal-anchor",
    )
    sibling = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Carrier",
        status=InvoiceStatus.PROCESSED,
        document_type_code="DT-06",
        invoice_no="SHARED-INV-77",
        invoice_date=date(2026, 6, 3),
        file_hash="bundle-universal-sibling",
    )
    db_session.add_all([anchor, sibling])
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 6, 1),
        date_to=date(2026, 6, 30),
    )
    header, data, _ws = _read_xlsx(payload.xlsx_bytes)
    row = next(
        r for r in data if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "SHARED-INV-77"
    )
    assert row[header.index("Universal match")] == "Yes"


@pytest.mark.asyncio
async def test_documents_bundle_export_with_empty_tenant_document_types(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    from app.models.tenant_rule_book_config import TenantRuleBookConfig
    from app.services.rule_book.rule_book_config_io import clear_posting_config_cache

    row = await db_session.get(TenantRuleBookConfig, TESTING_TENANT_UUID)
    assert row is not None
    row.config = {**row.config, "document_types": []}
    await db_session.flush()
    clear_posting_config_cache()

    db_session.add(
        Invoice(
            tenant_id=TESTING_TENANT_UUID,
            vendor="Supplier",
            document_type_code="DT-01",
            invoice_no="INV-CATALOG-FALLBACK",
            invoice_date=date(2026, 7, 8),
            status=InvoiceStatus.PROCESSED,
            file_hash="bundle-empty-catalogue",
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-07-01&date_to=2026-07-31"
    )
    assert res.status_code == 200
    assert res.headers.get("x-data-rows") == "0"
    header, data, _ws = _read_xlsx(res.content)
    assert len(header) == 13
    assert "Timestamp" in header
    assert "Uploaded by" in header
    assert "Source" in header
    assert data == []


@pytest.mark.asyncio
async def test_documents_bundle_export_dual_linkage_invoice_no_and_po_reference(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Buyer",
        document_type_code="DT-01",
        invoice_no="INV-DUAL-100",
        po_reference="PO-DUAL-99",
        purchase_document_type=PurchaseDocumentType.INVOICE.value,
        route_target=ROUTE_PURCHASE,
        invoice_date=date(2026, 7, 10),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-dual-anchor",
    )
    invoice_no_sibling = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Carrier",
        document_type_code="DT-06",
        invoice_no="INV-DUAL-100",
        po_reference="OTHER-PO",
        invoice_date=date(2026, 7, 9),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-dual-inv-sibling",
    )
    po_sibling = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Supplier",
        document_type_code="DT-02",
        invoice_no="PO-DOC-ONLY",
        po_reference="PO-DUAL-99",
        purchase_document_type=PurchaseDocumentType.PO.value,
        invoice_date=date(2026, 7, 8),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-dual-po-sibling",
    )
    db_session.add_all([anchor, invoice_no_sibling, po_sibling])
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-07-01&date_to=2026-07-31"
    )
    assert res.status_code == 200
    header, data, ws = _read_xlsx(res.content)
    row_index = next(
        i
        for i, r in enumerate(data)
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-DUAL-100"
    )

    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="Proforma / advance",
        url=vault_view_path(invoice_no_sibling.id),
    )
    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="PO (supporting)",
        url=vault_view_path(po_sibling.id),
    )


@pytest.mark.asyncio
async def test_documents_bundle_export_po_reference_case_insensitive(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Buyer",
        document_type_code="DT-01",
        invoice_no="INV-CASE-BUNDLE",
        po_reference="PO-Case-Bundle",
        purchase_document_type=PurchaseDocumentType.INVOICE.value,
        route_target=ROUTE_PURCHASE,
        invoice_date=date(2026, 7, 10),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-case-anchor",
    )
    po_sibling = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Supplier",
        document_type_code="DT-02",
        invoice_no="PO-CASE-ONLY",
        po_reference="po-case-bundle",
        purchase_document_type=PurchaseDocumentType.PO.value,
        invoice_date=date(2026, 7, 8),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-case-po-sibling",
    )
    db_session.add_all([anchor, po_sibling])
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-07-01&date_to=2026-07-31"
    )
    assert res.status_code == 200
    header, data, ws = _read_xlsx(res.content)
    row_index = next(
        i
        for i, r in enumerate(data)
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-CASE-BUNDLE"
    )
    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="PO (supporting)",
        url=vault_view_path(po_sibling.id),
    )


def test_linked_docs_helpers_still_support_audit_export() -> None:
    from app.schemas.dossier import DossierLinkedDocumentResponse

    linked = DossierLinkedDocumentsResponse(
        linkage_kind="invoice_no",
        linkage_label="Invoice no.",
        enforce_bundle=False,
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-01",
                label="Anchor",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=1,
            ),
            DossierLinkedDocumentResponse(
                id="sibling",
                document_type_code="DT-26",
                label="Sibling",
                document_ref="DOC-9",
                present=True,
                requirement="advisory",
                link_kind="invoice_no",
                invoice_id=9,
            ),
        ],
    )
    entries = collect_linked_doc_entries(1, linked)
    assert len(entries) == 1
    assert entries[0].invoice_id == 9

    audit_cell = format_linked_docs_export(1, linked)
    assert "DT-26:DOC-9" in audit_cell

    by_dt = linked_docs_by_dt_code(1, linked, label_fn="bundle")
    assert "DT-26" in by_dt
    assert "DOC-9" in by_dt["DT-26"]


def test_bundle_dt_cells_marks_missing_mandatory_slots() -> None:
    from app.schemas.dossier import DossierLinkedDocumentResponse

    linked = DossierLinkedDocumentsResponse(
        linkage_kind="po_reference",
        linkage_label="PO reference",
        enforce_bundle=True,
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-01",
                label="Anchor",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=1,
            ),
            DossierLinkedDocumentResponse(
                id="po-slot",
                document_type_code="DT-02",
                label="PO",
                present=False,
                requirement="mandatory",
            ),
            DossierLinkedDocumentResponse(
                id="grn-slot",
                document_type_code="DT-03",
                label="GRN",
                present=False,
                requirement="mandatory",
            ),
        ],
    )
    cells = bundle_dt_cells_by_code(1, linked, ["DT-02", "DT-03", "DT-26"])
    assert cells["DT-02"] == "Missing"
    assert cells["DT-03"] == "Missing"
    assert cells["DT-26"] == ""


def test_bundle_dt_cells_marks_advisory_slots() -> None:
    from app.schemas.dossier import DossierLinkedDocumentResponse

    linked = DossierLinkedDocumentsResponse(
        linkage_kind="invoice_no",
        linkage_label="Invoice no.",
        enforce_bundle=False,
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-01",
                label="Anchor",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=1,
            ),
            DossierLinkedDocumentResponse(
                id="advisory",
                document_type_code="DT-26",
                label="Customer invoice",
                present=False,
                requirement="advisory",
            ),
        ],
    )
    cells = bundle_dt_cells_by_code(1, linked, ["DT-26"])
    assert cells["DT-26"] == "Advisory"


def test_collect_linked_doc_entries_uses_slot_dt_code_for_manual_overlay() -> None:
    from app.schemas.dossier import (
        DossierLinkedDocumentResponse,
        DossierManualLinkInfoResponse,
    )

    linked = DossierLinkedDocumentsResponse(
        linkage_kind="shipment_ref",
        linkage_label="Import dossier",
        enforce_bundle=True,
        documents=[
            DossierLinkedDocumentResponse(
                id="DT-27-bundle",
                document_type_code="DT-27",
                label="Bill of lading",
                present=False,
                requirement="mandatory",
                manual_link=DossierManualLinkInfoResponse(
                    id=1,
                    invoice_id=42,
                    linked_dossier_id="DOC-42",
                    document_ref="BL-1",
                    label="Unclassified upload",
                    document_type_code="DT-06",
                    has_file=True,
                ),
            ),
        ],
    )
    entries = collect_linked_doc_entries(1, linked)
    assert len(entries) == 1
    assert entries[0].invoice_id == 42
    assert entries[0].document_type_code == "DT-27"
    assert entries[0].pin_dt_code is True


def test_bundle_dt_cells_shows_manual_link_instead_of_missing() -> None:
    from app.schemas.dossier import (
        DossierLinkedDocumentResponse,
        DossierManualLinkInfoResponse,
    )

    linked = DossierLinkedDocumentsResponse(
        linkage_kind="shipment_ref",
        linkage_label="Import dossier",
        enforce_bundle=True,
        documents=[
            DossierLinkedDocumentResponse(
                id="anchor",
                document_type_code="DT-10",
                label="Import invoice",
                present=True,
                requirement="mandatory",
                is_anchor=True,
                invoice_id=1,
            ),
            DossierLinkedDocumentResponse(
                id="DT-27-bundle",
                document_type_code="DT-27",
                label="Bill of lading",
                present=False,
                requirement="mandatory",
                manual_link=DossierManualLinkInfoResponse(
                    id=1,
                    invoice_id=42,
                    linked_dossier_id="DOC-42",
                    document_ref="BL-4402",
                    label="Bill of lading",
                    document_type_code="DT-27",
                    has_file=True,
                ),
            ),
        ],
    )
    cells = bundle_dt_cells_by_code(1, linked, ["DT-27"])
    assert cells["DT-27"] != "Missing"
    assert "BL-4402" in cells["DT-27"]
    assert_csv_hyperlink(cells["DT-27"], url=vault_view_path(42))


@pytest.mark.asyncio
async def test_documents_bundle_export_includes_manual_slot_link(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    from app.services.dossier.dossier_manual_link_service import create_manual_link

    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Importer",
        document_type_code="DT-01",
        invoice_no="INV-MANUAL-BUNDLE",
        invoice_date=date(2026, 9, 1),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-manual-anchor",
    )
    support = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Supplier",
        document_type_code="DT-02",
        invoice_no="PO-MANUAL-1",
        invoice_date=date(2026, 9, 1),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-manual-po",
    )
    db_session.add_all([anchor, support])
    await db_session.flush()

    await create_manual_link(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        anchor_invoice_id=anchor.id,
        linked_invoice_id=support.id,
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-09-01&date_to=2026-09-30"
    )
    assert res.status_code == 200
    header, data, ws = _read_xlsx(res.content)
    row_index = next(
        i
        for i, r in enumerate(data)
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-MANUAL-BUNDLE"
    )
    assert data[row_index][header.index("PO (supporting)")] != "Missing"
    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="PO (supporting)",
        url=vault_view_path(support.id),
    )


@pytest.mark.asyncio
async def test_documents_bundle_export_missing_mandatory_in_xlsx(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    db_session.add(
        Invoice(
            tenant_id=TESTING_TENANT_UUID,
            vendor="Supplier",
            document_type_code="DT-01",
            invoice_no="INV-MISSING-BUNDLE",
            invoice_date=date(2026, 8, 1),
            status=InvoiceStatus.PROCESSED,
            file_hash="bundle-missing-mandatory",
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-08-01&date_to=2026-08-31"
    )
    assert res.status_code == 200
    header, data, _ws = _read_xlsx(res.content)
    row = next(
        r for r in data if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-MISSING-BUNDLE"
    )
    assert row[header.index("PO (supporting)")] == "Missing"
    assert row[header.index("GRN (supporting)")] == "Missing"


@pytest.mark.asyncio
async def test_documents_bundle_export_plain_format(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    inv = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Plain Vendor",
        document_type_code="DT-01",
        invoice_no="INV-PLAIN-1",
        invoice_date=date(2026, 8, 5),
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-plain-format",
    )
    db_session.add(inv)
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-08-01&date_to=2026-08-31&format=plain"
    )
    assert res.status_code == 200
    header, data, ws = _read_xlsx(res.content)
    row_index = next(
        i
        for i, r in enumerate(data)
        if "INV-PLAIN-1" in _cell_text(r[header.index("Invoice no.")])
    )
    inv_cell = _cell_text(data[row_index][header.index("Invoice no.")])
    assert not inv_cell.startswith("=HYPERLINK(")
    assert "INV-PLAIN-1 |" in inv_cell
    assert "/vault?invoice=" in inv_cell
    excel_row = _HEADER_ROW + 1 + row_index
    cell = ws.cell(row=excel_row, column=header.index("Invoice no.") + 1)
    assert cell.hyperlink is None


@pytest.mark.asyncio
async def test_documents_bundle_export_includes_all_org_document_types(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    from app.services.invoice.invoice_evaluation_service import load_config_for_tenant

    org_types = await load_config_for_tenant(db_session, TESTING_TENANT_UUID)

    db_session.add(
        Invoice(
            tenant_id=TESTING_TENANT_UUID,
            vendor="Supplier",
            document_type_code="DT-01",
            invoice_no="INV-FULL-COLS",
            invoice_date=date(2026, 8, 10),
            status=InvoiceStatus.PROCESSED,
            file_hash="bundle-full-cols",
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-08-01&date_to=2026-08-31"
    )
    assert res.status_code == 200
    header, _data, _ws = _read_xlsx(res.content)
    fixed_count = 13
    dt_column_count = len(header) - fixed_count
    assert header[header.index("Timestamp")] == "Timestamp"
    assert "Uploaded by" in header
    assert "Source" in header
    assert dt_column_count == len(org_types.document_types)
    shipped_only_codes = {"DT-14", "DT-15", "DT-26", "DT-27", "DT-28"}
    header_labels = set(header[fixed_count:])
    for code in shipped_only_codes:
        assert not any(code in label for label in header_labels)


@pytest.mark.asyncio
async def test_documents_bundle_export_includes_custom_tenant_document_type_column(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    from app.models.tenant_rule_book_config import TenantRuleBookConfig
    from app.schemas.document_type import DocumentTypeClassifier, DocumentTypeDefinition
    from app.services.invoice.invoice_evaluation_service import load_config_for_tenant
    from app.services.rule_book.rule_book_config_io import clear_posting_config_cache

    base_config = await load_config_for_tenant(db_session, TESTING_TENANT_UUID)
    dt01 = next(dt for dt in base_config.document_types if dt.code == "DT-01")
    custom_dt = DocumentTypeDefinition(
        code="DT-99",
        title="Custom freight note",
        short_title="Freight note",
        klass="Supporting",
        posting="No",
        recognition_mode="signals",
        recognition_signals=["heading_invoice"],
        llm_prompt="",
        route_target="Vault",
        classifier=DocumentTypeClassifier(enabled=False, priority=100, confidence=0.85),
    )
    row = await db_session.get(TenantRuleBookConfig, TESTING_TENANT_UUID)
    assert row is not None
    row.config = {
        **row.config,
        "document_types": [
            dt01.model_dump(by_alias=True),
            custom_dt.model_dump(by_alias=True),
        ],
    }
    await db_session.flush()
    clear_posting_config_cache()

    db_session.add(
        Invoice(
            tenant_id=TESTING_TENANT_UUID,
            vendor="Supplier",
            document_type_code="DT-01",
            invoice_no="INV-CUSTOM-DT-COL",
            invoice_date=date(2026, 8, 11),
            status=InvoiceStatus.PROCESSED,
            file_hash="bundle-custom-dt-col",
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-08-01&date_to=2026-08-31"
    )
    assert res.status_code == 200
    header, data, _ws = _read_xlsx(res.content)
    assert "Freight note" in header
    bundle_row = next(
        r
        for r in data
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-CUSTOM-DT-COL"
    )
    assert bundle_row[header.index("Freight note")] in (None, "")


@pytest.mark.asyncio
async def test_documents_bundle_export_includes_exception_anchor_with_po_sibling(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    po_doc = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme Supplies",
        document_type_code="DT-02",
        po_reference="PO-EXC-BUNDLE",
        invoice_no="PO-EXC-BUNDLE",
        invoice_date=date(2026, 9, 5),
        route_target=ROUTE_PURCHASE,
        purchase_document_type=PurchaseDocumentType.PO.value,
        status=InvoiceStatus.PROCESSED,
        file_hash="bundle-exc-po-sibling",
    )
    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme Supplies",
        document_type_code="DT-01",
        po_reference="PO-EXC-BUNDLE",
        invoice_no="INV-EXC-BUNDLE",
        invoice_date=date(2026, 9, 10),
        route_target=ROUTE_PURCHASE,
        purchase_document_type=PurchaseDocumentType.INVOICE.value,
        status=InvoiceStatus.EXCEPTION,
        file_hash="bundle-exc-anchor",
    )
    db_session.add_all([po_doc, anchor])
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-09-01&date_to=2026-09-30"
    )
    assert res.status_code == 200
    header, data, ws = _read_xlsx(res.content)
    assert "Status" not in header
    row_index = next(
        i
        for i, r in enumerate(data)
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "INV-EXC-BUNDLE"
    )
    _assert_hyperlink_cell(
        ws,
        header=header,
        data_row_index=row_index,
        column_name="PO (supporting)",
        url=vault_view_path(po_doc.id),
    )


@pytest.mark.asyncio
async def test_documents_bundle_export_excludes_pending_anchor(
    client: AsyncClient,
    db_session: AsyncSession,
) -> None:
    db_session.add_all(
        [
            Invoice(
                tenant_id=TESTING_TENANT_UUID,
                vendor="Supplier",
                document_type_code="DT-01",
                invoice_no="INV-PENDING-BUNDLE",
                invoice_date=date(2026, 9, 12),
                status=InvoiceStatus.PENDING,
                file_hash="bundle-pending-anchor",
            ),
            Invoice(
                tenant_id=TESTING_TENANT_UUID,
                vendor="Supplier",
                document_type_code="DT-01",
                invoice_no="INV-PARSING-BUNDLE",
                invoice_date=date(2026, 9, 13),
                status=InvoiceStatus.PARSING,
                file_hash="bundle-parsing-anchor",
            ),
        ]
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/documents-bundle/export"
        "?date_from=2026-09-01&date_to=2026-09-30"
    )
    assert res.status_code == 200
    header, data, _ws = _read_xlsx(res.content)
    invoice_nos = _invoice_nos_from_rows(header, data)
    assert "INV-PENDING-BUNDLE" not in invoice_nos
    assert "INV-PARSING-BUNDLE" not in invoice_nos


@pytest.mark.asyncio
async def test_documents_bundle_export_includes_validating_anchor(
    db_session: AsyncSession,
) -> None:
    """Active-workflow anchors (e.g. validating) still export; Status is not a column."""
    anchor = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Importer",
        status=InvoiceStatus.VALIDATING,
        document_type_code="DT-01",
        invoice_no="SHARED-VAL-88",
        invoice_date=date(2026, 9, 20),
        file_hash="bundle-validating-anchor",
    )
    sibling = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Carrier",
        status=InvoiceStatus.PROCESSED,
        document_type_code="DT-06",
        invoice_no="SHARED-VAL-88",
        invoice_date=date(2026, 9, 20),
        file_hash="bundle-validating-sibling",
    )
    db_session.add_all([anchor, sibling])
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 9, 1),
        date_to=date(2026, 9, 30),
    )
    header, data, _ws = _read_xlsx(payload.xlsx_bytes)
    assert "Status" not in header
    row = next(
        r
        for r in data
        if _invoice_no_from_cell(r[header.index("Invoice no.")]) == "SHARED-VAL-88"
    )
    assert row[header.index("Universal match")] == "Yes"
