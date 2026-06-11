"""Tests for Excel workbook export."""

from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.models.invoice import Invoice, InvoiceStatus
from app.models.journal import EntryType, JournalEntry
from app.models.line_item import LineItem
from app.services.account_mapper import clear_rule_book_cache, map_with_details
from app.services.rule_book_mapper import clear_classification_config_cache
from app.services.workbook_writer import (
    SHEET_DAILY_RECON,
    SHEET_EXPENSE_SUMMARY,
    SHEET_INVOICES,
    SHEET_JOURNAL_ENTRIES,
    SHEET_LEDGER_MAPPING,
    SHEET_LINE_ITEMS,
    SHEET_PROCESSING_STATUS,
    SHEET_RULE_BOOK,
    workbook_filename,
    write_workbook,
)


@pytest.fixture(autouse=True)
def _clear_settings() -> None:
    get_settings.cache_clear()
    clear_rule_book_cache()
    clear_classification_config_cache()
    yield
    clear_classification_config_cache()
    clear_rule_book_cache()
    get_settings.cache_clear()


@pytest.mark.asyncio
async def test_map_with_details_po_match() -> None:
    inv = Invoice(org_id=1,
        vendor="Google Australia Pty Ltd",
        invoice_no="GOOG-AU-99102",
        po_reference="PO-MKT-2026-014",
        route_target="Purchase Management",
        status=InvoiceStatus.MAPPING,
        currency="AUD",
    )
    detail = map_with_details(inv, line_description="Campaign spend")
    assert detail.rule_type == "Purchase rule"
    assert detail.expense_category == "Marketing Expense"


@pytest.mark.asyncio
async def test_write_workbook_sheets(db_session: AsyncSession, tmp_path: Path) -> None:
    get_settings.cache_clear()
    import app.services.workbook_writer as ww

    original = Path(get_settings().upload_dir)

    inv = Invoice(org_id=1,
        vendor="Google Australia Pty Ltd",
        abn="33102417032",
        invoice_no="GOOG-INV-7781032",
        invoice_date=date(2026, 5, 6),
        due_date=date(2026, 6, 5),
        currency="AUD",
        subtotal=Decimal("5350.00"),
        gst=Decimal("535.00"),
        total=Decimal("5885.00"),
        status=InvoiceStatus.PROCESSED,
        file_hash="wb_test_hash",
        account_code="6130",
        account_name="Marketing Expense",
        validation_results='[{"rule":"VR03","passed":true,"message":"OK"}]',
    )
    db_session.add(inv)
    await db_session.flush()

    db_session.add_all(
        [
            LineItem(
                invoice_id=inv.id,
                description="Google Ads - May",
                qty=Decimal("1"),
                unit_price=Decimal("5350.00"),
                amount=Decimal("5350.00"),
            ),
            JournalEntry(
                invoice_id=inv.id,
                date=date(2026, 5, 6),
                account_code="6130",
                account_name="Marketing Expense",
                debit=Decimal("5350.00"),
                credit=Decimal("0"),
                entry_type=EntryType.DEBIT,
            ),
            JournalEntry(
                invoice_id=inv.id,
                date=date(2026, 5, 6),
                account_code="1400",
                account_name="GST Paid",
                debit=Decimal("535.00"),
                credit=Decimal("0"),
                entry_type=EntryType.DEBIT,
            ),
            JournalEntry(
                invoice_id=inv.id,
                date=date(2026, 5, 6),
                account_code="2000",
                account_name="Accounts Payable",
                debit=Decimal("0"),
                credit=Decimal("5885.00"),
                entry_type=EntryType.CREDIT,
            ),
        ]
    )
    await db_session.flush()

    reports_dir = tmp_path / "reports"
    reports_dir.mkdir(parents=True)

    def fake_reports_dir() -> Path:
        return reports_dir

    ww._reports_dir = fake_reports_dir  # type: ignore[method-assign]

    path = await write_workbook(db_session, 1, date(2026, 5, 6))
    assert path.is_file()
    assert path.name == "output_workbook_hv-org_2026-05-06.xlsx"

    wb = load_workbook(path, read_only=True)
    assert "README" in wb.sheetnames
    assert SHEET_INVOICES in wb.sheetnames
    assert SHEET_LINE_ITEMS in wb.sheetnames
    assert SHEET_LEDGER_MAPPING in wb.sheetnames
    assert SHEET_JOURNAL_ENTRIES in wb.sheetnames
    assert SHEET_DAILY_RECON in wb.sheetnames
    assert SHEET_EXPENSE_SUMMARY in wb.sheetnames
    assert SHEET_PROCESSING_STATUS in wb.sheetnames
    assert SHEET_RULE_BOOK in wb.sheetnames
    wb.close()


def test_workbook_filename_range() -> None:
    slug = "hv-org"
    assert workbook_filename(slug) == "output_workbook_hv-org.xlsx"
    assert (
        workbook_filename(slug, date(2026, 5, 6), date(2026, 5, 6))
        == "output_workbook_hv-org_2026-05-06.xlsx"
    )
    assert (
        workbook_filename(slug, date(2026, 5, 1), date(2026, 5, 31))
        == "output_workbook_hv-org_2026-05-01_to_2026-05-31.xlsx"
    )
