"""Catalog, favourites, preview, and export API tests."""

from datetime import date, timedelta
import calendar
from decimal import Decimal
from uuid import UUID

import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.invoice import Invoice, InvoiceStatus
from app.models.journal import EntryType, JournalEntry
from app.models.tenant import Tenant
from app.services.reports.report_catalog import CATALOG_BY_ID
from app.services.reports.report_export_service import PDF_ROWS_PER_PAGE
from app.tenant_ids import TESTING_TENANT_UUID


def _today() -> date:
    return date.today()


async def _processed_invoice(db: AsyncSession, **kwargs) -> Invoice:
    payload = {
        "tenant_id": TESTING_TENANT_UUID,
        "vendor": "Acme",
        "invoice_no": "INV-CAT-1",
        "invoice_date": _today(),
        "due_date": _today() + timedelta(days=14),
        "subtotal": Decimal("100.00"),
        "gst": Decimal("10.00"),
        "total": Decimal("110.00"),
        "currency": "AUD",
        "status": InvoiceStatus.PROCESSED,
        "file_hash": "cat-inv-1",
    }
    payload.update(kwargs)
    inv = Invoice(**payload)
    db.add(inv)
    await db.flush()
    return inv


async def _journal(
    db: AsyncSession,
    invoice: Invoice,
    *,
    account_code: str,
    account_name: str,
    debit: Decimal = Decimal("0"),
    credit: Decimal = Decimal("0"),
    on: date | None = None,
    tenant_id=TESTING_TENANT_UUID,
) -> JournalEntry:
    entry = JournalEntry(
        tenant_id=tenant_id,
        invoice_id=invoice.id,
        date=on or _today(),
        account_code=account_code,
        account_name=account_name,
        debit=debit,
        credit=credit,
        entry_type=EntryType.DEBIT if debit else EntryType.CREDIT,
    )
    db.add(entry)
    await db.flush()
    return entry


@pytest.mark.asyncio
async def test_reports_catalog_lists_seed_set(client: AsyncClient) -> None:
    res = await client.get("/api/reports/catalog")
    assert res.status_code == 200
    data = res.json()["data"]
    ids = [item["id"] for item in data["reports"]]
    assert ids == list(CATALOG_BY_ID)
    assert data["favourite_ids"] == []
    budget = next(item for item in data["reports"] if item["id"] == "budget-variance")
    assert budget["supports_compare"] is True
    aged = next(item for item in data["reports"] if item["id"] == "aged-payables")
    assert aged["supports_compare"] is False
    for removed in (
        "profit-and-loss",
        "balance-sheet",
        "cash-flow",
        "changes-in-equity",
        "aged-receivables",
        "business-snapshot",
        "trial-balance",
        "gl-summary",
        "sales-tax-summary",
        "account-transactions",
        "bank-reconciliation",
    ):
        assert removed not in ids


@pytest.mark.asyncio
async def test_favourites_persist_and_reject_unknown(client: AsyncClient) -> None:
    bad = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["invoice-register", "not-a-real-report"]},
    )
    assert bad.status_code == 422

    ok = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["invoice-register", "aged-payables"]},
    )
    assert ok.status_code == 200
    assert ok.json()["data"] == ["invoice-register", "aged-payables"]

    catalog = await client.get("/api/reports/catalog")
    assert catalog.status_code == 200
    assert set(catalog.json()["data"]["favourite_ids"]) == {
        "invoice-register",
        "aged-payables",
    }


@pytest.mark.asyncio
async def test_preview_and_export_unknown_report_id(client: AsyncClient) -> None:
    preview = await client.get("/api/reports/not-a-real-report/preview")
    assert preview.status_code == 404

    export = await client.post(
        "/api/reports/not-a-real-report/export",
        json={"format": "xlsx", "range": "month"},
    )
    assert export.status_code == 404


@pytest.mark.asyncio
async def test_preview_empty_then_invoice_register_with_invoice(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    empty = await client.get("/api/reports/invoice-register/preview?range=month")
    assert empty.status_code == 200
    payload = empty.json()["data"]
    assert payload["empty"] is True
    assert payload["rows"] == []

    await _processed_invoice(
        db_session,
        file_hash="cat-register",
        vendor="Acme Register Co",
        invoice_no="INV-REG-1",
    )
    await db_session.commit()

    filled = await client.get("/api/reports/invoice-register/preview?range=month")
    assert filled.status_code == 200
    data = filled.json()["data"]
    assert data["empty"] is False
    assert data["report_id"] == "invoice-register"
    joined = " ".join(cell for row in data["rows"] for cell in row["cells"])
    assert "Acme Register Co" in joined
    assert "INV-REG-1" in joined


@pytest.mark.asyncio
async def test_export_xlsx_and_pdf_content_types(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _processed_invoice(db_session, file_hash="cat-export", invoice_no="EXP-1")
    await db_session.commit()

    xlsx = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "xlsx", "range": "month"},
    )
    assert xlsx.status_code == 200
    assert "spreadsheetml" in xlsx.headers["content-type"]
    assert xlsx.headers["content-disposition"].endswith('.xlsx"')

    pdf = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "pdf", "range": "month"},
    )
    assert pdf.status_code == 200
    assert pdf.headers["content-type"] == "application/pdf"
    assert pdf.content[:4] == b"%PDF"


@pytest.mark.asyncio
async def test_pdf_export_paginates_when_rows_exceed_threshold(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    extra = PDF_ROWS_PER_PAGE + 8
    for i in range(extra):
        await _processed_invoice(
            db_session,
            file_hash=f"cat-pdf-pages-{i}",
            invoice_no=f"PDF-{i:03d}",
            vendor=f"Vendor {i:03d}",
        )
    await db_session.commit()

    res = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "pdf", "range": "month"},
    )
    assert res.status_code == 200
    assert int(res.headers["x-data-rows"]) == extra
    assert int(res.headers["x-page-count"]) >= 2

    import fitz

    doc = fitz.open(stream=res.content, filetype="pdf")
    try:
        assert doc.page_count >= 2
        assert extra > PDF_ROWS_PER_PAGE
    finally:
        doc.close()


@pytest.mark.asyncio
async def test_preview_ignores_other_tenant_invoices(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    other_id = UUID("aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeeee")
    other = Tenant(
        id=other_id,
        name="Other Co",
        slug="other-co",
        currency="AUD",
        settings_json={"country": "AU"},
    )
    db_session.add(other)
    await db_session.flush()
    foreign_inv = Invoice(
        tenant_id=other.id,
        vendor="Foreign",
        invoice_no="INV-OTHER",
        invoice_date=_today(),
        total=Decimal("999.00"),
        currency="AUD",
        status=InvoiceStatus.PROCESSED,
        file_hash="other-tenant-inv",
    )
    db_session.add(foreign_inv)
    await db_session.flush()
    await db_session.commit()

    res = await client.get("/api/reports/invoice-register/preview?range=month")
    assert res.status_code == 200
    joined = " ".join(
        cell for row in res.json()["data"]["rows"] for cell in row["cells"]
    )
    assert "Should not appear" not in joined
    assert "Foreign" not in joined
    assert "INV-OTHER" not in joined


def _joined(payload: dict) -> str:
    return " ".join(cell for row in payload["rows"] for cell in row["cells"])


@pytest.mark.asyncio
async def test_invoice_register_blank_payment_cells_without_payment_row(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    inv = await _processed_invoice(
        db_session,
        file_hash="reg-nopay",
        vendor="NoPay Co",
        invoice_no="NP-1",
        invoice_date=_today(),
        due_date=_today() + timedelta(days=7),
    )
    await db_session.commit()

    res = await client.get("/api/reports/invoice-register/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    match = next(row for row in data["rows"] if "NP-1" in row["cells"])
    assert match["cells"][-2] == ""
    assert match["cells"][-1] == ""
    assert "null" not in match["cells"][-2].lower()
    assert "none" not in match["cells"][-1].lower()
    assert inv.id is not None


@pytest.mark.asyncio
async def test_ap_reports_exclude_team_expenses_route(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.services.invoice.invoice_evaluation_service import ROUTE_TEAM

    ap = await _processed_invoice(
        db_session,
        file_hash="ap-keep",
        vendor="AP Keep Vendor",
        invoice_no="AP-KEEP-1",
        invoice_date=_today(),
        due_date=_today() - timedelta(days=3),
    )
    await _journal(
        db_session,
        ap,
        account_code="2000",
        account_name="Accounts Payable",
        credit=Decimal("110.00"),
    )
    te = await _processed_invoice(
        db_session,
        file_hash="te-excl",
        vendor="TE Exclude Vendor",
        invoice_no="TE-EXCL-1",
        invoice_date=_today(),
        due_date=_today() - timedelta(days=3),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
    )
    await _journal(
        db_session,
        te,
        account_code="2000",
        account_name="Accounts Payable",
        credit=Decimal("110.00"),
    )
    await db_session.commit()

    for report_id in (
        "aged-payables",
        "payment-schedule",
        "invoice-register",
        "vendor-spend-summary",
    ):
        res = await client.get(f"/api/reports/{report_id}/preview?range=month")
        assert res.status_code == 200, report_id
        blob = _joined(res.json()["data"])
        assert "TE-EXCL-1" not in blob, report_id
        assert "TE Exclude Vendor" not in blob, report_id
        assert "AP-KEEP-1" in blob or "AP Keep Vendor" in blob, report_id


@pytest.mark.asyncio
async def test_aged_payables_fills_document_column(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    inv = await _processed_invoice(
        db_session,
        file_hash="aged-doc",
        vendor="Aged Doc Co",
        invoice_no="AGED-DOC-9",
        due_date=_today() - timedelta(days=10),
    )
    await _journal(
        db_session,
        inv,
        account_code="2000",
        account_name="Accounts Payable",
        credit=Decimal("50.00"),
    )
    await db_session.commit()

    res = await client.get("/api/reports/aged-payables/preview?range=month")
    assert res.status_code == 200
    assert "AGED-DOC-9" in _joined(res.json()["data"])


@pytest.mark.asyncio
async def test_cash_forecast_notes_and_payment_schedule_buckets(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    today = _today()
    month_end = date(today.year, today.month, calendar.monthrange(today.year, today.month)[1])
    overdue = await _processed_invoice(
        db_session,
        file_hash="sched-over",
        vendor="Sched Co",
        invoice_no="SCH-OVER",
        invoice_date=today,
        due_date=month_end - timedelta(days=5),
    )
    soon = await _processed_invoice(
        db_session,
        file_hash="sched-soon",
        vendor="Sched Co",
        invoice_no="SCH-SOON",
        invoice_date=today,
        due_date=month_end + timedelta(days=2),
    )
    later = await _processed_invoice(
        db_session,
        file_hash="sched-up",
        vendor="Sched Co",
        invoice_no="SCH-UP",
        invoice_date=today,
        due_date=month_end + timedelta(days=21),
    )
    for inv in (overdue, soon, later):
        await _journal(
            db_session,
            inv,
            account_code="2000",
            account_name="Accounts Payable",
            credit=Decimal("110.00"),
        )
    await db_session.commit()
    assert overdue.id and soon.id and later.id

    sched = await client.get("/api/reports/payment-schedule/preview?range=month")
    assert sched.status_code == 200
    blob = _joined(sched.json()["data"])
    assert "OVERDUE" in blob
    assert "Due within 7 days" in blob
    assert "Upcoming" in blob
    notes = sched.json()["data"].get("notes") or ""
    assert "Overdue total" in notes
    assert "Due within 7 days" in notes

    forecast = await client.get("/api/reports/cash-forecast/preview?range=month")
    assert forecast.status_code == 200
    assert forecast.json()["data"]["notes"]
    assert "reimbursement" in forecast.json()["data"]["notes"].lower()


@pytest.mark.asyncio
async def test_budget_variance_committed_column(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.services.invoice.invoice_evaluation_service import ROUTE_TEAM
    from app.services.purchase.team_expense_spend_service import current_period_keys

    today = _today()
    keys = current_period_keys(today)
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="Travel",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("1000.00"),
            enforcement="soft",
        )
    )
    actual = await _processed_invoice(
        db_session,
        file_hash="bv-actual",
        vendor="Staff",
        invoice_no="BV-ACT",
        invoice_date=today,
        total=Decimal("100.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        employee_email="traveller@example.com",
    )
    committed = await _processed_invoice(
        db_session,
        file_hash="bv-commit",
        vendor="Staff",
        invoice_no="BV-COM",
        invoice_date=today,
        total=Decimal("40.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        evaluation_status="",
        employee_email="traveller@example.com",
    )
    await db_session.commit()
    assert actual.status == InvoiceStatus.PROCESSED
    assert committed.status == InvoiceStatus.VALIDATING

    res = await client.get("/api/reports/budget-variance/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert "Committed" in data["columns"]
    assert "team expense" in (data.get("notes") or "").lower()
    cols = data["columns"]
    row = next(
        r["cells"] for r in data["rows"] if r["cells"][cols.index("Category")] == "Travel"
    )
    assert row[cols.index("Budget")] == "1,000.00"
    assert row[cols.index("Committed")] == "40.00"
    assert row[cols.index("Actual")] == "100.00"


@pytest.mark.asyncio
async def test_advance_aging_skips_overclaimed_employee(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.employee_master import EmployeeMasterRecord
    from app.services.invoice.invoice_evaluation_service import ROUTE_TEAM
    from app.services.master_data.party_coa_subledger_service import party_sub_ledger_code

    aged = EmployeeMasterRecord(
        tenant_id=TESTING_TENANT_UUID,
        master_id="emaged",
        name="Aged Emp",
        email="aged@example.com",
        advance_parent_ledger="Staff Advance",
        status="Active",
    )
    over = EmployeeMasterRecord(
        tenant_id=TESTING_TENANT_UUID,
        master_id="emover",
        name="Over Claimed",
        email="over@example.com",
        advance_parent_ledger="Staff Advance",
        status="Active",
    )
    db_session.add_all([aged, over])
    await db_session.flush()

    aged_inv = await _processed_invoice(
        db_session,
        file_hash="adv-aged",
        vendor="Aged Emp",
        invoice_no="ADV-AGED",
        invoice_date=_today() - timedelta(days=100),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="aged@example.com",
    )
    over_take = await _processed_invoice(
        db_session,
        file_hash="adv-over-take",
        vendor="Over Claimed",
        invoice_no="ADV-TAKE",
        invoice_date=_today() - timedelta(days=20),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="over@example.com",
    )
    over_use = await _processed_invoice(
        db_session,
        file_hash="adv-over-use",
        vendor="Over Claimed",
        invoice_no="ADV-USE",
        invoice_date=_today() - timedelta(days=5),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="over@example.com",
    )
    aged_code = party_sub_ledger_code("emaged")
    over_code = party_sub_ledger_code("emover")
    await _journal(
        db_session,
        aged_inv,
        account_code=aged_code,
        account_name="Staff Advance",
        debit=Decimal("80.00"),
        on=_today() - timedelta(days=100),
    )
    await _journal(
        db_session,
        over_take,
        account_code=over_code,
        account_name="Staff Advance",
        debit=Decimal("100.00"),
        on=_today() - timedelta(days=20),
    )
    await _journal(
        db_session,
        over_use,
        account_code=over_code,
        account_name="Staff Advance",
        credit=Decimal("150.00"),
        on=_today() - timedelta(days=5),
    )
    await db_session.commit()

    res = await client.get("/api/reports/advance-aging/preview?range=month")
    assert res.status_code == 200
    blob = _joined(res.json()["data"])
    assert "Aged Emp" in blob
    assert "90+" in res.json()["data"]["columns"]
    assert "80.00" in blob
    assert "Over Claimed" not in blob


@pytest.mark.asyncio
async def test_expense_claims_register_xlsx_respects_custom_date_range(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models.employee_master import EmployeeMasterRecord
    from app.services.invoice.invoice_evaluation_service import ROUTE_TEAM

    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emrange",
            name="Range Emp",
            email="range@example.com",
            status="Active",
        )
    )
    await _processed_invoice(
        db_session,
        file_hash="te-in-range",
        vendor="Range Emp",
        invoice_no="IN-RANGE",
        invoice_date=date(2026, 1, 15),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="range@example.com",
        total=Decimal("25.00"),
    )
    await _processed_invoice(
        db_session,
        file_hash="te-out-range",
        vendor="Range Emp",
        invoice_no="OUT-RANGE",
        invoice_date=date(2026, 6, 15),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="range@example.com",
        total=Decimal("99.00"),
    )
    await db_session.commit()

    res = await client.post(
        "/api/reports/expense-claims-register/export",
        json={
            "format": "xlsx",
            "range": "custom",
            "from": "2026-01-01",
            "to": "2026-01-31",
        },
    )
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    wb = load_workbook(BytesIO(res.content))
    values = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            values.extend("" if cell is None else str(cell) for cell in row)
    blob = " ".join(values)
    assert "IN-RANGE" in blob
    assert "OUT-RANGE" not in blob
    assert "Profit & Loss" not in blob


@pytest.mark.asyncio
async def test_catalog_includes_new_report_ids(client: AsyncClient) -> None:
    res = await client.get("/api/reports/catalog")
    assert res.status_code == 200
    ids = [item["id"] for item in res.json()["data"]["reports"]]
    for report_id in (
        "aged-payables",
        "budget-variance",
        "payment-schedule",
        "invoice-register",
        "vendor-spend-summary",
        "cash-forecast",
        "advance-reconciliation",
        "advance-aging",
        "expense-claims-register",
        "reimbursement-due",
        "missing-documents",
        "claim-status",
        "policy-exceptions",
        "invoice-exception",
        "process-efficiency",
        "control-centre",
    ):
        assert report_id in ids
    assert len(ids) == 16
    assert "aging-ap" not in ids
    control = next(
        item for item in res.json()["data"]["reports"] if item["id"] == "control-centre"
    )
    assert control["category"] == "exceptions_controls"

