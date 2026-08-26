"""Edge-case coverage for Invoice-to-Pay and Team Expense catalog reports."""

from __future__ import annotations

import calendar
import json
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from io import BytesIO
from uuid import UUID

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm.attributes import flag_modified

from app.models.audit import AuditLog
from app.models.employee_master import EmployeeMasterRecord
from app.models.invoice import Invoice, InvoiceStatus
from app.models.journal import EntryType, JournalEntry
from app.models.payment import Payment, PaymentStatus
from app.models.purchase_order import PurchaseOrder, PurchaseOrderStatus
from app.models.tenant import Tenant
from app.models.tenant_rule_book_config import TenantRuleBookConfig
from app.models.user_report_favourite import UserReportFavourite
from app.services.invoice.invoice_evaluation_service import (
    EVAL_PENDING_APPROVAL,
    ROUTE_PURCHASE,
    ROUTE_SALES,
    ROUTE_TEAM,
    ROUTE_VAULT,
)
from app.services.master_data.party_coa_subledger_service import party_sub_ledger_code
from app.services.purchase.team_expense_validator import vr_te03_receipt
from app.services.reports.payables_catalog_builders import (
    days_to_due,
    due_bucket,
    timing_label,
)
from app.services.reports.report_catalog import CATALOG_BY_ID
from app.services.reports.report_export_service import PDF_ROWS_PER_PAGE
from app.services.reports.team_expense_catalog_builders import fifo_age_outstanding
from app.schemas.rule_book_config import TeamExpensePolicy, TeamExpenseRule
from app.tenant_ids import TESTING_TENANT_UUID

NEW_REPORT_IDS = (
    "aged-payables",
    "payment-schedule",
    "invoice-register",
    "vendor-spend-summary",
    "cash-forecast",
    "advance-reconciliation",
    "advance-aging",
    "expense-claims-register",
    "reimbursement-due",
    "missing-documents",
    "claim-status",
    "policy-exceptions",
    "invoice-exception",
    "control-centre",
)
KPI_REPORT_IDS = ("process-efficiency",)


def _today() -> date:
    return date.today()


def _month_end(day: date | None = None) -> date:
    day = day or _today()
    return date(day.year, day.month, calendar.monthrange(day.year, day.month)[1])


def _joined(payload: dict) -> str:
    return " ".join(cell for row in payload["rows"] for cell in row["cells"])


def _row_for(payload: dict, token: str) -> list[str]:
    return next(row["cells"] for row in payload["rows"] if token in row["cells"])


async def _invoice(db: AsyncSession, **kwargs) -> Invoice:
    payload = {
        "tenant_id": TESTING_TENANT_UUID,
        "vendor": "Acme",
        "invoice_no": "INV-EDGE-1",
        "invoice_date": _today(),
        "due_date": _today() + timedelta(days=14),
        "subtotal": Decimal("100.00"),
        "gst": Decimal("10.00"),
        "total": Decimal("110.00"),
        "currency": "AUD",
        "status": InvoiceStatus.PROCESSED,
        "file_hash": kwargs.pop("file_hash", f"edge-{id(kwargs)}"),
    }
    payload.update(kwargs)
    if "file_hash" not in payload or not payload["file_hash"]:
        payload["file_hash"] = f"edge-{payload['invoice_no']}"
    inv = Invoice(**payload)
    db.add(inv)
    await db.flush()
    return inv


async def _journal(
    db: AsyncSession,
    invoice: Invoice,
    *,
    account_code: str = "2000",
    account_name: str = "Accounts Payable",
    debit: Decimal = Decimal("0"),
    credit: Decimal = Decimal("110.00"),
    on: date | None = None,
    tenant_id=TESTING_TENANT_UUID,
) -> JournalEntry:
    entry = JournalEntry(
        tenant_id=tenant_id,
        invoice_id=invoice.id,
        date=on or invoice.invoice_date or _today(),
        account_code=account_code,
        account_name=account_name,
        debit=debit,
        credit=credit,
        entry_type=EntryType.DEBIT if debit else EntryType.CREDIT,
    )
    db.add(entry)
    await db.flush()
    return entry


async def _payment(db: AsyncSession, invoice: Invoice, **kwargs) -> Payment:
    payload = {
        "tenant_id": invoice.tenant_id,
        "invoice_id": invoice.id,
        "vendor": invoice.vendor,
        "amount": invoice.total or Decimal("0"),
        "currency": invoice.currency or "AUD",
        "status": PaymentStatus.QUEUE,
    }
    payload.update(kwargs)
    row = Payment(**payload)
    db.add(row)
    await db.flush()
    return row


async def _add_te_rules(db: AsyncSession, *rules: dict) -> None:
    row = await db.get(TenantRuleBookConfig, TESTING_TENANT_UUID)
    assert row is not None
    config = dict(row.config)
    existing = list(config.get("team_expense_rules") or [])
    config["team_expense_rules"] = list(rules) + existing
    row.config = config
    flag_modified(row, "config")
    await db.flush()


def test_due_bucket_seven_day_boundary() -> None:
    as_of = date(2026, 8, 31)
    assert due_bucket(as_of, as_of) == "Due within 7 days"
    assert due_bucket(as_of + timedelta(days=7), as_of) == "Due within 7 days"
    assert due_bucket(as_of + timedelta(days=8), as_of) == "Upcoming"
    assert due_bucket(as_of - timedelta(days=1), as_of) == "Overdue"
    assert days_to_due(date(2026, 7, 15), date(2026, 8, 25)) == -41
    assert timing_label(date(2026, 7, 15), date(2026, 8, 25)) == "OVERDUE"
    assert timing_label(as_of + timedelta(days=7), as_of) == "Due within 7 days"


def test_fifo_age_outstanding_zero_exact_credit_and_boundaries() -> None:
    as_of = date(2026, 8, 24)
    assert fifo_age_outstanding([(as_of, Decimal("50"), Decimal("50"))], as_of) is None
    leftover = fifo_age_outstanding(
        [
            (as_of - timedelta(days=100), Decimal("40"), Decimal("0")),
            (as_of - timedelta(days=10), Decimal("60"), Decimal("0")),
            (as_of - timedelta(days=5), Decimal("0"), Decimal("40")),
        ],
        as_of,
    )
    assert leftover is not None
    assert leftover["90+"] == Decimal("0")
    assert leftover["0–30"] == Decimal("60.00") or leftover["0–30"] == Decimal("60")
    at_30 = fifo_age_outstanding([(as_of - timedelta(days=30), Decimal("10"), Decimal("0"))], as_of)
    at_60 = fifo_age_outstanding([(as_of - timedelta(days=60), Decimal("10"), Decimal("0"))], as_of)
    at_90 = fifo_age_outstanding([(as_of - timedelta(days=90), Decimal("10"), Decimal("0"))], as_of)
    assert at_30 is not None and at_30["0–30"] > 0
    assert at_60 is not None and at_60["31–60"] > 0
    assert at_90 is not None and at_90["61–90"] > 0


def test_vr_te03_passes_when_receipt_file_present() -> None:
    rule = TeamExpenseRule(
        id="te-req",
        name="Needs receipt",
        post_to={"ledger": "Travel"},
        policy=TeamExpensePolicy(require_receipt=True, receipt_threshold=0),
    )
    result = vr_te03_receipt(
        rule, 80.0, has_receipt_file=True, team_expense_kind="expense_claim"
    )
    assert result.passed is True


@pytest.mark.asyncio
async def test_new_reports_empty_state_on_fresh_tenant(client: AsyncClient) -> None:
    for report_id in NEW_REPORT_IDS:
        res = await client.get(f"/api/reports/{report_id}/preview?range=month")
        assert res.status_code == 200, report_id
        data = res.json()["data"]
        assert data["empty"] is True, report_id
        assert data["rows"] == [], report_id
        assert data["columns"], report_id


@pytest.mark.asyncio
async def test_process_efficiency_empty_tenant_still_shows_kpi_rows(
    client: AsyncClient,
) -> None:
    res = await client.get("/api/reports/process-efficiency/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["empty"] is False
    assert data["columns"] == [
        "Metric",
        "Current Month",
        "Previous Month",
        "Target",
        "Trend",
        "Definition",
    ]
    metrics = [row["cells"][0] for row in data["rows"]]
    assert metrics == [
        "Invoices processed",
        "Avg receipt-to-approval days",
        "Straight-through processing %",
        "Manual intervention %",
        "First-pass validation %",
        "Cost per invoice",
    ]
    by_metric = {row["cells"][0]: row["cells"] for row in data["rows"]}
    assert by_metric["Invoices processed"][1:5] == ["0", "0", "-", "Unchanged"]
    assert by_metric["Cost per invoice"][1:5] == ["-", "-", "-", "-"]
    assert all(row["cells"][3] == "-" for row in data["rows"])
    notes = data.get("notes") or ""
    assert "proxy" in notes.lower()
    assert "no data source" in notes.lower()


@pytest.mark.asyncio
async def test_custom_range_from_after_to_is_400_not_500(client: AsyncClient) -> None:
    preview = await client.get(
        "/api/reports/invoice-register/preview",
        params={"range": "custom", "from": "2026-08-20", "to": "2026-08-01"},
    )
    assert preview.status_code == 400
    export = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "xlsx", "range": "custom", "from": "2026-08-20", "to": "2026-08-01"},
    )
    assert export.status_code == 400
    te_xlsx = await client.post(
        "/api/reports/expense-claims-register/export",
        json={"format": "xlsx", "range": "custom", "from": "2026-08-20", "to": "2026-08-01"},
    )
    assert te_xlsx.status_code == 400


@pytest.mark.asyncio
async def test_custom_range_with_zero_matching_rows(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _invoice(
        db_session,
        file_hash="zero-range",
        invoice_no="IN-2026",
        invoice_date=date(2026, 8, 10),
    )
    await db_session.commit()
    res = await client.get(
        "/api/reports/invoice-register/preview",
        params={"range": "custom", "from": "2020-01-01", "to": "2020-01-31"},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["empty"] is True
    assert data["rows"] == []
    assert "IN-2026" not in _joined(data)


@pytest.mark.asyncio
async def test_aged_payables_boundary_null_due_and_same_vendor(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    on_as_of = await _invoice(
        db_session,
        file_hash="aged-on-asof",
        vendor="Twin Co",
        invoice_no="AGED-ON-ASOF",
        due_date=as_of,
        total=Decimal("10.00"),
    )
    await _journal(db_session, on_as_of, credit=Decimal("10.00"))
    null_due = await _invoice(
        db_session,
        file_hash="aged-null-due",
        vendor="Null Due Co",
        invoice_no="AGED-NULL-DUE",
        due_date=None,
        total=Decimal("20.00"),
    )
    await _journal(db_session, null_due, credit=Decimal("20.00"))
    sibling = await _invoice(
        db_session,
        file_hash="aged-twin-2",
        vendor="Twin Co",
        invoice_no="AGED-TWIN-2",
        due_date=as_of - timedelta(days=10),
        total=Decimal("30.00"),
    )
    await _journal(db_session, sibling, credit=Decimal("30.00"))
    await db_session.commit()

    res = await client.get("/api/reports/aged-payables/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    current_idx = data["columns"].index("Current")
    on_row = _row_for(data, "AGED-ON-ASOF")
    null_row = _row_for(data, "AGED-NULL-DUE")
    assert on_row[current_idx] == "10.00"
    assert null_row[current_idx] == "20.00"
    docs = [row["cells"][1] for row in data["rows"] if row["cells"][0] == "Twin Co"]
    assert "AGED-ON-ASOF" in docs
    assert "AGED-TWIN-2" in docs
    assert data["columns"][:8] == [
        "Vendor",
        "Invoice No",
        "Invoice Date",
        "Due Date",
        "Invoice Amount",
        "Amount Paid",
        "Balance Due",
        "Days Overdue",
    ]


@pytest.mark.asyncio
async def test_aged_payables_pack_layout_partial_payment_and_days_overdue(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    inv = await _invoice(
        db_session,
        file_hash="aged-pack-1042",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-1042",
        invoice_date=as_of - timedelta(days=71),
        due_date=as_of - timedelta(days=41),
        total=Decimal("125000.00"),
    )
    await _journal(db_session, inv, credit=Decimal("125000.00"))
    await _journal(
        db_session,
        inv,
        debit=Decimal("25000.00"),
        credit=Decimal("0"),
    )
    await db_session.commit()
    assert inv.id

    res = await client.get("/api/reports/aged-payables/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    row = _row_for(data, "INV-1042")
    cols = data["columns"]
    assert row[cols.index("Invoice Amount")] == "125,000.00"
    assert row[cols.index("Amount Paid")] == "25,000.00"
    assert row[cols.index("Balance Due")] == "100,000.00"
    assert row[cols.index("Days Overdue")] == "41"
    assert row[cols.index("31–60")] == "100,000.00"
    assert row[cols.index("Current")] == "-"
    assert row[cols.index("1–30")] == "-"
    assert row[cols.index("61–90")] == "-"
    assert row[cols.index("90+")] == "-"


@pytest.mark.asyncio
async def test_payment_schedule_seven_day_unpaid_and_past_scheduled(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    exact_7 = await _invoice(
        db_session,
        file_hash="sched-7",
        vendor="Sched Edge",
        invoice_no="SCH-EXACT-7",
        due_date=as_of + timedelta(days=7),
    )
    unpaid = await _invoice(
        db_session,
        file_hash="sched-unpaid",
        vendor="Sched Edge",
        invoice_no="SCH-UNPAID",
        due_date=as_of + timedelta(days=3),
    )
    await _payment(db_session, unpaid, status=PaymentStatus.QUEUE)
    overdue = await _invoice(
        db_session,
        file_hash="sched-past",
        vendor="Sched Edge",
        invoice_no="SCH-PAST-SCHED",
        due_date=as_of - timedelta(days=4),
    )
    await _payment(
        db_session,
        overdue,
        status=PaymentStatus.SCHEDULED,
        scheduled_date=as_of - timedelta(days=2),
    )
    for inv, credit in (
        (exact_7, Decimal("110.00")),
        (unpaid, Decimal("110.00")),
        (overdue, Decimal("110.00")),
    ):
        await _journal(db_session, inv, credit=credit)
    await db_session.commit()
    assert exact_7.id

    res = await client.get("/api/reports/payment-schedule/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    timing_idx = data["columns"].index("Timing")
    sched_idx = data["columns"].index("Scheduled Pay Date")
    days_idx = data["columns"].index("Days to Due")
    assert _row_for(data, "SCH-EXACT-7")[timing_idx] == "Due within 7 days"
    assert _row_for(data, "SCH-EXACT-7")[days_idx] == "7"
    unpaid_row = _row_for(data, "SCH-UNPAID")
    assert unpaid_row[timing_idx] == "Due within 7 days"
    past_row = _row_for(data, "SCH-PAST-SCHED")
    assert past_row[timing_idx] == "OVERDUE"
    assert past_row[days_idx] == "-4"
    assert past_row[sched_idx] == (as_of - timedelta(days=2)).isoformat()


@pytest.mark.asyncio
async def test_payment_schedule_pack_layout_overdue_days_and_totals(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    inv = await _invoice(
        db_session,
        file_hash="sched-pack-1042",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-1042",
        invoice_date=as_of - timedelta(days=71),
        due_date=as_of - timedelta(days=41),
        total=Decimal("125000.00"),
        approval_chain={"approvals": [{"name": "AP Manager", "at": "2026-07-10T10:00:00Z"}]},
    )
    await _journal(db_session, inv, credit=Decimal("125000.00"))
    await _journal(
        db_session,
        inv,
        debit=Decimal("25000.00"),
        credit=Decimal("0"),
    )
    await _payment(
        db_session,
        inv,
        status=PaymentStatus.SCHEDULED,
        scheduled_date=as_of - timedelta(days=42),
        amount=Decimal("100000.00"),
    )
    await db_session.commit()

    res = await client.get("/api/reports/payment-schedule/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Vendor",
        "Invoice No",
        "Due Date",
        "Amount Due",
        "Approval Status",
        "Scheduled Pay Date",
        "Days to Due",
        "Timing",
        "Notes",
    ]
    row = _row_for(data, "INV-1042")
    cols = data["columns"]
    assert row[cols.index("Vendor")] == "Acme Supplies Pvt Ltd"
    assert row[cols.index("Due Date")] == (as_of - timedelta(days=41)).isoformat()
    assert row[cols.index("Amount Due")] == "100,000.00"
    assert row[cols.index("Approval Status")] == "Approved"
    assert row[cols.index("Scheduled Pay Date")] == (as_of - timedelta(days=42)).isoformat()
    assert row[cols.index("Days to Due")] == "-41"
    assert row[cols.index("Timing")] == "OVERDUE"
    assert row[cols.index("Notes")] == ""
    notes = data.get("notes") or ""
    assert "Overdue total: 100,000.00" in notes
    assert "Due within 7 days: -" in notes


@pytest.mark.asyncio
async def test_invoice_register_last_approval_empty_chain_and_multi_payment(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    last = await _invoice(
        db_session,
        file_hash="reg-last-appr",
        vendor="Approval Co",
        invoice_no="REG-LAST",
        approval_chain={
            "approvals": [
                {"name": "First Approver", "at": "2026-01-01T10:00:00Z"},
                {"name": "Last Approver", "at": "2026-01-08T10:00:00Z"},
            ]
        },
    )
    empty = await _invoice(
        db_session,
        file_hash="reg-empty-appr",
        vendor="Approval Co",
        invoice_no="REG-EMPTY",
        approval_chain={"approvals": []},
    )
    multi = await _invoice(
        db_session,
        file_hash="reg-multi-pay",
        vendor="Approval Co",
        invoice_no="REG-MULTI-PAY",
    )
    await _payment(
        db_session,
        multi,
        status=PaymentStatus.QUEUE,
        payment_intent="QUEUE-REF",
    )
    await _payment(
        db_session,
        multi,
        status=PaymentStatus.PAID,
        payment_intent="PAID-REF",
        paid_date=datetime(2026, 8, 12, tzinfo=timezone.utc),
    )
    await db_session.commit()
    assert last.id and empty.id

    res = await client.get("/api/reports/invoice-register/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    matches = [row for row in data["rows"] if "REG-MULTI-PAY" in row["cells"]]
    assert len(matches) == 1
    multi_row = matches[0]["cells"]
    assert "Last Approver" in _row_for(data, "REG-LAST")
    assert "First Approver" not in _row_for(data, "REG-LAST")
    empty_row = _row_for(data, "REG-EMPTY")
    approver_idx = data["columns"].index("Approver")
    assert empty_row[approver_idx] == ""
    assert "PAID-REF" in multi_row
    assert "QUEUE-REF" not in multi_row
    pay_date_idx = data["columns"].index("Payment Date")
    pay_ref_idx = data["columns"].index("Payment Ref")
    assert multi_row[pay_date_idx] == "2026-08-12"
    assert multi_row[pay_ref_idx] == "PAID-REF"


@pytest.mark.asyncio
async def test_invoice_register_pack_layout_tax_approval_and_blank_payment(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    inv = await _invoice(
        db_session,
        file_hash="reg-pack-1042",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-1042",
        invoice_date=date(2026, 6, 15),
        due_date=date(2026, 7, 15),
        subtotal=Decimal("125000.00"),
        gst=Decimal("22500.00"),
        total=Decimal("147500.00"),
        approval_chain={
            "approvals": [{"name": "A. Rao", "at": "2026-06-18T10:00:00Z"}]
        },
    )
    await db_session.commit()
    assert inv.id

    res = await client.get(
        "/api/reports/invoice-register/preview",
        params={"range": "custom", "from": "2026-06-01", "to": "2026-07-31"},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Vendor",
        "Invoice No",
        "Invoice Date",
        "Due Date",
        "Amount (excl. tax)",
        "Tax / GST",
        "Total",
        "Status",
        "Approver",
        "Approval Date",
        "Payment Date",
        "Payment Ref",
    ]
    row = _row_for(data, "INV-1042")
    cols = data["columns"]
    assert row[cols.index("Vendor")] == "Acme Supplies Pvt Ltd"
    assert row[cols.index("Invoice Date")] == "2026-06-15"
    assert row[cols.index("Due Date")] == "2026-07-15"
    assert row[cols.index("Amount (excl. tax)")] == "125,000.00"
    assert row[cols.index("Tax / GST")] == "22,500.00"
    assert row[cols.index("Total")] == "147,500.00"
    assert row[cols.index("Status")] == "Approved"
    assert row[cols.index("Approver")] == "A. Rao"
    assert row[cols.index("Approval Date")] == "2026-06-18"
    assert row[cols.index("Payment Date")] == ""
    assert row[cols.index("Payment Ref")] == ""


@pytest.mark.asyncio
async def test_invoice_register_month_includes_prior_dated_invoices(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    prior = await _invoice(
        db_session,
        file_hash="reg-prior-date",
        vendor="Prior Register Co",
        invoice_no="REG-PRIOR",
        invoice_date=as_of - timedelta(days=50),
        due_date=as_of - timedelta(days=20),
    )
    future = await _invoice(
        db_session,
        file_hash="reg-future-date",
        vendor="Future Register Co",
        invoice_no="REG-FUTURE",
        invoice_date=as_of + timedelta(days=3),
        due_date=as_of + timedelta(days=17),
    )
    await db_session.commit()
    assert prior.id and future.id

    month = await client.get("/api/reports/invoice-register/preview?range=month")
    assert month.status_code == 200
    month_blob = _joined(month.json()["data"])
    assert "REG-PRIOR" in month_blob
    assert "REG-FUTURE" not in month_blob
    assert month.json()["data"]["period_label"].startswith("As of ")

    custom = await client.get(
        "/api/reports/invoice-register/preview",
        params={
            "range": "custom",
            "from": as_of.replace(day=1).isoformat(),
            "to": as_of.isoformat(),
        },
    )
    assert custom.status_code == 200
    custom_blob = _joined(custom.json()["data"])
    assert "REG-PRIOR" not in custom_blob
    assert "REG-FUTURE" not in custom_blob


@pytest.mark.asyncio
async def test_vendor_spend_groups_whitespace_and_case(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _invoice(
        db_session,
        file_hash="vend-a",
        vendor="  Acme Co  ",
        invoice_no="VS-1",
        total=Decimal("10.00"),
    )
    await _invoice(
        db_session,
        file_hash="vend-b",
        vendor="acme co",
        invoice_no="VS-2",
        total=Decimal("15.00"),
    )
    await db_session.commit()
    res = await client.get("/api/reports/vendor-spend-summary/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    named = [row for row in data["rows"] if "acme" in row["cells"][0].casefold()]
    assert len(named) == 1
    assert named[0]["cells"][1] == "2"
    assert "25.00" in named[0]["cells"][2]


@pytest.mark.asyncio
async def test_vendor_spend_pack_layout_unpaid_dash_and_total(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    inv = await _invoice(
        db_session,
        file_hash="vs-pack-1042",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-1042",
        invoice_date=date(2026, 6, 15),
        due_date=date(2026, 7, 15),
        subtotal=Decimal("125000.00"),
        gst=Decimal("22500.00"),
        total=Decimal("147500.00"),
    )
    await db_session.commit()
    assert inv.id

    res = await client.get("/api/reports/vendor-spend-summary/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Vendor",
        "# Invoices",
        "Total Invoiced",
        "Total Paid",
        "Outstanding",
    ]
    assert "Invoice Register" in (data.get("notes") or "")
    row = _row_for(data, "Acme Supplies Pvt Ltd")
    cols = data["columns"]
    assert row[cols.index("# Invoices")] == "1"
    assert row[cols.index("Total Invoiced")] == "147,500.00"
    assert row[cols.index("Total Paid")] == "-"
    assert row[cols.index("Outstanding")] == "147,500.00"
    total = next(item for item in data["rows"] if item["cells"][0] == "TOTAL")
    assert total["emphasize"] is True
    assert total["cells"][cols.index("Total Invoiced")] == "147,500.00"
    assert total["cells"][cols.index("Total Paid")] == "-"
    assert total["cells"][cols.index("Outstanding")] == "147,500.00"


@pytest.mark.asyncio
async def test_cash_forecast_empty_overdue_scheduled_and_cumulative(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    empty = await client.get("/api/reports/cash-forecast/preview?range=month")
    assert empty.status_code == 200
    assert empty.json()["data"]["empty"] is True

    as_of = _month_end()
    overdue = await _invoice(
        db_session,
        file_hash="fc-overdue",
        vendor="Forecast Co",
        invoice_no="FC-OVERDUE",
        due_date=as_of - timedelta(days=9),
        total=Decimal("40.00"),
    )
    await _payment(
        db_session,
        overdue,
        status=PaymentStatus.SCHEDULED,
        scheduled_date=as_of + timedelta(days=20),
    )
    upcoming = await _invoice(
        db_session,
        file_hash="fc-up",
        vendor="Forecast Co",
        invoice_no="FC-UP",
        due_date=as_of + timedelta(days=12),
        total=Decimal("10.00"),
    )
    await _journal(db_session, overdue, credit=Decimal("40.00"))
    await _journal(db_session, upcoming, credit=Decimal("10.00"))
    await db_session.commit()
    assert upcoming.id

    res = await client.get("/api/reports/cash-forecast/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    date_idx = data["columns"].index("Date")
    amt_idx = data["columns"].index("Amount due")
    cum_idx = data["columns"].index("Cumulative")
    dates = [row["cells"][date_idx] for row in data["rows"]]
    assert as_of.isoformat() in dates
    assert (as_of + timedelta(days=20)).isoformat() not in dates
    overdue_row = next(row["cells"] for row in data["rows"] if row["cells"][date_idx] == as_of.isoformat())
    assert overdue_row[amt_idx] == "40.00"
    parsed = [
        (row["cells"][date_idx], Decimal(row["cells"][amt_idx].replace(",", "")), Decimal(row["cells"][cum_idx].replace(",", "")))
        for row in data["rows"]
    ]
    running = Decimal("0")
    prev = ""
    for day, amount, cumulative in parsed:
        assert day >= prev
        running += amount
        assert cumulative == running
        prev = day


@pytest.mark.asyncio
async def test_budget_variance_zero_unmatched_and_negative_remaining(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.services.purchase.team_expense_spend_service import current_period_keys

    today = _today()
    keys = current_period_keys(today)
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="ZeroGL",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("0.00"),
            enforcement="soft",
        )
    )
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="Travel",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("100.00"),
            enforcement="soft",
        )
    )
    await _invoice(
        db_session,
        file_hash="bv-act",
        vendor="Staff",
        invoice_no="BV-ACT-2",
        total=Decimal("80.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        employee_email="traveller@example.com",
    )
    await _invoice(
        db_session,
        file_hash="bv-com",
        vendor="Staff",
        invoice_no="BV-COM-2",
        total=Decimal("40.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        evaluation_status="",
        employee_email="traveller@example.com",
    )
    await _invoice(
        db_session,
        file_hash="bv-orphan",
        vendor="Staff",
        invoice_no="BV-ORPHAN",
        total=Decimal("99.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="NoSuchGL",
        evaluation_status="",
        employee_email="traveller@example.com",
    )
    await db_session.commit()

    res = await client.get("/api/reports/budget-variance/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert "NoSuchGL" not in blob
    cols = data["columns"]
    cat_idx = cols.index("Category")
    zero = next(row["cells"] for row in data["rows"] if row["cells"][cat_idx] == "ZeroGL")
    assert zero[cols.index("Budget")] == "0.00"
    assert zero[cols.index("Variance %")] == "-"
    assert zero[cols.index("% Utilise")] == "-"
    travel = next(row["cells"] for row in data["rows"] if row["cells"][cat_idx] == "Travel")
    assert travel[cols.index("Variance")] == "20.00"
    assert travel[cols.index("% Utilise")] == "120.0%"


@pytest.mark.asyncio
async def test_budget_variance_pack_layout_variance_and_utilise(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.services.purchase.team_expense_spend_service import current_period_keys

    today = _today()
    keys = current_period_keys(today)
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            department="Sales — West",
            gl_ledger="Travel",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("500000.00"),
            enforcement="soft",
        )
    )
    await _invoice(
        db_session,
        file_hash="bv-pack-actual",
        vendor="Staff",
        invoice_no="BV-PACK-ACT",
        total=Decimal("380000.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        employee_email="traveller@example.com",
    )
    await _invoice(
        db_session,
        file_hash="bv-pack-commit",
        vendor="Staff",
        invoice_no="BV-PACK-COM",
        total=Decimal("50000.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="Travel",
        evaluation_status="",
        employee_email="traveller@example.com",
    )
    await db_session.commit()

    res = await client.get("/api/reports/budget-variance/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"][:8] == [
        "Department / Project",
        "Category",
        "Budget",
        "Committed",
        "Actual",
        "Variance",
        "Variance %",
        "% Utilise",
    ]
    cols = data["columns"]
    row = next(
        item["cells"]
        for item in data["rows"]
        if item["cells"][cols.index("Category")] == "Travel"
    )
    assert row[cols.index("Department / Project")] == "Sales — West"
    assert row[cols.index("Budget")] == "500,000.00"
    assert row[cols.index("Committed")] == "50,000.00"
    assert row[cols.index("Actual")] == "380,000.00"
    assert row[cols.index("Variance")] == "120,000.00"
    assert row[cols.index("Variance %")] == "24.0%"
    assert row[cols.index("% Utilise")] == "86.0%"
    total = next(item for item in data["rows"] if item["cells"][0] == "TOTAL")
    assert total["emphasize"] is True
    assert total["cells"][cols.index("Budget")] == "500,000.00"
    assert total["cells"][cols.index("Variance")] == "120,000.00"
    assert total["cells"][cols.index("% Utilise")] == "86.0%"


@pytest.mark.asyncio
async def test_budget_variance_month_includes_prior_period_budgets(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.services.purchase.team_expense_spend_service import current_period_keys

    as_of = _month_end()
    prior_day = date(as_of.year, as_of.month, 1) - timedelta(days=1)
    future_day = as_of + timedelta(days=1)
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="PriorTravel",
            period_kind="monthly",
            period_key=current_period_keys(prior_day)["monthly"],
            allocated=Decimal("10.00"),
            enforcement="soft",
        )
    )
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="FutureTravel",
            period_kind="monthly",
            period_key=current_period_keys(future_day)["monthly"],
            allocated=Decimal("20.00"),
            enforcement="soft",
        )
    )
    await db_session.commit()

    month = await client.get("/api/reports/budget-variance/preview?range=month")
    assert month.status_code == 200
    month_blob = _joined(month.json()["data"])
    assert "PriorTravel" in month_blob
    assert "FutureTravel" not in month_blob

    custom = await client.get(
        "/api/reports/budget-variance/preview",
        params={
            "range": "custom",
            "from": date(as_of.year, as_of.month, 1).isoformat(),
            "to": as_of.isoformat(),
        },
    )
    assert custom.status_code == 200
    custom_blob = _joined(custom.json()["data"])
    assert "PriorTravel" not in custom_blob
    assert "FutureTravel" not in custom_blob
    assert "Team Expenses" in (custom.json()["data"].get("notes") or "")


@pytest.mark.asyncio
async def test_ap_reports_exclude_sales_and_vault(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    keep = await _invoice(
        db_session,
        file_hash="ap-keep-sv",
        vendor="Keep Vendor",
        invoice_no="AP-KEEP-SV",
        due_date=_today() - timedelta(days=2),
    )
    await _journal(db_session, keep)
    sales = await _invoice(
        db_session,
        file_hash="sales-excl",
        vendor="Sales Exclude",
        invoice_no="SALES-EXCL-1",
        due_date=_today() - timedelta(days=2),
        route_target=ROUTE_SALES,
    )
    await _journal(db_session, sales)
    vault = await _invoice(
        db_session,
        file_hash="vault-excl",
        vendor="Vault Exclude",
        invoice_no="VAULT-EXCL-1",
        due_date=_today() - timedelta(days=2),
        route_target=ROUTE_VAULT,
    )
    await _journal(db_session, vault)
    await db_session.commit()

    for report_id in (
        "aged-payables",
        "payment-schedule",
        "invoice-register",
        "vendor-spend-summary",
        "cash-forecast",
    ):
        res = await client.get(f"/api/reports/{report_id}/preview?range=month")
        assert res.status_code == 200, report_id
        blob = _joined(res.json()["data"])
        assert "SALES-EXCL-1" not in blob, report_id
        assert "VAULT-EXCL-1" not in blob, report_id
        assert "Sales Exclude" not in blob, report_id
        assert "Vault Exclude" not in blob, report_id
        if report_id == "cash-forecast":
            amounts = [row["cells"][2] for row in res.json()["data"]["rows"]]
            assert "110.00" in amounts, report_id
            assert "330.00" not in amounts, report_id
        else:
            assert "AP-KEEP-SV" in blob or "Keep Vendor" in blob, report_id


@pytest.mark.asyncio
async def test_advance_aging_zero_fifo_and_boundaries(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    db_session.add_all(
        [
            EmployeeMasterRecord(
                tenant_id=TESTING_TENANT_UUID,
                master_id="emzero",
                name="Zero Outstanding",
                email="zero@example.com",
                advance_parent_ledger="Staff Advance",
                status="Active",
            ),
            EmployeeMasterRecord(
                tenant_id=TESTING_TENANT_UUID,
                master_id="emfifo",
                name="Fifo Emp",
                email="fifo@example.com",
                advance_parent_ledger="Staff Advance",
                status="Active",
            ),
            EmployeeMasterRecord(
                tenant_id=TESTING_TENANT_UUID,
                master_id="embound",
                name="Boundary Emp",
                email="bound@example.com",
                advance_parent_ledger="Staff Advance",
                status="Active",
            ),
        ]
    )
    await db_session.flush()
    zero_inv = await _invoice(
        db_session,
        file_hash="adv-zero",
        vendor="Zero Outstanding",
        invoice_no="ADV-ZERO",
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="zero@example.com",
        invoice_date=as_of - timedelta(days=20),
    )
    fifo_inv = await _invoice(
        db_session,
        file_hash="adv-fifo",
        vendor="Fifo Emp",
        invoice_no="ADV-FIFO",
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="fifo@example.com",
        invoice_date=as_of - timedelta(days=100),
    )
    bound_inv = await _invoice(
        db_session,
        file_hash="adv-bound",
        vendor="Boundary Emp",
        invoice_no="ADV-BOUND",
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="bound@example.com",
        invoice_date=as_of - timedelta(days=90),
    )
    zero_code = party_sub_ledger_code("emzero")
    fifo_code = party_sub_ledger_code("emfifo")
    bound_code = party_sub_ledger_code("embound")
    await _journal(
        db_session,
        zero_inv,
        account_code=zero_code,
        account_name="Staff Advance",
        debit=Decimal("50.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=20),
    )
    await _journal(
        db_session,
        zero_inv,
        account_code=zero_code,
        account_name="Staff Advance",
        debit=Decimal("0"),
        credit=Decimal("50.00"),
        on=as_of - timedelta(days=5),
    )
    await _journal(
        db_session,
        fifo_inv,
        account_code=fifo_code,
        account_name="Staff Advance",
        debit=Decimal("40.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=100),
    )
    await _journal(
        db_session,
        fifo_inv,
        account_code=fifo_code,
        account_name="Staff Advance",
        debit=Decimal("60.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=10),
    )
    await _journal(
        db_session,
        fifo_inv,
        account_code=fifo_code,
        account_name="Staff Advance",
        debit=Decimal("0"),
        credit=Decimal("40.00"),
        on=as_of - timedelta(days=5),
    )
    await _journal(
        db_session,
        bound_inv,
        account_code=bound_code,
        account_name="Staff Advance",
        debit=Decimal("12.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=90),
    )
    await db_session.commit()

    res = await client.get("/api/reports/advance-aging/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    cols = data["columns"]
    zero_row = _row_for(data, "ADV-ZERO")
    assert zero_row[cols.index("Outstanding")] == "-"
    assert zero_row[cols.index("Advance Issued")] == "50.00"
    assert zero_row[cols.index("Amount Settled")] == "50.00"
    fifo_row = _row_for(data, "ADV-FIFO")
    assert fifo_row[cols.index("90+")] == "60.00"
    assert fifo_row[cols.index("0–30")] == "-"
    assert fifo_row[cols.index("Outstanding")] == "60.00"
    assert fifo_row[cols.index("Advance Issued")] == "100.00"
    bound_row = _row_for(data, "ADV-BOUND")
    assert bound_row[cols.index("61–90")] == "12.00"
    assert bound_row[cols.index("Days Outstanding")] == "90"


@pytest.mark.asyncio
async def test_advance_aging_pack_layout_settled_days_and_total(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="empriya",
            name="Priya Sharma",
            email="priya@example.com",
            advance_parent_ledger="Staff Advance",
            status="Active",
        )
    )
    await db_session.flush()
    inv = await _invoice(
        db_session,
        file_hash="adv-pack-207",
        vendor="Priya Sharma",
        invoice_no="ADV-207",
        invoice_date=as_of - timedelta(days=85),
        total=Decimal("60000.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="priya@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    code = party_sub_ledger_code("empriya")
    await _journal(
        db_session,
        inv,
        account_code=code,
        account_name="Staff Advance",
        debit=Decimal("60000.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=85),
    )
    await _journal(
        db_session,
        inv,
        account_code=code,
        account_name="Staff Advance",
        debit=Decimal("0"),
        credit=Decimal("60000.00"),
        on=as_of - timedelta(days=10),
    )
    await db_session.commit()
    assert inv.id

    res = await client.get("/api/reports/advance-aging/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Employee",
        "Advance Ref",
        "Date Issued",
        "Advance Issued",
        "Amount Settled",
        "Outstanding",
        "Days Outstanding",
        "0–30",
        "31–60",
        "61–90",
        "90+",
    ]
    row = _row_for(data, "ADV-207")
    cols = data["columns"]
    assert row[cols.index("Employee")] == "Priya Sharma"
    assert row[cols.index("Advance Issued")] == "60,000.00"
    assert row[cols.index("Amount Settled")] == "60,000.00"
    assert row[cols.index("Outstanding")] == "-"
    assert row[cols.index("Days Outstanding")] == "85"
    assert row[cols.index("0–30")] == "-"
    assert row[cols.index("61–90")] == "-"
    assert row[cols.index("90+")] == "-"
    total = next(item for item in data["rows"] if item["cells"][0] == "TOTAL")
    assert total["emphasize"] is True
    assert total["cells"][cols.index("Advance Issued")] == "60,000.00"
    assert total["cells"][cols.index("Outstanding")] == "-"


@pytest.mark.asyncio
async def test_expense_claims_register_kind_filter(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emkinds",
            name="Kind Emp",
            email="kinds@example.com",
            status="Active",
        )
    )
    await _invoice(
        db_session,
        file_hash="kind-against",
        vendor="Kind Emp",
        invoice_no="KIND-AGAINST",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_against_advance",
        employee_email="kinds@example.com",
    )
    await _invoice(
        db_session,
        file_hash="kind-adv",
        vendor="Kind Emp",
        invoice_no="KIND-ADV",
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="kinds@example.com",
    )
    await _invoice(
        db_session,
        file_hash="kind-direct",
        vendor="Kind Emp",
        invoice_no="KIND-DIRECT",
        route_target=ROUTE_TEAM,
        team_expense_kind="direct_payment",
        employee_email="kinds@example.com",
    )
    await db_session.commit()
    res = await client.get("/api/reports/expense-claims-register/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert "KIND-AGAINST" in blob
    assert "KIND-ADV" not in blob
    assert "KIND-DIRECT" not in blob
    row = _row_for(data, "KIND-AGAINST")
    cols = data["columns"]
    assert row[cols.index("Against Advance?")] == "Yes"


@pytest.mark.asyncio
async def test_expense_claims_register_pack_layout_and_as_of_month(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="empriyaclm",
            name="Priya Sharma",
            email="priya.claim@example.com",
            department="Sales — West",
            status="Active",
        )
    )
    await _invoice(
        db_session,
        file_hash="clm-3301",
        vendor="Priya Sharma",
        invoice_no="CLM-3301",
        invoice_date=date(2026, 6, 5),
        total=Decimal("18000.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_against_advance",
        employee_email="priya.claim@example.com",
        account_name="Accommodation",
        raw_file_path="/tmp/receipt.pdf",
        status=InvoiceStatus.PROCESSED,
    )
    await db_session.commit()

    res = await client.get("/api/reports/expense-claims-register/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Claim ID",
        "Employee",
        "Date",
        "Dept / Project",
        "Category",
        "Amount",
        "Against Advance?",
        "Receipt Attached?",
        "Within Policy",
        "Status",
        "Approver",
    ]
    assert "Every claim line" in (data.get("notes") or "")
    row = _row_for(data, "CLM-3301")
    cols = data["columns"]
    assert row[cols.index("Claim ID")] == "CLM-3301"
    assert row[cols.index("Employee")] == "Priya Sharma"
    assert row[cols.index("Date")] == "2026-06-05"
    assert row[cols.index("Dept / Project")] == "Sales — West"
    assert row[cols.index("Category")] == "Accommodation"
    assert row[cols.index("Amount")] == "18,000.00"
    assert row[cols.index("Against Advance?")] == "Yes"
    assert row[cols.index("Receipt Attached?")] == "Yes"
    assert row[cols.index("Within Policy")] == "Yes"
    assert row[cols.index("Status")] == "Approved"
    assert row[cols.index("Approver")] == ""
    assert all(item["cells"][0] != "TOTAL" for item in data["rows"])


@pytest.mark.asyncio
async def test_reimbursement_due_excludes_pending_rejected_duplicate_processed(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emreimb",
            name="Reimb Emp",
            email="reimb@example.com",
            status="Active",
        )
    )
    keep = await _invoice(
        db_session,
        file_hash="reimb-keep",
        vendor="Reimb Emp",
        invoice_no="REIMB-KEEP",
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        evaluation_status="",
        employee_email="reimb@example.com",
        total=Decimal("33.00"),
    )
    await _invoice(
        db_session,
        file_hash="reimb-pend",
        vendor="Reimb Emp",
        invoice_no="REIMB-PEND",
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        evaluation_status=EVAL_PENDING_APPROVAL,
        employee_email="reimb@example.com",
    )
    await _invoice(
        db_session,
        file_hash="reimb-rej",
        vendor="Reimb Emp",
        invoice_no="REIMB-REJ",
        status=InvoiceStatus.REJECTED,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="reimb@example.com",
    )
    await _invoice(
        db_session,
        file_hash="reimb-dup",
        vendor="Reimb Emp",
        invoice_no="REIMB-DUP",
        status=InvoiceStatus.DUPLICATE_SKIPPED,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="reimb@example.com",
    )
    await _invoice(
        db_session,
        file_hash="reimb-proc",
        vendor="Reimb Emp",
        invoice_no="REIMB-PROC",
        status=InvoiceStatus.PROCESSED,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="reimb@example.com",
    )
    await db_session.commit()
    assert keep.id
    res = await client.get("/api/reports/reimbursement-due/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert "Reimb Emp" in blob
    assert "REIMB-PEND" not in blob
    assert "REIMB-REJ" not in blob
    assert "REIMB-DUP" not in blob
    assert "REIMB-PROC" not in blob
    cols = data["columns"]
    row = _row_for(data, "Reimb Emp")
    assert row[cols.index("Out-of-pocket Claims")] == "33.00"
    assert row[cols.index("From Advances (Co. owes)")] == "-"
    assert row[cols.index("Total Due")] == "33.00"
    assert "out-of-pocket" in (data.get("notes") or "").lower()


@pytest.mark.asyncio
async def test_reimbursement_due_pack_layout_advance_net_and_total(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="empriyareimb",
            name="Priya Sharma",
            email="priya.reimb@example.com",
            advance_parent_ledger="Staff Advance",
            status="Active",
        )
    )
    await db_session.flush()
    adv = await _invoice(
        db_session,
        file_hash="reimb-pack-adv",
        vendor="Priya Sharma",
        invoice_no="ADV-207",
        total=Decimal("60000.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="priya.reimb@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    code = party_sub_ledger_code("empriyareimb")
    await _journal(
        db_session,
        adv,
        account_code=code,
        account_name="Staff Advance",
        debit=Decimal("60000.00"),
        credit=Decimal("0"),
    )
    await _invoice(
        db_session,
        file_hash="reimb-pack-clm",
        vendor="Priya Sharma",
        invoice_no="CLM-3301",
        total=Decimal("72000.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_against_advance",
        evaluation_status="",
        employee_email="priya.reimb@example.com",
    )
    await db_session.commit()

    res = await client.get("/api/reports/reimbursement-due/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Employee",
        "From Advances (Co. owes)",
        "Out-of-pocket Claims",
        "Total Due",
        "Payment Status",
    ]
    row = _row_for(data, "Priya Sharma")
    cols = data["columns"]
    assert row[cols.index("From Advances (Co. owes)")] == "12,000.00"
    assert row[cols.index("Out-of-pocket Claims")] == "-"
    assert row[cols.index("Total Due")] == "12,000.00"
    assert row[cols.index("Payment Status")] == ""
    total = next(item for item in data["rows"] if item["cells"][0] == "TOTAL")
    assert total["emphasize"] is True
    assert total["cells"][cols.index("From Advances (Co. owes)")] == "12,000.00"
    assert total["cells"][cols.index("Total Due")] == "12,000.00"


@pytest.mark.asyncio
async def test_missing_documents_receipt_policy_cases(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _add_te_rules(
        db_session,
        {
            "id": "tr-optional",
            "name": "Optional receipt",
            "enabled": True,
            "priority": 1,
            "match_on": {"merchant_contains": "OptionalReceipt"},
            "post_to": {"ledger": "Travel Expense"},
            "policy": {
                "require_receipt": False,
                "receipt_threshold": 0,
                "auto_approve_below": 0,
            },
        },
        {
            "id": "tr-required",
            "name": "Required receipt",
            "enabled": True,
            "priority": 2,
            "match_on": {"merchant_contains": "NeedsReceipt"},
            "post_to": {"ledger": "Travel Expense"},
            "policy": {
                "require_receipt": True,
                "receipt_threshold": 0,
                "auto_approve_below": 0,
            },
        },
    )
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emdocs",
            name="Docs Emp",
            email="docs@example.com",
            status="Active",
        )
    )
    await _invoice(
        db_session,
        file_hash="md-optional",
        vendor="OptionalReceipt Cafe",
        invoice_no="MD-OPTIONAL",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="docs@example.com",
        raw_file_path=None,
        total=Decimal("70.00"),
    )
    await _invoice(
        db_session,
        file_hash="md-both",
        vendor="NeedsReceipt Cafe",
        invoice_no="MD-BOTH",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="docs@example.com",
        raw_file_path=None,
        total=Decimal("70.00"),
    )
    await _invoice(
        db_session,
        file_hash="md-hasfile",
        vendor="NeedsReceipt Cafe",
        invoice_no="MD-HASFILE",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="docs@example.com",
        raw_file_path="/tmp/receipt.pdf",
        total=Decimal("70.00"),
    )
    await db_session.commit()

    res = await client.get("/api/reports/missing-documents/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert "MD-OPTIONAL" not in blob
    assert "MD-HASFILE" not in blob
    both = [row for row in data["rows"] if "MD-BOTH" in row["cells"]]
    assert len(both) == 1
    cols = data["columns"]
    row = both[0]["cells"]
    assert row[cols.index("Type")] == "Expense claim"
    assert row[cols.index("Owner")] == "Docs Emp"
    assert row[cols.index("Expected Document")] == "Receipt"
    assert row[cols.index("Attached")] == "N"
    assert row[cols.index("Flag")] == "MISSING"
    notes = row[cols.index("Notes")]
    assert "No stored file" in notes
    assert "VR-TE03" in notes


@pytest.mark.asyncio
async def test_missing_documents_pack_layout_invoice_without_file(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _invoice(
        db_session,
        file_hash="md-pack-1042",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-1042",
        raw_file_path=None,
    )
    await _invoice(
        db_session,
        file_hash="md-pack-has",
        vendor="Acme Supplies Pvt Ltd",
        invoice_no="INV-HASFILE",
        raw_file_path="/tmp/tax-invoice.pdf",
    )
    await _invoice(
        db_session,
        file_hash="md-pack-sales",
        vendor="Sales Co",
        invoice_no="SO-SKIP",
        route_target=ROUTE_SALES,
        raw_file_path=None,
    )
    await db_session.commit()

    res = await client.get("/api/reports/missing-documents/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Type",
        "Reference ID",
        "Owner",
        "Expected Document",
        "Attached",
        "Flag",
        "Notes",
    ]
    assert "audit-ready" in (data.get("notes") or "")
    blob = _joined(data)
    assert "INV-HASFILE" not in blob
    assert "SO-SKIP" not in blob
    row = _row_for(data, "INV-1042")
    cols = data["columns"]
    assert row[cols.index("Type")] == "Invoice"
    assert row[cols.index("Owner")] == "Acme Supplies Pvt Ltd"
    assert row[cols.index("Expected Document")] == "Tax invoice PDF"
    assert row[cols.index("Attached")] == "N"
    assert row[cols.index("Flag")] == "MISSING"
    assert row[cols.index("Notes")] == "Vendor to resend"


@pytest.mark.asyncio
async def test_new_reports_ignore_other_tenant(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    other_id = UUID("aaaaaaaa-bbbb-cccc-dddd-eeeeeeeeeee1")
    db_session.add(
        Tenant(
            id=other_id,
            name="Other Co",
            slug="other-edge",
            currency="AUD",
            settings_json={"country": "AU"},
        )
    )
    await db_session.flush()
    foreign_ap = Invoice(
        tenant_id=other_id,
        vendor="FOREIGN-ISO-VENDOR",
        invoice_no="FOREIGN-ISO-1",
        invoice_date=_today(),
        due_date=_today() - timedelta(days=3),
        total=Decimal("999.00"),
        currency="AUD",
        status=InvoiceStatus.PROCESSED,
        file_hash="foreign-iso-ap",
        duplicate_review_suggested=True,
        validation_results=json.dumps(
            [{"rule": "VR02", "passed": False, "skipped": False, "message": "dup"}]
        ),
    )
    db_session.add(foreign_ap)
    await db_session.flush()
    await _journal(
        db_session,
        foreign_ap,
        credit=Decimal("999.00"),
        tenant_id=other_id,
    )
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=other_id,
            master_id="emforeign",
            name="FOREIGN-ISO-EMP",
            email="foreign@example.com",
            advance_parent_ledger="Staff Advance",
            status="Active",
        )
    )
    foreign_te = Invoice(
        tenant_id=other_id,
        vendor="FOREIGN-ISO-EMP",
        invoice_no="FOREIGN-TE-1",
        invoice_date=_today(),
        total=Decimal("50.00"),
        currency="AUD",
        status=InvoiceStatus.VALIDATING,
        file_hash="foreign-iso-te",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="foreign@example.com",
        evaluation_status=EVAL_PENDING_APPROVAL,
        validation_results=json.dumps(
            [
                {
                    "rule": "VR-TE04",
                    "passed": False,
                    "skipped": False,
                    "message": "Missing bank",
                    "severity": "block",
                }
            ]
        ),
    )
    db_session.add(foreign_te)
    await db_session.commit()

    for report_id in (*NEW_REPORT_IDS, *KPI_REPORT_IDS):
        res = await client.get(f"/api/reports/{report_id}/preview?range=month")
        assert res.status_code == 200, report_id
        blob = _joined(res.json()["data"])
        assert "FOREIGN-ISO-1" not in blob, report_id
        assert "FOREIGN-ISO-VENDOR" not in blob, report_id
        assert "FOREIGN-TE-1" not in blob, report_id
        assert "FOREIGN-ISO-EMP" not in blob, report_id
    efficiency = await client.get("/api/reports/process-efficiency/preview?range=month")
    assert efficiency.status_code == 200
    processed_row = next(
        row["cells"]
        for row in efficiency.json()["data"]["rows"]
        if row["cells"][0] == "Invoices processed"
    )
    assert processed_row[1] == "0"


@pytest.mark.asyncio
async def test_invoice_register_pdf_paginates_past_28_rows(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    extra = PDF_ROWS_PER_PAGE + 8
    for i in range(extra):
        await _invoice(
            db_session,
            file_hash=f"pdf-reg-{i}",
            vendor=f"Pager {i:03d}",
            invoice_no=f"PDF-REG-{i:03d}",
        )
    await db_session.commit()
    res = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "pdf", "range": "month"},
    )
    assert res.status_code == 200
    assert int(res.headers["x-data-rows"]) >= extra
    assert int(res.headers["x-page-count"]) >= 2
    import fitz

    doc = fitz.open(stream=res.content, filetype="pdf")
    try:
        assert doc.page_count >= 2
        text = "".join(page.get_text() for page in doc)
        assert "Pager 000" in text
        assert "Pager 035" in text
    finally:
        doc.close()


@pytest.mark.asyncio
async def test_unicode_vendor_and_employee_names_export(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _invoice(
        db_session,
        file_hash="uni-ap",
        vendor="Café José 李",
        invoice_no="UNI-AP-1",
    )
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emuni",
            name="Renée Müller",
            email="renee@example.com",
            status="Active",
        )
    )
    await _invoice(
        db_session,
        file_hash="uni-te",
        vendor="Renée Müller",
        invoice_no="UNI-TE-1",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="renee@example.com",
    )
    await db_session.commit()

    xlsx = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "xlsx", "range": "month"},
    )
    assert xlsx.status_code == 200
    wb = load_workbook(BytesIO(xlsx.content))
    values = [
        str(cell) if cell is not None else ""
        for ws in wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for cell in row
    ]
    assert any("Café José 李" in value for value in values)

    pdf = await client.post(
        "/api/reports/invoice-register/export",
        json={"format": "pdf", "range": "month"},
    )
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"
    import fitz

    doc = fitz.open(stream=pdf.content, filetype="pdf")
    try:
        text = "".join(page.get_text() for page in doc)
        assert "Café" in text or "Jose" in text or "李" in text or "UNI-AP-1" in text
        assert "UNI-AP-1" in text
    finally:
        doc.close()

    te_xlsx = await client.post(
        "/api/reports/expense-claims-register/export",
        json={"format": "xlsx", "range": "month"},
    )
    assert te_xlsx.status_code == 200
    te_wb = load_workbook(BytesIO(te_xlsx.content))
    te_values = [
        str(cell) if cell is not None else ""
        for ws in te_wb.worksheets
        for row in ws.iter_rows(values_only=True)
        for cell in row
    ]
    assert any("Renée" in value or "Müller" in value or "UNI-TE-1" in value for value in te_values)


@pytest.mark.asyncio
async def test_favourites_rapid_replace_has_no_duplicates(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from sqlalchemy import select

    first = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["cash-forecast"]},
    )
    assert first.status_code == 200
    second = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["cash-forecast"]},
    )
    assert second.status_code == 200
    third = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["invoice-register", "cash-forecast", "invoice-register"]},
    )
    assert third.status_code == 200
    assert third.json()["data"] == ["invoice-register", "cash-forecast"]
    fourth = await client.put(
        "/api/reports/favourites",
        json={"report_ids": ["vendor-spend-summary"]},
    )
    assert fourth.status_code == 200
    catalog = await client.get("/api/reports/catalog")
    ids = catalog.json()["data"]["favourite_ids"]
    assert ids == ["vendor-spend-summary"]
    rows = (
        await db_session.execute(select(UserReportFavourite.report_id))
    ).scalars().all()
    assert rows == ["vendor-spend-summary"]
    assert set(ids).issubset(CATALOG_BY_ID)


def _validation(*rows: dict) -> str:
    return json.dumps(list(rows))


@pytest.mark.asyncio
async def test_claim_status_age_only_for_pending_approval(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    start, end = date(2026, 8, 1), date(2026, 8, 31)
    pending = await _invoice(
        db_session,
        file_hash="cs-pending",
        invoice_no="CS-PENDING",
        invoice_date=date(2026, 8, 10),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        status=InvoiceStatus.VALIDATING,
        evaluation_status=EVAL_PENDING_APPROVAL,
        approval_chain={
            "module": "team_expenses",
            "mode": "two_way",
            "required": 2,
            "approvals": [{"name": "Pat Signer", "user_id": 11}],
        },
    )
    posted = await _invoice(
        db_session,
        file_hash="cs-posted",
        invoice_no="CS-POSTED",
        invoice_date=date(2026, 8, 11),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        status=InvoiceStatus.PROCESSED,
        evaluation_status="posted",
    )
    no_hold = await _invoice(
        db_session,
        file_hash="cs-no-hold",
        invoice_no="CS-NO-HOLD",
        invoice_date=date(2026, 8, 12),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        status=InvoiceStatus.VALIDATING,
        evaluation_status=EVAL_PENDING_APPROVAL,
    )
    db_session.add(
        AuditLog(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=pending.id,
            event="team_expense_approval_required",
            created_at=datetime(2026, 8, 24, tzinfo=timezone.utc),
        )
    )
    db_session.add(
        AuditLog(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=posted.id,
            event="team_expense_approval_required",
            created_at=datetime(2026, 8, 20, tzinfo=timezone.utc),
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/claim-status/preview",
        params={"range": "custom", "from": start.isoformat(), "to": end.isoformat()},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    notes = data.get("notes") or ""
    assert "approval-hold timestamp" in notes
    cols = data["columns"]
    assert cols[-1] == "Status / Reason"
    age_idx = cols.index("Age in Stage")
    approver_idx = cols.index("Current Approver")
    stage_idx = cols.index("Current Stage")
    pending_row = _row_for(data, "CS-PENDING")
    posted_row = _row_for(data, "CS-POSTED")
    no_hold_row = _row_for(data, "CS-NO-HOLD")
    assert pending_row[age_idx] == str((end - date(2026, 8, 24)).days)
    assert posted_row[age_idx] == "0"
    assert no_hold_row[age_idx] == str((end - date(2026, 8, 12)).days)
    assert pending_row[stage_idx] == "Finance Review"
    assert posted_row[stage_idx] == "Approved"
    assert posted_row[approver_idx] == "—"
    assert "Pat Signer" in pending_row[approver_idx]
    assert "awaiting pool (1 remaining)" in pending_row[approver_idx]


@pytest.mark.asyncio
async def test_claim_status_pack_layout_stages_and_reimbursement(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emkumar",
            name="A. Kumar",
            email="kumar@example.com",
            department="Sales",
            status="Active",
        )
    )
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="emsingh",
            name="M. Singh",
            email="singh@example.com",
            department="Operations",
            status="Active",
        )
    )
    db_session.add(
        EmployeeMasterRecord(
            tenant_id=TESTING_TENANT_UUID,
            master_id="empatel",
            name="R. Patel",
            email="patel@example.com",
            department="Technology",
            status="Active",
        )
    )
    pending = await _invoice(
        db_session,
        file_hash="clm-8801",
        vendor="A. Kumar",
        invoice_no="CLM-8801",
        invoice_date=date(2026, 8, 18),
        total=Decimal("1240.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="kumar@example.com",
        status=InvoiceStatus.VALIDATING,
        evaluation_status=EVAL_PENDING_APPROVAL,
    )
    approved = await _invoice(
        db_session,
        file_hash="clm-8803",
        vendor="M. Singh",
        invoice_no="CLM-8803",
        invoice_date=date(2026, 8, 12),
        total=Decimal("860.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="singh@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    await _payment(
        db_session,
        approved,
        status=PaymentStatus.SCHEDULED,
        scheduled_date=date(2026, 8, 23),
    )
    await _invoice(
        db_session,
        file_hash="clm-8804",
        vendor="R. Patel",
        invoice_no="CLM-8804",
        invoice_date=date(2026, 8, 11),
        total=Decimal("460.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        employee_email="patel@example.com",
        status=InvoiceStatus.REJECTED,
        validation_results=_validation(
            {
                "rule": "VR-TE03",
                "passed": False,
                "skipped": False,
                "message": "Receipt required",
                "severity": "block",
            }
        ),
    )
    await _invoice(
        db_session,
        file_hash="clm-adv-skip",
        vendor="A. Kumar",
        invoice_no="ADV-SKIP",
        invoice_date=date(2026, 8, 10),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="kumar@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    db_session.add(
        AuditLog(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=pending.id,
            event="team_expense_approval_required",
            created_at=datetime(2026, 8, 18, tzinfo=timezone.utc),
        )
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/claim-status/preview",
        params={"range": "custom", "from": "2026-08-01", "to": "2026-08-21"},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Claim ID",
        "Employee",
        "Department",
        "Submitted Date",
        "Amount",
        "Current Stage",
        "Current Approver",
        "Age in Stage",
        "Reimbursement Date",
        "Status / Reason",
    ]
    blob = _joined(data)
    assert "ADV-SKIP" not in blob
    cols = data["columns"]
    one = _row_for(data, "CLM-8801")
    assert one[cols.index("Employee")] == "A. Kumar"
    assert one[cols.index("Department")] == "Sales"
    assert one[cols.index("Submitted Date")] == "2026-08-18"
    assert one[cols.index("Amount")] == "1,240.00"
    assert one[cols.index("Current Stage")] == "Manager Approval"
    assert one[cols.index("Age in Stage")] == "3"
    assert one[cols.index("Status / Reason")] == "Pending"
    three = _row_for(data, "CLM-8803")
    assert three[cols.index("Current Stage")] == "Approved"
    assert three[cols.index("Current Approver")] == "—"
    assert three[cols.index("Age in Stage")] == "0"
    assert three[cols.index("Reimbursement Date")] == "2026-08-23"
    assert three[cols.index("Status / Reason")] == "Scheduled"
    four = _row_for(data, "CLM-8804")
    assert four[cols.index("Current Stage")] == "Rejected"
    assert four[cols.index("Age in Stage")] == "0"
    assert four[cols.index("Status / Reason")] == "Missing receipt"


@pytest.mark.asyncio
async def test_policy_exceptions_exclude_vr_te03(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    await _invoice(
        db_session,
        file_hash="pe-te03-only",
        invoice_no="PE-TE03-ONLY",
        vendor="Cafe TE03",
        invoice_date=_today(),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        raw_file_path=None,
        validation_results=_validation(
            {
                "rule": "VR-TE03",
                "passed": False,
                "skipped": False,
                "message": "Receipt required",
                "severity": "block",
            }
        ),
    )
    await _invoice(
        db_session,
        file_hash="pe-te04",
        invoice_no="PE-TE04",
        vendor="Cafe TE04",
        invoice_date=_today(),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        validation_results=_validation(
            {
                "rule": "VR-TE04",
                "passed": False,
                "skipped": False,
                "message": "Employee bank missing",
                "severity": "block",
            }
        ),
    )
    await _invoice(
        db_session,
        file_hash="pe-both",
        invoice_no="PE-BOTH",
        vendor="Cafe Both",
        invoice_date=_today(),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        raw_file_path=None,
        validation_results=_validation(
            {
                "rule": "VR-TE03",
                "passed": False,
                "skipped": False,
                "message": "Receipt required",
                "severity": "block",
            },
            {
                "rule": "VR-TE08",
                "passed": False,
                "skipped": False,
                "message": "Budget overrun",
                "severity": "block",
            },
        ),
    )
    await db_session.commit()

    res = await client.get("/api/reports/policy-exceptions/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert "PE-TE03-ONLY" not in blob
    assert "VR-TE03" not in blob
    assert "PE-TE04" in blob
    assert "VR-TE04" in blob
    both_rows = [row["cells"] for row in data["rows"] if "PE-BOTH" in row["cells"]]
    assert len(both_rows) == 1
    assert "VR-TE08" in both_rows[0]
    assert "VR-TE03" not in both_rows[0]
    assert "VR-TE03" in (data.get("notes") or "")


@pytest.mark.asyncio
async def test_invoice_exception_po_variance_requires_purchase_linked_po(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    dup = await _invoice(
        db_session,
        file_hash="ie-vr02",
        invoice_no="IE-VR02",
        vendor="Dup Vendor",
        invoice_date=_today(),
        status=InvoiceStatus.DUPLICATE_SKIPPED,
        validation_results=_validation(
            {"rule": "VR02", "passed": False, "skipped": False, "message": "Duplicate"}
        ),
    )
    db_session.add(
        AuditLog(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=dup.id,
            event="duplicate_skipped",
            detail={"actor_name": "Reviewer One"},
        )
    )
    po_var = await _invoice(
        db_session,
        file_hash="ie-po-var",
        invoice_no="IE-PO-VAR",
        vendor="PO Vendor",
        invoice_date=_today(),
        route_target=ROUTE_PURCHASE,
        status=InvoiceStatus.PROCESSED,
    )
    db_session.add(
        PurchaseOrder(
            tenant_id=TESTING_TENANT_UUID,
            po_number="PO-VAR-1",
            vendor="PO Vendor",
            invoice_id=po_var.id,
            po_qty=Decimal("1"),
            po_unit_price=Decimal("10"),
            status=PurchaseOrderStatus.VARIANCE_PENDING,
            three_way_match_status="mismatch",
        )
    )
    non_po = await _invoice(
        db_session,
        file_hash="ie-non-po",
        invoice_no="IE-NON-PO",
        vendor="Non PO Vendor",
        invoice_date=_today(),
        route_target=ROUTE_PURCHASE,
        status=InvoiceStatus.PROCESSED,
    )
    sales_with_po = await _invoice(
        db_session,
        file_hash="ie-sales-po",
        invoice_no="IE-SALES-PO",
        vendor="Sales Vendor",
        invoice_date=_today(),
        route_target=ROUTE_SALES,
        status=InvoiceStatus.PROCESSED,
    )
    db_session.add(
        PurchaseOrder(
            tenant_id=TESTING_TENANT_UUID,
            po_number="PO-SALES-1",
            vendor="Sales Vendor",
            invoice_id=sales_with_po.id,
            po_qty=Decimal("1"),
            po_unit_price=Decimal("10"),
            status=PurchaseOrderStatus.VARIANCE_PENDING,
            three_way_match_status="variance_pending",
        )
    )
    open_po = await _invoice(
        db_session,
        file_hash="ie-open-po",
        invoice_no="IE-OPEN-PO",
        vendor="Open PO Vendor",
        invoice_date=_today(),
        route_target=ROUTE_PURCHASE,
        status=InvoiceStatus.PROCESSED,
    )
    db_session.add(
        PurchaseOrder(
            tenant_id=TESTING_TENANT_UUID,
            po_number="PO-OPEN-1",
            vendor="Open PO Vendor",
            invoice_id=open_po.id,
            po_qty=Decimal("1"),
            po_unit_price=Decimal("10"),
            status=PurchaseOrderStatus.OPEN,
        )
    )
    await db_session.commit()
    assert non_po.id

    res = await client.get("/api/reports/invoice-exception/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    blob = _joined(data)
    assert data["columns"] == [
        "Risk",
        "Exception Type",
        "Invoice ID",
        "Vendor",
        "Amount",
        "Detected On",
        "Reason",
        "Owner",
        "Status",
        "Action",
    ]
    assert "IE-VR02" in blob
    assert "Potential Duplicate" in blob
    assert "Same amount/vendor and similar invoice number" in blob
    assert "Compare source document and hash" in blob
    assert "Under review" in blob
    assert "Reviewer One" in blob
    assert "IE-PO-VAR" in blob
    assert "PO Mismatch" in blob
    assert "Invoice exceeds PO by 100.00" in blob
    assert "Approve variance or request credit" in blob
    assert "Duplicate invoice (VR02)" not in blob
    assert "PO amount variance" not in blob
    assert "IE-NON-PO" not in blob
    assert "IE-SALES-PO" not in blob
    assert "IE-OPEN-PO" not in blob
    assert "N/A" not in blob
    assert "bank-change" not in blob.lower()
    assert "Bank change" not in blob
    assert "Bank Detail Change" not in blob


@pytest.mark.asyncio
async def test_invoice_exception_pack_layout(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    dup = await _invoice(
        db_session,
        file_hash="ie-pack-dup",
        invoice_no="INV-10482",
        vendor="Metro Media",
        invoice_date=date(2026, 8, 18),
        total=Decimal("8960.00"),
        status=InvoiceStatus.DUPLICATE_SKIPPED,
        validation_results=_validation(
            {"rule": "VR02", "passed": False, "skipped": False, "message": "Duplicate"}
        ),
    )
    po_inv = await _invoice(
        db_session,
        file_hash="ie-pack-po",
        invoice_no="INV-10484",
        vendor="OfficeHub",
        invoice_date=date(2026, 8, 21),
        total=Decimal("2360.00"),
        route_target=ROUTE_PURCHASE,
        status=InvoiceStatus.PROCESSED,
    )
    db_session.add(
        PurchaseOrder(
            tenant_id=TESTING_TENANT_UUID,
            po_number="PO-10484",
            vendor="OfficeHub",
            invoice_id=po_inv.id,
            po_qty=Decimal("1"),
            po_unit_price=Decimal("2000"),
            status=PurchaseOrderStatus.VARIANCE_PENDING,
            three_way_match_status="mismatch",
        )
    )
    bank = await _invoice(
        db_session,
        file_hash="ie-pack-bank",
        invoice_no="INV-10479",
        vendor="NorthStar Services",
        invoice_date=date(2026, 8, 22),
        total=Decimal("24500.00"),
        email_attachment_name="bank_details_change.pdf",
        status=InvoiceStatus.PROCESSED,
    )
    await _payment(
        db_session,
        bank,
        status=PaymentStatus.SCHEDULED,
        scheduled_date=date(2026, 8, 23),
    )
    await _invoice(
        db_session,
        file_hash="ie-pack-te-bank",
        invoice_no="INV-TE-BANK",
        vendor="Team Bank Doc",
        invoice_date=date(2026, 8, 22),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        email_attachment_name="bank_details_change.pdf",
    )
    await db_session.commit()
    assert dup.id and po_inv.id and bank.id

    res = await client.get(
        "/api/reports/invoice-exception/preview",
        params={"range": "custom", "from": "2026-08-01", "to": "2026-08-31"},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    by_id = {row["cells"][2]: row["cells"] for row in data["rows"]}
    assert by_id["INV-10482"] == [
        "High",
        "Potential Duplicate",
        "INV-10482",
        "Metro Media",
        "8,960.00",
        "2026-08-18",
        "Same amount/vendor and similar invoice number",
        "",
        "Under review",
        "Compare source document and hash",
    ]
    assert by_id["INV-10484"] == [
        "Medium",
        "PO Mismatch",
        "INV-10484",
        "OfficeHub",
        "2,360.00",
        "2026-08-21",
        "Invoice exceeds PO by 360.00",
        "",
        "Open",
        "Approve variance or request credit",
    ]
    assert by_id["INV-10479"] == [
        "Critical",
        "Bank Detail Change",
        "INV-10479",
        "NorthStar Services",
        "24,500.00",
        "2026-08-22",
        "Bank account changed 1 day before scheduled payment",
        "",
        "Blocked",
        "Independent bank verification",
    ]
    assert "INV-TE-BANK" not in by_id
    assert [row["cells"][2] for row in data["rows"]] == [
        "INV-10479",
        "INV-10482",
        "INV-10484",
    ]


@pytest.mark.asyncio
async def test_process_efficiency_computes_trend_without_target(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    current_created = datetime(2026, 8, 10, 9, 0, tzinfo=timezone.utc)
    prior_created = datetime(2026, 7, 10, 9, 0, tzinfo=timezone.utc)
    stp = await _invoice(
        db_session,
        file_hash="peff-stp",
        invoice_no="PEFF-STP",
        invoice_date=date(2026, 8, 10),
        status=InvoiceStatus.PROCESSED,
        created_at=current_created,
    )
    edited = await _invoice(
        db_session,
        file_hash="peff-edit",
        invoice_no="PEFF-EDIT",
        invoice_date=date(2026, 8, 10),
        status=InvoiceStatus.PROCESSED,
        created_at=current_created,
    )
    prior = await _invoice(
        db_session,
        file_hash="peff-prior",
        invoice_no="PEFF-PRIOR",
        invoice_date=date(2026, 7, 10),
        status=InvoiceStatus.PROCESSED,
        created_at=prior_created,
    )
    stp.created_at = current_created
    edited.created_at = current_created
    prior.created_at = prior_created
    db_session.add_all(
        [
            AuditLog(
                tenant_id=TESTING_TENANT_UUID,
                invoice_id=stp.id,
                event="team_expense_auto_approved",
                created_at=datetime(2026, 8, 12, tzinfo=timezone.utc),
            ),
            AuditLog(
                tenant_id=TESTING_TENANT_UUID,
                invoice_id=edited.id,
                event="invoice_fields_updated",
                created_at=datetime(2026, 8, 11, tzinfo=timezone.utc),
            ),
            AuditLog(
                tenant_id=TESTING_TENANT_UUID,
                invoice_id=edited.id,
                event="invoice_approved",
                created_at=datetime(2026, 8, 15, tzinfo=timezone.utc),
            ),
            AuditLog(
                tenant_id=TESTING_TENANT_UUID,
                invoice_id=prior.id,
                event="invoice_approved",
                created_at=datetime(2026, 7, 12, tzinfo=timezone.utc),
            ),
        ]
    )
    await db_session.commit()

    res = await client.get(
        "/api/reports/process-efficiency/preview",
        params={"range": "custom", "from": "2026-08-01", "to": "2026-08-31"},
    )
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["columns"] == [
        "Metric",
        "Current Month",
        "Previous Month",
        "Target",
        "Trend",
        "Definition",
    ]
    by_metric = {row["cells"][0]: row["cells"] for row in data["rows"]}
    assert by_metric["Invoices processed"] == [
        "Invoices processed",
        "2",
        "1",
        "-",
        "Up",
        "Count of invoices completed",
    ]
    assert by_metric["Avg receipt-to-approval days"][1:5] == ["3.5", "2.0", "-", "Worsening"]
    assert by_metric["Straight-through processing %"][1:5] == ["50.0%", "0.0%", "-", "Improving"]
    assert by_metric["Manual intervention %"][1:5] == ["50.0%", "100.0%", "-", "Improving"]
    assert by_metric["First-pass validation %"][1:5] == ["50.0%", "100.0%", "-", "Worsening"]
    assert by_metric["Cost per invoice"][1:5] == ["-", "-", "-", "-"]
    assert by_metric["Cost per invoice"][5] == "Estimated AP operating cost / invoices processed"
    notes = data.get("notes") or ""
    assert "proxy" in notes.lower()
    assert "no data source" in notes.lower()
    assert not any(row["cells"][0] == "First-Pass" for row in data["rows"])


@pytest.mark.asyncio
async def test_control_centre_unions_confirmed_sources_only(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.models.employee_master import EmployeeMasterRecord
    from app.services.master_data.party_coa_subledger_service import party_sub_ledger_code
    from app.services.purchase.team_expense_spend_service import current_period_keys

    as_of = _month_end()
    overdue = await _invoice(
        db_session,
        file_hash="cc-overdue-ap",
        vendor="CC Overdue Co",
        invoice_no="CC-OVERDUE-AP",
        due_date=as_of - timedelta(days=4),
        status=InvoiceStatus.PROCESSED,
    )
    upcoming = await _invoice(
        db_session,
        file_hash="cc-upcoming-ap",
        vendor="CC Upcoming Co",
        invoice_no="CC-UPCOMING-AP",
        due_date=as_of + timedelta(days=21),
        status=InvoiceStatus.PROCESSED,
    )
    await _journal(db_session, overdue, credit=Decimal("110.00"))
    await _journal(db_session, upcoming, credit=Decimal("110.00"))
    aged = EmployeeMasterRecord(
        tenant_id=TESTING_TENANT_UUID,
        master_id="emccadv",
        name="CC Advance Emp",
        email="ccadvance@example.com",
        advance_parent_ledger="Staff Advance",
        status="Active",
    )
    db_session.add(aged)
    await db_session.flush()
    adv_inv = await _invoice(
        db_session,
        file_hash="cc-adv",
        vendor="CC Advance Emp",
        invoice_no="CC-ADV",
        invoice_date=as_of - timedelta(days=74),
        route_target=ROUTE_TEAM,
        team_expense_kind="advance_requisition",
        employee_email="ccadvance@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    await _journal(
        db_session,
        adv_inv,
        account_code=party_sub_ledger_code("emccadv"),
        account_name="Staff Advance",
        debit=Decimal("80.00"),
        credit=Decimal("0"),
        on=as_of - timedelta(days=74),
    )
    keys = current_period_keys(_today())
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="CC Travel",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("10.00"),
            enforcement="soft",
        )
    )
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="CC Under",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("5000.00"),
            enforcement="soft",
        )
    )
    await _invoice(
        db_session,
        file_hash="cc-over-budget",
        vendor="CC Spend",
        invoice_no="CC-OVER-BUDGET",
        invoice_date=_today(),
        total=Decimal("80.00"),
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        account_name="CC Travel",
        employee_email="ccadvance@example.com",
        status=InvoiceStatus.PROCESSED,
    )
    await _invoice(
        db_session,
        file_hash="cc-te03",
        invoice_no="CC-TE03-ONLY",
        vendor="CC Cafe TE03",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        validation_results=_validation(
            {
                "rule": "VR-TE03",
                "passed": False,
                "skipped": False,
                "message": "Receipt required",
                "severity": "block",
            }
        ),
    )
    await _invoice(
        db_session,
        file_hash="cc-te04",
        invoice_no="CC-TE04",
        vendor="CC Cafe TE04",
        route_target=ROUTE_TEAM,
        team_expense_kind="expense_claim",
        validation_results=_validation(
            {
                "rule": "VR-TE04",
                "passed": False,
                "skipped": False,
                "message": "Employee bank missing",
                "severity": "block",
            }
        ),
    )
    await _invoice(
        db_session,
        file_hash="cc-vr02",
        invoice_no="CC-VR02",
        vendor="CC Dup Vendor",
        status=InvoiceStatus.DUPLICATE_SKIPPED,
        validation_results=_validation(
            {"rule": "VR02", "passed": False, "skipped": False, "message": "Duplicate"}
        ),
    )
    non_po = await _invoice(
        db_session,
        file_hash="cc-non-po",
        invoice_no="CC-NON-PO",
        vendor="CC Non PO",
        route_target=ROUTE_PURCHASE,
        status=InvoiceStatus.PROCESSED,
        due_date=as_of + timedelta(days=40),
    )
    await db_session.commit()
    assert overdue.id and upcoming.id and non_po.id

    res = await client.get("/api/reports/control-centre/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    assert data["report_id"] == "control-centre"
    assert data["columns"] == [
        "Priority",
        "Alert",
        "Module",
        "Amount",
        "Owner",
        "Age / Due",
        "Recommended Action",
        "Status",
    ]
    blob = _joined(data)
    notes = data.get("notes") or ""
    assert "CC-OVERDUE-AP" in blob
    assert "Vendor invoice overdue" in blob
    assert "Invoice-to-Pay" in blob
    assert "CC-UPCOMING-AP" not in blob
    assert "CC Advance Emp" in blob
    assert "Employee advance outstanding > 60 days" in blob
    assert "74 days" in blob
    assert "CC Travel forecast above budget" in blob
    assert "CC Under" not in blob
    assert "CC-TE04" in blob
    assert "VR-TE04" in blob
    assert "CC-TE03-ONLY" in blob
    assert "Expense receipt missing" in blob
    assert "CC-VR02" in blob
    assert "Potential duplicate invoice" in blob
    assert "CC-NON-PO" not in blob
    assert "Straight-through" not in blob
    assert "Process Efficiency" not in blob
    assert "Invoice Processing Efficiency" not in blob
    assert "bank-change" not in blob.lower()
    assert "Bank change" not in blob
    assert "Bank Detail Change" not in blob
    assert "insurance" in notes.lower()
    alerts = [row["cells"][1] for row in data["rows"]]
    overdue_idx = next(i for i, alert in enumerate(alerts) if "Vendor invoice overdue" in alert)
    budget_idx = next(i for i, alert in enumerate(alerts) if "forecast above budget" in alert)
    assert overdue_idx < budget_idx


@pytest.mark.asyncio
async def test_ap_remaining_nets_partial_paid_without_journal_and_negative(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    as_of = _month_end()
    partial = await _invoice(
        db_session,
        file_hash="ap-partial",
        vendor="Partial Co",
        invoice_no="AP-PARTIAL",
        due_date=as_of - timedelta(days=3),
        total=Decimal("110.00"),
    )
    await _journal(db_session, partial, credit=Decimal("110.00"))
    await _journal(
        db_session,
        partial,
        debit=Decimal("40.00"),
        credit=Decimal("0"),
    )

    paid_only = await _invoice(
        db_session,
        file_hash="ap-paid-flag",
        vendor="Paid Flag Co",
        invoice_no="AP-PAID-FLAG",
        due_date=as_of - timedelta(days=2),
        total=Decimal("110.00"),
    )
    await _journal(db_session, paid_only, credit=Decimal("110.00"))
    await _payment(db_session, paid_only, status=PaymentStatus.PAID)

    settled = await _invoice(
        db_session,
        file_hash="ap-settled",
        vendor="Settled Co",
        invoice_no="AP-SETTLED",
        due_date=as_of - timedelta(days=1),
        total=Decimal("110.00"),
    )
    await _journal(db_session, settled, credit=Decimal("110.00"))
    await _journal(
        db_session,
        settled,
        debit=Decimal("110.00"),
        credit=Decimal("0"),
    )

    overpaid = await _invoice(
        db_session,
        file_hash="ap-overpaid",
        vendor="Overpaid Co",
        invoice_no="AP-OVERPAID",
        due_date=as_of - timedelta(days=5),
        total=Decimal("110.00"),
    )
    await _journal(db_session, overpaid, credit=Decimal("110.00"))
    await _journal(
        db_session,
        overpaid,
        debit=Decimal("150.00"),
        credit=Decimal("0"),
    )
    await db_session.commit()

    aged = await client.get("/api/reports/aged-payables/preview?range=month")
    assert aged.status_code == 200
    aged_blob = _joined(aged.json()["data"])
    assert "AP-PARTIAL" in aged_blob
    assert "70.00" in aged_blob
    assert "AP-PAID-FLAG" in aged_blob
    assert "AP-SETTLED" not in aged_blob
    assert "AP-OVERPAID" in aged_blob
    assert "-40.00" in aged_blob

    sched = await client.get("/api/reports/payment-schedule/preview?range=month")
    assert sched.status_code == 200
    sched_data = sched.json()["data"]
    amt_idx = sched_data["columns"].index("Amount Due")
    assert _row_for(sched_data, "AP-PARTIAL")[amt_idx] == "70.00"
    assert _row_for(sched_data, "AP-PAID-FLAG")[amt_idx] == "110.00"
    with pytest.raises(StopIteration):
        _row_for(sched_data, "AP-SETTLED")
    with pytest.raises(StopIteration):
        _row_for(sched_data, "AP-OVERPAID")

    forecast = await client.get("/api/reports/cash-forecast/preview?range=month")
    assert forecast.status_code == 200
    fc = forecast.json()["data"]
    overdue_row = next(
        row["cells"] for row in fc["rows"] if row["cells"][0] == as_of.isoformat()
    )
    overdue_amt = Decimal(overdue_row[fc["columns"].index("Amount due")].replace(",", ""))
    assert overdue_amt == Decimal("180.00")

    control = await client.get("/api/reports/control-centre/preview?range=month")
    assert control.status_code == 200
    cc_blob = _joined(control.json()["data"])
    assert "AP-PARTIAL" in cc_blob
    assert "70.00" in cc_blob
    assert "AP-OVERPAID" not in cc_blob
    assert "AP-SETTLED" not in cc_blob


def test_paid_payment_amount_null_and_zero_are_not_invoice_total() -> None:
    from types import SimpleNamespace

    from app.models.payment import PaymentStatus
    from app.services.reports.payables_catalog_builders import _paid_payment_amount

    zero, missing_zero = _paid_payment_amount(
        SimpleNamespace(status=PaymentStatus.PAID, amount=Decimal("0.00"))
    )
    assert zero == Decimal("0.00")
    assert missing_zero is False

    null_amt, missing_null = _paid_payment_amount(
        SimpleNamespace(status=PaymentStatus.PAID, amount=None)
    )
    assert null_amt == Decimal("0.00")
    assert missing_null is True


@pytest.mark.asyncio
async def test_vendor_spend_paid_zero_amount_is_not_invoice_total(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    zero = await _invoice(
        db_session,
        file_hash="vs-zero-paid",
        vendor="Zero Paid Co",
        invoice_no="VS-ZERO",
        total=Decimal("80.00"),
    )
    await _payment(db_session, zero, status=PaymentStatus.PAID, amount=Decimal("0.00"))
    await db_session.commit()

    res = await client.get("/api/reports/vendor-spend-summary/preview?range=month")
    assert res.status_code == 200
    data = res.json()["data"]
    paid_idx = data["columns"].index("Total Paid")
    out_idx = data["columns"].index("Outstanding")
    zero_row = _row_for(data, "Zero Paid Co")
    assert zero_row[paid_idx] == "-"
    assert zero_row[out_idx] == "80.00"
    assert "80.00" not in zero_row[paid_idx]
    notes = data.get("notes") or ""
    assert "Invoice Register" in notes
    assert "no amount" not in notes.lower()


@pytest.mark.asyncio
async def test_ensure_payment_for_invoice_refuses_before_processed(
    db_session: AsyncSession,
) -> None:
    from app.services.payments.payment_service import ensure_payment_for_invoice

    inv = await _invoice(
        db_session,
        file_hash="pay-pre-processed",
        vendor="Guard Co",
        invoice_no="PAY-GUARD",
        due_date=_today() + timedelta(days=7),
        total=Decimal("110.00"),
        status=InvoiceStatus.VALIDATING,
        route_target=ROUTE_PURCHASE,
    )
    payment = await ensure_payment_for_invoice(db_session, inv)
    assert payment is None
    assert inv.status != InvoiceStatus.PROCESSED

