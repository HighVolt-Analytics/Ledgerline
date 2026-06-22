"""Employee bulk import from register / payment templates."""

import io

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.audit import AuditLog
from app.models.employee_master import EmployeeMasterRecord
from app.services.employee_import_service import (
    build_import_template,
    import_employee_masters,
    parse_employee_import_file,
)
from app.services.rule_book_mapper import clear_classification_config_cache


def _register_csv(rows: list[list[str]]) -> bytes:
    lines = ["name,email,whatsapp_number,role,status"]
    for row in rows:
        lines.append(",".join(row))
    return "\n".join(lines).encode("utf-8")


def _payment_xlsx(rows: list[list[str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(
        [
            "email",
            "bsb",
            "account_number",
            "account_name",
            "bank_name",
            "budget_monthly",
            "budget_quarterly",
            "budget_annual",
            "status",
        ]
    )
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.asyncio
async def test_parse_register_csv_aliases() -> None:
    data = (
        "Full Name,Work Email,Mobile,Department\n"
        "Sam Lee,sam.lee@example.com,+61411112222,Sales\n"
    ).encode("utf-8")
    rows = parse_employee_import_file(data, "staff.csv")
    assert len(rows) == 1
    assert rows[0]["Full Name"] == "Sam Lee"
    assert rows[0]["Work Email"] == "sam.lee@example.com"


@pytest.mark.asyncio
async def test_register_import_creates_and_updates(db_session: AsyncSession) -> None:
    clear_classification_config_cache()
    csv_bytes = _register_csv(
        [
            ["Alex Chen", "alex@example.com", "+61400000001", "Ops", "Pending verification"],
            ["Alex Chen", "alex@example.com", "+61400000099", "Engineering", "Active"],
        ]
    )
    rows = parse_employee_import_file(csv_bytes, "register.csv")

    result = await import_employee_masters(
        db_session, 1, mode="register", rows=rows, dry_run=False
    )
    assert result.created == 1
    assert result.updated == 1
    assert result.errors == []

    row = (
        await db_session.execute(
            select(EmployeeMasterRecord).where(
                EmployeeMasterRecord.tenant_id == 1,
                EmployeeMasterRecord.email == "alex@example.com",
            )
        )
    ).scalar_one()
    assert row.role == "Engineering"
    assert row.whatsapp_number == "+61400000099"
    assert row.status == "Active"


@pytest.mark.asyncio
async def test_payment_import_requires_existing_employee(db_session: AsyncSession) -> None:
    clear_classification_config_cache()
    register_rows = parse_employee_import_file(
        _register_csv([["Pat Jones", "pat@example.com", "", "HR", "Pending verification"]]),
        "register.csv",
    )
    await import_employee_masters(db_session, 1, mode="register", rows=register_rows)

    missing = await import_employee_masters(
        db_session,
        1,
        mode="payment",
        rows=parse_employee_import_file(
            _payment_xlsx([["ghost@example.com", "062-001", "12345678", "Ghost", "CBA", "", "", "", "Active"]]),
            "payment.xlsx",
        ),
        dry_run=False,
    )
    assert missing.skipped == 1
    assert missing.errors[0].message.startswith("no employee found")

    ok = await import_employee_masters(
        db_session,
        1,
        mode="payment",
        rows=parse_employee_import_file(
            _payment_xlsx(
                [["pat@example.com", "062-001", "99887766", "Pat Jones", "CBA", "1500", "4000", "12000", "Active"]]
            ),
            "payment.xlsx",
        ),
        dry_run=False,
    )
    assert ok.updated == 1
    row = (
        await db_session.execute(
            select(EmployeeMasterRecord).where(EmployeeMasterRecord.email == "pat@example.com")
        )
    ).scalar_one()
    assert row.bank["account_number"] == "99887766"
    assert row.budget["monthly"] == 1500
    assert row.status == "Active"


@pytest.mark.asyncio
async def test_register_dry_run_no_persist(db_session: AsyncSession) -> None:
    clear_classification_config_cache()
    rows = parse_employee_import_file(
        _register_csv([["Dry Run", "dry@example.com", "", "", "Pending verification"]]),
        "register.csv",
    )
    result = await import_employee_masters(db_session, 1, mode="register", rows=rows, dry_run=True)
    assert result.created == 1
    row = (
        await db_session.execute(
            select(EmployeeMasterRecord).where(EmployeeMasterRecord.email == "dry@example.com")
        )
    ).scalar_one_or_none()
    assert row is None


@pytest.mark.asyncio
async def test_build_import_templates() -> None:
    register = build_import_template("register")
    payment = build_import_template("payment")
    assert register[:2] == b"PK"
    assert payment[:2] == b"PK"

    register_wb = load_workbook(io.BytesIO(register), read_only=True)
    assert register_wb.sheetnames == ["Employee data", "Instructions"]
    payment_wb = load_workbook(io.BytesIO(payment), read_only=True)
    assert payment_wb.sheetnames == ["Employee data", "Instructions"]


@pytest.mark.asyncio
async def test_parse_styled_register_template() -> None:
    xlsx = build_import_template("register")
    rows = parse_employee_import_file(xlsx, "register.xlsx")
    assert rows == []


@pytest.mark.asyncio
async def test_employee_import_api(client: AsyncClient, db_session: AsyncSession) -> None:
    clear_classification_config_cache()

    res = await client.get("/api/employee-masters/import/templates/register")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]

    csv_bytes = _register_csv([["API User", "api.user@example.com", "+61422223333", "Finance", "Active"]])
    res = await client.post(
        "/api/employee-masters/import?mode=register&dry_run=true",
        files={"file": ("register.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    preview = res.json()["data"]
    assert preview["created"] == 1
    assert preview["dry_run"] is True

    res = await client.post(
        "/api/employee-masters/import?mode=register&dry_run=false",
        files={"file": ("register.csv", csv_bytes, "text/csv")},
    )
    assert res.status_code == 200
    assert res.json()["data"]["created"] == 1

    payment_bytes = _payment_xlsx(
        [["api.user@example.com", "062-000", "55443322", "API User", "NAB", "800", "", "", "Active"]]
    )
    res = await client.post(
        "/api/employee-masters/import?mode=payment&dry_run=false",
        files={"file": ("payment.xlsx", payment_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 200
    assert res.json()["data"]["updated"] == 1

    res = await client.get("/api/employee-masters")
    match = next(e for e in res.json()["data"] if e["email"] == "api.user@example.com")
    assert match["bank"]["account_number"] == "55443322"

    events = (
        await db_session.execute(
            select(AuditLog.event).where(AuditLog.event == "employee_import_completed")
        )
    ).scalars().all()
    assert events

    await client.delete(f"/api/employee-masters/{match['id']}")
