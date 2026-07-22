"""Unit and export tests for invoice-family vault labels on Documents Bundle sheet."""

from __future__ import annotations

import io
from datetime import date

import pytest
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.invoice import Invoice, InvoiceStatus
from app.services.dossier.vision_bundle_linkage import (
    persist_vision_bundle_snapshot,
    resolve_vision_bundle_key,
)
from app.services.extraction.document_heading_utils import is_invoice_family_vault_label
from app.services.reports.documents_bundle_export_service import (
    _HEADER_ROW,
    _SHEET_POSTING,
    _SHEET_UNLINKED,
    _UNLINKED_COLUMNS,
    build_documents_bundle_export,
    is_understood_invoice_family_anchor,
    vault_folder_label_for_export,
)
from app.tenant_ids import TESTING_TENANT_UUID


def test_proforma_invoice_no_label() -> None:
    from app.services.reports.documents_bundle_export_service import _proforma_invoice_no_label

    with_val = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        extracted_fields={"proforma_invoice_no": " PI-900 "},
    )
    assert _proforma_invoice_no_label(with_val) == "PI-900"
    empty = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        extracted_fields={},
    )
    assert _proforma_invoice_no_label(empty) == ""


@pytest.mark.asyncio
async def test_understood_bundle_includes_proforma_column(
    db_session: AsyncSession,
) -> None:
    inv = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="",
        invoice_date=date(2026, 7, 18),
        document_heading="TAX INVOICE",
        extracted_fields={
            "canonical_document_type": "Tax Invoice",
            "proforma_invoice_no": "PI-UB-77",
            "vision_bundle_kind": "proforma_invoice_no",
            "vision_bundle_key": "PI-UB-77",
        },
        route_target="Tax Invoice",
        raw_file_path="/tmp/tax-pi.pdf",
        file_hash="ub-tax-pi",
    )
    db_session.add(inv)
    await db_session.flush()
    persist_vision_bundle_snapshot(inv, resolve_vision_bundle_key(inv))
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 7, 1),
        date_to=date(2026, 7, 31),
    )
    header, data = _read_understood_sheet(payload.xlsx_bytes)
    assert "Proforma invoice no." in header
    assert "Invoice no." in header
    col = header.index("Proforma invoice no.")
    assert any(row[col] == "PI-UB-77" for row in data)


def test_invoice_family_vault_labels() -> None:
    assert is_invoice_family_vault_label("Tax Invoice") is True
    assert is_invoice_family_vault_label("Commercial Invoice") is True
    assert is_invoice_family_vault_label("Customer Invoice") is True
    assert is_invoice_family_vault_label("Sales Invoice") is True
    assert is_invoice_family_vault_label("Credit Note") is True
    assert is_invoice_family_vault_label("Debit Note") is True
    assert is_invoice_family_vault_label("Invoice") is True

    assert is_invoice_family_vault_label("Packing List") is False
    assert is_invoice_family_vault_label("Air Waybill") is False
    assert is_invoice_family_vault_label("Bill of Lading") is False
    assert is_invoice_family_vault_label("Purchase Order") is False
    assert is_invoice_family_vault_label("Proforma Invoice") is False
    assert is_invoice_family_vault_label("") is False


def test_vault_folder_label_prefers_canonical() -> None:
    inv = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        document_heading="TAX INVOICE",
        extracted_fields={"canonical_document_type": "Tax Invoice"},
        route_target="Tax Invoice",
    )
    assert vault_folder_label_for_export(inv) == "Tax Invoice"
    assert is_understood_invoice_family_anchor(inv) is True


def test_packing_list_not_understood_anchor() -> None:
    inv = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        document_heading="PACKING LIST",
        extracted_fields={"canonical_document_type": "Packing List"},
        route_target="Packing List",
    )
    assert vault_folder_label_for_export(inv) == "Packing List"
    assert is_understood_invoice_family_anchor(inv) is False


def _read_understood_sheet(content: bytes) -> tuple[list[str], list[list[object]]]:
    wb = load_workbook(io.BytesIO(content), data_only=False)
    assert _SHEET_POSTING in wb.sheetnames
    assert _SHEET_UNLINKED in wb.sheetnames
    ws = wb[_SHEET_POSTING]
    header = [
        ws.cell(row=_HEADER_ROW, column=col).value
        for col in range(1, ws.max_column + 1)
    ]
    assert header and header[0] is not None
    headers = [str(value) for value in header]
    data: list[list[object]] = []
    for row_idx in range(_HEADER_ROW + 1, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
        if all(value is None or value == "" for value in values):
            continue
        data.append(values)
    return headers, data


def _read_unlinked_sheet(content: bytes) -> tuple[list[str], list[list[object]]]:
    wb = load_workbook(io.BytesIO(content), data_only=False)
    ws = wb[_SHEET_UNLINKED]
    header = [
        ws.cell(row=_HEADER_ROW, column=col).value
        for col in range(1, ws.max_column + 1)
    ]
    assert header and header[0] is not None
    headers = [str(value) for value in header]
    data: list[list[object]] = []
    for row_idx in range(_HEADER_ROW + 1, ws.max_row + 1):
        values = [ws.cell(row=row_idx, column=col).value for col in range(1, ws.max_column + 1)]
        if all(value is None or value == "" for value in values):
            continue
        data.append(values)
    return headers, data


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def _invoice_no_from_cell(value: object) -> str:
    text = _cell_text(value)
    if " | " in text:
        return text.split(" | ", 1)[0]
    return text


@pytest.mark.asyncio
async def test_understood_bundle_sheet_invoice_row_with_siblings(
    db_session: AsyncSession,
) -> None:
    tax = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-1",
        invoice_date=date(2026, 7, 10),
        document_heading="TAX INVOICE",
        document_ref="DOC-TAX-1",
        extracted_fields={
            "canonical_document_type": "Tax Invoice",
            "vision_bundle_kind": "invoice_no",
            "vision_bundle_key": "INV-UB-1",
        },
        route_target="Tax Invoice",
        raw_file_path="/tmp/tax.pdf",
        file_hash="ub-tax-1",
    )
    packing = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-1",
        invoice_date=date(2026, 7, 10),
        document_heading="PACKING LIST",
        document_ref="DOC-PL-1",
        extracted_fields={
            "canonical_document_type": "Packing List",
            "vision_bundle_kind": "invoice_no",
            "vision_bundle_key": "INV-UB-1",
        },
        route_target="Packing List",
        raw_file_path="/tmp/pl.pdf",
        file_hash="ub-pl-1",
    )
    awb = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-1",
        invoice_date=date(2026, 7, 10),
        document_heading="AIR WAYBILL",
        document_ref="DOC-AWB-1",
        extracted_fields={
            "canonical_document_type": "Air Waybill",
            "vision_bundle_kind": "invoice_no",
            "vision_bundle_key": "INV-UB-1",
        },
        route_target="Air Waybill",
        raw_file_path="/tmp/awb.pdf",
        file_hash="ub-awb-1",
    )
    db_session.add_all([tax, packing, awb])
    await db_session.flush()
    for inv in (tax, packing, awb):
        persist_vision_bundle_snapshot(inv, resolve_vision_bundle_key(inv))
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 7, 1),
        date_to=date(2026, 7, 31),
    )
    assert payload.data_rows >= 1
    header, data = _read_understood_sheet(payload.xlsx_bytes)
    assert "Document number" in header
    assert header.index("Document number") + 1 == header.index("DT type")
    assert "Packing List" in header
    assert "Air Waybill" in header
    assert len(data) == 1
    row = data[0]
    assert _cell_text(row[header.index("Document number")]) == "DOC-TAX-1"
    assert _cell_text(row[header.index("DT type")]) == "Tax Invoice"
    pl_cell = _cell_text(row[header.index("Packing List")])
    awb_cell = _cell_text(row[header.index("Air Waybill")])
    assert pl_cell
    assert awb_cell
    assert "DOC-PL-1" in pl_cell or pl_cell.startswith("DOC-")
    assert "DOC-AWB-1" in awb_cell or awb_cell.startswith("DOC-")

    unlinked_header, unlinked_data = _read_unlinked_sheet(payload.xlsx_bytes)
    assert unlinked_header == list(_UNLINKED_COLUMNS)
    assert unlinked_data == []


@pytest.mark.asyncio
async def test_understood_bundle_excludes_packing_list_only(
    db_session: AsyncSession,
) -> None:
    packing = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Carrier",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-ONLY-PL",
        invoice_date=date(2026, 7, 12),
        document_heading="PACKING LIST",
        document_ref="DOC-PL-ONLY",
        extracted_fields={
            "canonical_document_type": "Packing List",
            "vision_bundle_kind": "invoice_no",
            "vision_bundle_key": "INV-UB-ONLY-PL",
        },
        route_target="Packing List",
        raw_file_path="/tmp/pl-only.pdf",
        file_hash="ub-pl-only",
    )
    db_session.add(packing)
    await db_session.flush()
    persist_vision_bundle_snapshot(packing, resolve_vision_bundle_key(packing))
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 7, 1),
        date_to=date(2026, 7, 31),
    )
    _header, data = _read_understood_sheet(payload.xlsx_bytes)
    assert data == []
    unlinked_header, unlinked_data = _read_unlinked_sheet(payload.xlsx_bytes)
    assert unlinked_header == list(_UNLINKED_COLUMNS)
    assert unlinked_header.index("Document number") + 1 == unlinked_header.index("DT type")
    assert len(unlinked_data) == 1
    assert _cell_text(unlinked_data[0][unlinked_header.index("Document number")]) == "DOC-PL-ONLY"
    assert _cell_text(unlinked_data[0][unlinked_header.index("DT type")]) == "Packing List"
    assert _invoice_no_from_cell(unlinked_data[0][unlinked_header.index("Invoice no.")]) == (
        "INV-UB-ONLY-PL"
    )


@pytest.mark.asyncio
async def test_understood_bundle_distinct_invoice_kinds(
    db_session: AsyncSession,
) -> None:
    tax = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Acme",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-TAX",
        invoice_date=date(2026, 7, 15),
        document_heading="TAX INVOICE",
        extracted_fields={"canonical_document_type": "Tax Invoice"},
        route_target="Tax Invoice",
        raw_file_path="/tmp/tax2.pdf",
        file_hash="ub-tax-2",
    )
    commercial = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        vendor="Beta",
        status=InvoiceStatus.EXCEPTION,
        evaluation_status="vision_vaulted",
        currency="AUD",
        invoice_no="INV-UB-COMM",
        invoice_date=date(2026, 7, 16),
        document_heading="COMMERCIAL INVOICE",
        extracted_fields={"canonical_document_type": "Commercial Invoice"},
        route_target="Commercial Invoice",
        raw_file_path="/tmp/comm.pdf",
        file_hash="ub-comm-1",
    )
    db_session.add_all([tax, commercial])
    await db_session.flush()
    for inv in (tax, commercial):
        persist_vision_bundle_snapshot(inv, resolve_vision_bundle_key(inv))
    await db_session.commit()

    payload = await build_documents_bundle_export(
        db_session,
        tenant_id=TESTING_TENANT_UUID,
        date_from=date(2026, 7, 1),
        date_to=date(2026, 7, 31),
    )
    header, data = _read_understood_sheet(payload.xlsx_bytes)
    types = {_cell_text(row[header.index("DT type")]) for row in data}
    assert "Tax Invoice" in types
    assert "Commercial Invoice" in types
    assert len(data) >= 2

    unlinked_header, unlinked_data = _read_unlinked_sheet(payload.xlsx_bytes)
    unlinked_types = {
        _cell_text(row[unlinked_header.index("DT type")]) for row in unlinked_data
    }
    assert "Tax Invoice" in unlinked_types
    assert "Commercial Invoice" in unlinked_types
