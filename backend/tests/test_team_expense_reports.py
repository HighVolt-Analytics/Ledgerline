"""Team expense reports: advance settlement, budget utilization, expense summary."""

from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest
from httpx import AsyncClient
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.invoice import Invoice, InvoiceStatus
from app.models.journal import EntryType, JournalEntry
from app.models.line_item import LineItem
from app.schemas.master_data import EmployeeMasterCreate
from app.schemas.rule_book_config import (
    TEAM_EXPENSE_KIND_ADVANCE,
    TEAM_EXPENSE_KIND_CLAIM,
    BudgetCategoryCap,
    ChartOfAccountEntry,
    EmployeeBudget,
    PostingDefaults,
    RuleBookConfigPayload,
    TeamExpensePostingDefaults,
    validate_rule_book_config_payload,
)
from app.services.master_data.master_data_service import create_employee_master
from app.services.master_data.party_coa_subledger_service import party_sub_ledger_code
from app.services.reports.team_expense_reports_service import (
    build_advance_settlement_rows,
    build_budget_utilization_rows,
    build_employee_expense_summary_rows,
)
from app.services.rule_book.rule_book_config_io import (
    load_rule_book_config_dict,
    save_rule_book_config,
)
from app.services.rule_book.rule_book_mapper import ROUTE_TEAM
from app.tenant_ids import TESTING_TENANT_UUID

EMPLOYEE_ID = "em-rpt-marcus"
EMPLOYEE_EMAIL = "marcus-rpt@example.com"


def _config() -> RuleBookConfigPayload:
    return RuleBookConfigPayload(
        posting_defaults=PostingDefaults(),
        team_expense_posting=TeamExpensePostingDefaults(
            default_advance_parent_ledger="Staff Advance",
            settlement_account="Bank Account",
        ),
        chart_of_accounts=[
            ChartOfAccountEntry(code="1000", name="Bank Account", type="Asset"),
            ChartOfAccountEntry(code="1300", name="Staff Advance", type="Asset"),
            ChartOfAccountEntry(code="5000", name="Operating Expenses", type="Expense"),
            ChartOfAccountEntry(code="2000", name="Accounts Payable", type="Liability"),
        ],
    )


async def _setup_employee(
    db_session: AsyncSession,
    *,
    mtd: float = 100.0,
    qtd: float = 200.0,
    ytd: float = 500.0,
    claim_count: int = 2,
) -> None:
    await save_rule_book_config(db_session, _config(), TESTING_TENANT_UUID)
    await create_employee_master(
        db_session,
        TESTING_TENANT_UUID,
        EmployeeMasterCreate(
            master_id=EMPLOYEE_ID,
            name="Marcus Webb",
            email=EMPLOYEE_EMAIL,
            whatsapp_number="+6591110001",
            division="Ops",
            location="Singapore",
            department="Ops",
            advance_parent_ledger="Staff Advance",
            budget=EmployeeBudget(
                monthly=1000,
                quarterly=3000,
                annual=12000,
                categories=[BudgetCategoryCap(ledger="Travel", cap=400)],
            ),
            mtd_spent=mtd,
            qtd_spent=qtd,
            ytd_spent=ytd,
            claim_count=claim_count,
            last_claim="2026-05-01",
        ),
    )


@pytest.mark.asyncio
async def test_advance_settlement_ledger_pending_and_available(
    db_session: AsyncSession,
) -> None:
    await _setup_employee(db_session)
    raw = await load_rule_book_config_dict(db_session, TESTING_TENANT_UUID)
    validate_rule_book_config_payload(raw)
    code = party_sub_ledger_code(EMPLOYEE_ID)

    processed_advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 5, 1),
        total=Decimal("1000"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-adv-1",
    )
    db_session.add(processed_advance)
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=processed_advance.id,
            date=date(2026, 5, 1),
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("1000"),
            credit=Decimal("0"),
            entry_type=EntryType.DEBIT,
        )
    )

    # Settlement credit on the employee child (historical clear / repayment).
    cleared = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 5, 2),
        total=Decimal("400"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-adv-2",
    )
    db_session.add(cleared)
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=cleared.id,
            date=date(2026, 5, 2),
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("0"),
            credit=Decimal("400"),
            entry_type=EntryType.CREDIT,
        )
    )

    # Open expense claims reserve Staff Advance float until posted.
    open_claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        currency="SGD",
        invoice_date=date(2026, 5, 3),
        total=Decimal("150"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-adv-3",
    )
    db_session.add(open_claim)
    await db_session.flush()

    rows = await build_advance_settlement_rows(db_session, TESTING_TENANT_UUID)
    row = next(r for r in rows if r.employee_id == EMPLOYEE_ID)
    assert row.advance_ledger_balance == Decimal("600")
    assert row.advance_taken == Decimal("1000")
    assert row.advance_used == Decimal("400")
    assert row.pending_against_advance == Decimal("150")
    assert row.available_advance == Decimal("450")
    assert row.claim_ytd_spent == 500.0
    assert row.name == "Marcus Webb"


@pytest.mark.asyncio
async def test_budget_utilization_math(db_session: AsyncSession) -> None:
    await _setup_employee(db_session, mtd=0, qtd=0, ytd=0, claim_count=5)
    today = date.today()
    db_session.add(
        Invoice(
            tenant_id=TESTING_TENANT_UUID,
            status=InvoiceStatus.PROCESSED,
            currency="SGD",
            invoice_date=today,
            total=Decimal("250"),
            route_target=ROUTE_TEAM,
            team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
            employee_email=EMPLOYEE_EMAIL,
            email_sender=EMPLOYEE_EMAIL,
            file_hash="rpt-budget-mtd",
        )
    )
    await db_session.flush()
    rows = await build_budget_utilization_rows(db_session, TESTING_TENANT_UUID)
    row = next(r for r in rows if r.employee_id == EMPLOYEE_ID)
    assert row.budget_monthly == 1000
    assert row.mtd_spent == 250
    assert row.monthly_remaining == 750
    assert row.monthly_utilization_pct == 25.0
    assert row.quarterly_remaining == 2750
    assert row.annual_remaining == 11750
    assert "Travel:400" in row.category_caps
    assert row.advance_float == 0
    assert row.monthly_cash_committed == 250
    assert row.monthly_cash_remaining == 750


@pytest.mark.asyncio
async def test_budget_utilization_cash_reserves_advance_float(
    db_session: AsyncSession,
) -> None:
    """Accrual remaining ignores advances; cash remaining reserves outstanding float."""
    await _setup_employee(db_session, mtd=0, qtd=0, ytd=0, claim_count=1)
    raw = await load_rule_book_config_dict(db_session, TESTING_TENANT_UUID)
    validate_rule_book_config_payload(raw)
    code = party_sub_ledger_code(EMPLOYEE_ID)
    today = date.today()

    claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("200"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-cash-claim",
    )
    advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("300"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-cash-adv",
    )
    db_session.add_all([claim, advance])
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=advance.id,
            date=today,
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("300"),
            credit=Decimal("0"),
            entry_type=EntryType.DEBIT,
        )
    )
    await db_session.flush()

    rows = await build_budget_utilization_rows(db_session, TESTING_TENANT_UUID)
    row = next(r for r in rows if r.employee_id == EMPLOYEE_ID)
    assert row.mtd_spent == 200
    assert row.monthly_remaining == 800
    assert row.advance_float == 300
    assert row.monthly_cash_committed == 500
    assert row.monthly_cash_remaining == 500
    assert row.monthly_cash_utilization_pct == 50.0
    # Advances still do not inflate accrual spend.
    assert row.mtd_spent == 200


@pytest.mark.asyncio
async def test_expense_summary_lines_employee_and_date_filter(
    db_session: AsyncSession,
) -> None:
    await _setup_employee(db_session)

    in_range = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 5, 10),
        invoice_no="CLM-10",
        document_ref="DOC-10",
        total=Decimal("80"),
        account_code="5000",
        account_name="Operating Expenses",
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        document_type_code="DT-08",
        email_sender=EMPLOYEE_EMAIL,
        employee_email=EMPLOYEE_EMAIL,
        evaluation_status="auto_coded",
        vendor="Grab",
        file_hash="rpt-sum-1",
    )
    out_of_range = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.EXCEPTION,
        currency="SGD",
        invoice_date=date(2026, 4, 1),
        invoice_no="CLM-01",
        total=Decimal("20"),
        account_code="5000",
        account_name="Operating Expenses",
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-sum-2",
    )
    db_session.add_all([in_range, out_of_range])
    await db_session.flush()
    db_session.add(
        LineItem(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=in_range.id,
            description="Airport transfer",
            qty=Decimal("1"),
            amount=Decimal("80"),
            sub_ledger="Travel",
        )
    )
    await db_session.flush()

    rows = await build_employee_expense_summary_rows(
        db_session,
        TESTING_TENANT_UUID,
        date_from=date(2026, 5, 1),
        date_to=date(2026, 5, 31),
    )
    assert len(rows) == 1
    row = rows[0]
    assert row.employee_id == EMPLOYEE_ID
    assert row.employee_name == "Marcus Webb"
    assert row.employee_email == EMPLOYEE_EMAIL
    assert row.mobile == "+6591110001"
    assert row.department == "Ops"
    assert row.division == "Ops"
    assert row.location == "Singapore"
    assert row.document_no == "DOC-10"
    assert row.line_description == "Airport transfer"
    assert row.line_amount == Decimal("80")
    assert row.ledger_code == "5000"
    assert row.main_gl == "Operating Expenses"
    assert row.sub_ledger == "Travel"
    assert row.status == InvoiceStatus.PROCESSED.value
    assert row.team_expense_kind == TEAM_EXPENSE_KIND_CLAIM


@pytest.mark.asyncio
async def test_gl_account_budget_utilization(
    db_session: AsyncSession,
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.services.master_data.department_budget_service import (
        build_department_budget_utilization_rows,
    )
    from app.services.purchase.team_expense_spend_service import current_period_keys

    await _setup_employee(db_session, mtd=0, qtd=0, ytd=0)
    today = date.today()
    keys = current_period_keys(today)

    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            department="",
            gl_ledger="Travel",
            period_kind="monthly",
            period_key=keys["monthly"],
            allocated=Decimal("1000"),
        )
    )
    claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("100"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-gl-claim",
        account_name="Travel",
    )
    # Advances must not count toward GL spend.
    advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("400"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="rpt-gl-adv",
        account_name="Travel",
    )
    db_session.add_all([claim, advance])
    await db_session.flush()

    rows = await build_department_budget_utilization_rows(db_session, TESTING_TENANT_UUID)
    row = next(r for r in rows if r.gl_ledger == "Travel")
    assert row.allocated == 1000
    assert row.consumed == 100
    assert row.remaining == 900
    assert row.utilization_pct == 10.0


@pytest.mark.asyncio
async def test_team_expense_report_apis(client: AsyncClient, db_session: AsyncSession) -> None:
    await _setup_employee(db_session)
    await db_session.flush()

    adv = await client.get("/api/reports/team-expenses/advance-settlement")
    assert adv.status_code == 200
    adv_rows = adv.json()["data"]
    assert any(r["employee_id"] == EMPLOYEE_ID for r in adv_rows)

    bud = await client.get("/api/reports/team-expenses/budget-utilization")
    assert bud.status_code == 200
    bud_rows = bud.json()["data"]
    marcus = next(r for r in bud_rows if r["employee_id"] == EMPLOYEE_ID)
    assert marcus["budget_monthly"] == 1000
    assert marcus["mtd_spent"] == 0
    assert "Travel:400" in marcus["category_caps"]
    assert marcus["advance_float"] == 0
    assert marcus["monthly_cash_committed"] == 0
    assert marcus["monthly_cash_remaining"] == 1000

    dept = await client.get("/api/reports/team-expenses/department-budget-utilization")
    assert dept.status_code == 200
    assert isinstance(dept.json()["data"], list)

    summary = await client.get(
        "/api/reports/team-expenses/expense-summary?date_from=2026-01-01&date_to=2026-12-31"
    )
    assert summary.status_code == 200
    assert isinstance(summary.json()["data"], list)


@pytest.mark.asyncio
async def test_team_expense_gl_budget_excel_export_has_readable_headers(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from openpyxl import load_workbook
    from io import BytesIO

    await _setup_employee(db_session)
    await db_session.flush()

    res = await client.get(
        "/api/reports/team-expenses/department-budget-utilization/export"
    )
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in (res.headers.get("content-type") or "")
    wb = load_workbook(BytesIO(res.content))
    ws = wb["GL Budgets"]
    headers = [cell.value for cell in ws[4]]
    assert "Parent GL" in headers
    assert "Budget" in headers
    assert "Enforcement" in headers


@pytest.mark.asyncio
async def test_employee_spend_detail_one_row_per_employee_sub_gl(
    db_session: AsyncSession,
) -> None:
    from app.models.department_budget import DepartmentBudget
    from app.schemas.rule_book_config import ChartOfAccountEntry, SubLedgerEntry
    from app.services.purchase.team_expense_spend_service import current_period_keys
    from app.services.reports.team_expense_reports_service import (
        build_employee_spend_detail_rows,
    )

    cfg = _config()
    cfg.chart_of_accounts = [
        ChartOfAccountEntry(code="1000", name="Bank Account", type="Asset"),
        ChartOfAccountEntry(code="1300", name="Staff Advance", type="Asset"),
        ChartOfAccountEntry(
            code="6000",
            name="Marketing Expenses",
            type="Expense",
            sub_ledgers=[
                SubLedgerEntry(code="6001", name="Traveling"),
                SubLedgerEntry(code="6002", name="Hotel"),
            ],
        ),
        ChartOfAccountEntry(code="2000", name="Accounts Payable", type="Liability"),
    ]
    await save_rule_book_config(db_session, cfg, TESTING_TENANT_UUID)
    await create_employee_master(
        db_session,
        TESTING_TENANT_UUID,
        EmployeeMasterCreate(
            master_id=EMPLOYEE_ID,
            name="Marcus Webb",
            email=EMPLOYEE_EMAIL,
            role="Marketing Executive",
            whatsapp_number="+6591110001",
            whatsapp_number_2="+6591110002",
            department="Marketing",
            division="B2B Sales",
            location="Singapore",
            supervisor_1="Priya Sharma",
            supervisor_2="Amit Rao",
            advance_parent_ledger="Staff Advance",
            status="Active",
            budget=EmployeeBudget(
                monthly=120000,
                quarterly=300000,
                annual=500000,
            ),
        ),
    )

    today = date.today()
    keys = current_period_keys(today)
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="Traveling",
            period_kind="annual",
            period_key=keys["annual"],
            allocated=Decimal("200000"),
        )
    )
    db_session.add(
        DepartmentBudget(
            tenant_id=TESTING_TENANT_UUID,
            gl_ledger="Hotel",
            period_kind="annual",
            period_key=keys["annual"],
            allocated=Decimal("100000"),
        )
    )

    travel_claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("80000"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="spend-travel",
        account_name="Traveling",
    )
    hotel_claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("20000"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="spend-hotel",
        account_name="Hotel",
    )
    # Advance must not create a spend-detail row.
    advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=today,
        total=Decimal("5000"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="spend-adv",
        account_name="Traveling",
    )
    db_session.add_all([travel_claim, hotel_claim, advance])
    await db_session.flush()

    # Cash reimbursed = settlement credit on the travel claim only.
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=travel_claim.id,
            date=today,
            account_code="1000",
            account_name="Bank Account",
            debit=Decimal("0"),
            credit=Decimal("75000"),
            entry_type=EntryType.CREDIT,
        )
    )
    await db_session.flush()

    rows = await build_employee_spend_detail_rows(db_session, TESTING_TENANT_UUID)
    assert len(rows) == 2
    by_sub = {r.sub_ledger: r for r in rows}

    travel = by_sub["Traveling"]
    assert travel.employee_id == EMPLOYEE_ID
    assert travel.name == "Marcus Webb"
    assert travel.role == "Marketing Executive"
    assert travel.whatsapp_number == "+6591110001"
    assert travel.department == "Marketing"
    assert travel.supervisor_1 == "Priya Sharma"
    assert travel.main_gl == "Marketing Expenses"
    assert travel.sub_gl_budget == 200000
    assert travel.employee_spend_ytd == 80000
    assert travel.pct_of_sub_gl_used == 40.0
    assert travel.claim_count == 1
    assert travel.cash_reimbursed_ytd == 75000
    assert travel.last_claim_date == today.isoformat()
    # Employee-level spending limits (same on every Sub-GL row for this employee).
    assert travel.budget_monthly == 120000
    assert travel.budget_quarterly == 300000
    assert travel.budget_annual == 500000
    assert travel.mtd_spent == 100000  # 80k Traveling + 20k Hotel
    assert travel.qtd_spent == 100000
    assert travel.ytd_spent_total == 100000
    assert travel.monthly_remaining == 20000
    assert travel.monthly_utilization_pct == round((100000 / 120000) * 100, 2)
    assert travel.annual_remaining == 400000
    assert travel.annual_utilization_pct == 20.0

    hotel = by_sub["Hotel"]
    assert hotel.main_gl == "Marketing Expenses"
    assert hotel.sub_gl_budget == 100000
    assert hotel.employee_spend_ytd == 20000
    assert hotel.pct_of_sub_gl_used == 20.0
    assert hotel.cash_reimbursed_ytd == 0
    assert hotel.mtd_spent == 100000
    assert hotel.budget_annual == 500000


@pytest.mark.asyncio
async def test_employee_spend_detail_excel_has_master_and_spend_sheets(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    await _setup_employee(db_session)
    await db_session.flush()

    res = await client.get("/api/reports/team-expenses/employee-spend-detail/export")
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Employee Spend Detail"]

    spend_headers = [c.value for c in wb["Employee Spend Detail"][4]]
    assert spend_headers[0] == "Employee ID"
    assert spend_headers[21] == "Employee Status"
    assert spend_headers[22:31] == [
        "Main GL",
        "Sub-Ledger",
        "Sub-GL Budget",
        "Employee Spend (YTD)",
        "% of Sub-GL Used by Employee",
        "No. of Claims",
        "Advance Pending",
        "Cash Reimbursed YTD",
        "Last Claim Date",
    ]
    assert spend_headers[31:43] == [
        "Monthly Spending Limit",
        "Monthly Spent",
        "Monthly Remaining",
        "Monthly Utilization %",
        "Quarterly Spending Limit",
        "Quarterly Spent",
        "Quarterly Remaining",
        "Quarterly Utilization %",
        "Annual Spending Limit",
        "Annual Spent",
        "Annual Remaining",
        "Annual Utilization %",
    ]
    # Header row only — continuous horizontal scroll (no column freeze).
    assert wb["Employee Spend Detail"].freeze_panes == "A5"


@pytest.mark.asyncio
async def test_employee_advance_detail_take_and_spend_rows(
    db_session: AsyncSession,
) -> None:
    from app.models.audit import AuditLog
    from app.services.reports.team_expense_reports_service import (
        build_employee_advance_detail_rows,
    )

    await _setup_employee(db_session)
    raw = await load_rule_book_config_dict(db_session, TESTING_TENANT_UUID)
    validate_rule_book_config_payload(raw)
    code = party_sub_ledger_code(EMPLOYEE_ID)

    advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 7, 1),
        invoice_no="ADV-100",
        total=Decimal("500"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="adv-detail-1",
    )
    db_session.add(advance)
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=advance.id,
            date=date(2026, 7, 1),
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("500"),
            credit=Decimal("0"),
            entry_type=EntryType.DEBIT,
        )
    )
    db_session.add(
        AuditLog(
            event="invoice_approved",
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=advance.id,
            detail={"actor_name": "Priya Sharma", "actor_email": "priya@example.com"},
        )
    )

    claim = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 7, 15),
        invoice_no="CLM-200",
        total=Decimal("300"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_CLAIM,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="adv-detail-2",
        account_name="Operating Expenses",
    )
    db_session.add(claim)
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=claim.id,
            date=date(2026, 7, 15),
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("0"),
            credit=Decimal("300"),
            entry_type=EntryType.CREDIT,
        )
    )
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=claim.id,
            date=date(2026, 7, 15),
            account_code="1000",
            account_name="Bank Account",
            debit=Decimal("0"),
            credit=Decimal("0"),  # fully netted — no cash; still a Used row
            entry_type=EntryType.CREDIT,
        )
    )
    await db_session.flush()

    rows = await build_employee_advance_detail_rows(db_session, TESTING_TENANT_UUID)
    assert len(rows) == 2
    assert rows[0].movement_type == "Advance"
    assert rows[0].took == Decimal("500")
    assert rows[0].used == Decimal("0")
    assert rows[0].outstanding_after == Decimal("500")
    assert rows[0].approved_by == "Priya Sharma"
    assert rows[0].document_date == date(2026, 7, 1)

    assert rows[1].movement_type == "Claim"
    assert rows[1].took == Decimal("0")
    assert rows[1].used == Decimal("300")
    assert rows[1].outstanding_after == Decimal("200")
    assert rows[1].available == Decimal("200")
    assert rows[1].pending_claims == Decimal("0")


@pytest.mark.asyncio
async def test_employee_advance_detail_excel_headers(
    client: AsyncClient, db_session: AsyncSession
) -> None:
    from io import BytesIO

    from openpyxl import load_workbook

    await _setup_employee(db_session)
    code = party_sub_ledger_code(EMPLOYEE_ID)
    advance = Invoice(
        tenant_id=TESTING_TENANT_UUID,
        status=InvoiceStatus.PROCESSED,
        currency="SGD",
        invoice_date=date(2026, 8, 1),
        total=Decimal("100"),
        route_target=ROUTE_TEAM,
        team_expense_kind=TEAM_EXPENSE_KIND_ADVANCE,
        employee_email=EMPLOYEE_EMAIL,
        email_sender=EMPLOYEE_EMAIL,
        file_hash="adv-xlsx-1",
    )
    db_session.add(advance)
    await db_session.flush()
    db_session.add(
        JournalEntry(
            tenant_id=TESTING_TENANT_UUID,
            invoice_id=advance.id,
            date=date(2026, 8, 1),
            account_code=code,
            account_name="Marcus Webb",
            debit=Decimal("100"),
            credit=Decimal("0"),
            entry_type=EntryType.DEBIT,
        )
    )
    await db_session.flush()

    res = await client.get("/api/reports/team-expenses/employee-advance-detail/export")
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["Employee Advance Detail"]
    headers = [c.value for c in wb["Employee Advance Detail"][4]]
    assert headers[0] == "Employee ID"
    assert headers[21] == "Employee Status"
    assert headers[22:34] == [
        "Movement Type",
        "Document No.",
        "Document Date",
        "Took",
        "Used",
        "Outstanding After",
        "Pending Claims",
        "Available",
        "Cash Reimbursed",
        "Document Status",
        "Approved By",
        "Approved On",
    ]

